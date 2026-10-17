"""ร่างบันทึกความเห็นทางกฎหมายเบื้องต้น (.docx) จากสถานการณ์ที่ผู้ใช้เล่าในแชท — แยกจากระบบแก้ไฟล์ Excel ทั้งหมด
(ไม่แตะ EditableDocument/DocumentVersion/focus เลย) รองรับกฎหมายไทยทุกหมวด ไม่จำกัดแค่ธุรกิจต่างด้าว/ภาษี

ขั้นตอน: (1) Claude ดึง slots จากข้อความ+ประวัติ ไม่เดาข้อเท็จจริงที่ไม่ได้พูดถึง (2) ค้น legal_retrieval.legal_search
หา chunk ที่เกี่ยวข้อง (3) Claude เรียบเรียงเนื้อหาบันทึกเป็น JSON มีโครงสร้าง อ้าง chunk_ids ได้ (4) โค้ดตรวจ
chunk_ids ให้อยู่ในชุดที่ค้นเจอจริงเท่านั้น ใส่ป้าย 📚/💡 เอง แล้วสร้างไฟล์ .docx ด้วย python-docx ล้วนๆ

import จาก main.py แบบ deferred (หลังนิยาม client/ROUTER_MODEL/ADVISOR_MODEL/save_legal_memo) เหมือน
legal_retrieval.py — ดู main.py บรรทัดที่ import โมดูลนี้"""
import io
import datetime

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import legal_retrieval
from llm_tools import call_with_required_tool
from main import client, get_chat_messages, save_legal_memo, ROUTER_MODEL, ADVISOR_MODEL

MEMO_FONT_NAME = "TH Sarabun New"
MEMO_BODY_SIZE_PT = 16
MEMO_HEADING_SIZE_PT = 20
_MISSING_PLACEHOLDER = "............"

MEMO_SLOTS_TOOL = {
    "name": "extract_memo_slots",
    "description": (
        "ดึงข้อมูลสถานการณ์ของผู้ใช้จากข้อความล่าสุด+ประวัติสนทนา เพื่อร่างบันทึกความเห็นทางกฎหมาย ห้ามเดา/แต่ง"
        "ข้อมูลที่ผู้ใช้ไม่ได้พูดถึงเด็ดขาด ไม่แน่ใจให้ใส่ null เสมอ ต้องเรียกเครื่องมือนี้เสมอ"
    ),
    "input_schema": {
        "type": "object",
        "properties": {
            "client_name": {"type": ["string", "null"], "description": "ชื่อลูกค้า/ผู้ขอคำปรึกษา ไม่ได้บอกใส่ null"},
            "legal_area": {
                "type": ["string", "null"],
                "description": "หมวดกฎหมายที่เกี่ยวข้อง เช่น ธุรกิจต่างด้าว/ภาษี/ที่ดิน/แรงงาน/สัญญา/ครอบครัว ฯลฯ ไม่ชัดเจนใส่ null",
            },
            "facts_summary": {
                "type": ["string", "null"],
                "description": "สรุปสถานการณ์/ข้อเท็จจริงทั้งหมดที่ผู้ใช้เล่ามาเป็นย่อหน้าเดียว ไม่มีเลยใส่ null",
            },
            "key_facts": {
                "type": ["array", "null"],
                "description": "ข้อเท็จจริงสำคัญแยกเป็นข้อสั้นๆ (สัดส่วนหุ้น/ประเภทธุรกิจ/จำนวนเงิน/วันที่ ฯลฯ) ไม่มีใส่ null",
                "items": {"type": "string"},
            },
            "user_questions": {
                "type": ["array", "null"],
                "description": "คำถามที่ผู้ใช้ถามไว้ตรงๆ (ถ้ามี) ไม่มีใส่ null",
                "items": {"type": "string"},
            },
        },
        "required": ["client_name", "legal_area", "facts_summary", "key_facts", "user_questions"],
    },
}

_MEMO_PARAGRAPH_ITEM_SCHEMA = {
    "type": "object",
    "properties": {
        "text": {"type": "string"},
        "chunk_ids": {
            "type": "array", "items": {"type": "integer"},
            "description": "id ของ chunk จากฐานความรู้ที่ใช้เขียนย่อหน้านี้จริง (จากรายการที่ให้มา) ไม่มีใส่ []",
        },
    },
    "required": ["text", "chunk_ids"],
}

MEMO_CONTENT_TOOL = {
    "name": "draft_memo_content",
    "description": (
        "เรียบเรียงเนื้อหาบันทึกความเห็นทางกฎหมายเบื้องต้นจากสถานการณ์และ chunk ฐานความรู้ที่ให้มา รองรับกฎหมายไทย"
        "ทุกหมวด ไม่จำกัดเฉพาะหมวดใดหมวดหนึ่ง ห้ามเขียนทำนอง \"นอกเหนือความเชี่ยวชาญ\" ต้องเรียกเครื่องมือนี้เสมอ"
    ),
    "input_schema": {
        "type": "object",
        "properties": {
            "issues": {"type": "array", "items": _MEMO_PARAGRAPH_ITEM_SCHEMA},
            "analysis": {"type": "array", "items": _MEMO_PARAGRAPH_ITEM_SCHEMA},
            "conclusion": _MEMO_PARAGRAPH_ITEM_SCHEMA,
            "next_steps": {"type": "array", "items": {"type": "string"}},
        },
        "required": ["issues", "analysis", "conclusion", "next_steps"],
    },
}

MEMO_CONTENT_SYSTEM_PROMPT = (
    "คุณกำลังเรียบเรียงบันทึกความเห็นทางกฎหมายเบื้องต้นให้ทนายความใช้ประกอบการให้คำปรึกษา รองรับกฎหมายไทยทุกหมวด "
    "(ธุรกิจต่างด้าว/ภาษี/ที่ดิน/แรงงาน/สัญญา/ครอบครัว ฯลฯ) ไม่จำกัดเฉพาะหมวดใดหมวดหนึ่ง ห้ามเขียนว่าเรื่องนี้ "
    "\"นอกเหนือความเชี่ยวชาญ\" เด็ดขาด ไม่ว่าจะเป็นกฎหมายหมวดใด\n"
    "- ได้รับ chunk จากฐานความรู้มาพร้อม id — ถ้า chunk ตรงประเด็น ให้ใช้เป็นหลักและใส่ id นั้นใน chunk_ids ของ"
    "ย่อหน้าที่ใช้ข้อมูลนั้นจริง\n"
    "- chunk ที่ให้มาไม่ครอบคลุมบางประเด็น ใช้ความรู้ทั่วไปของคุณเขียนเสริมได้ (ใส่ chunk_ids เป็น [] สำหรับย่อหน้า"
    "นั้น) ระบุชื่อกฎหมาย/หน่วยงานให้ถูกต้องเท่าที่มั่นใจ\n"
    "- ไม่แน่ใจตัวเลข/อัตรา/มาตราที่แน่นอน ให้เขียนเป็นภาพรวมแทนการระบุตัวเลขที่ไม่แน่ใจ ห้ามแต่งตัวเลข/มาตราขึ้นมา\n"
    "- ห้ามเดา/แต่งข้อเท็จจริงของผู้ใช้ที่ไม่ได้อยู่ในสถานการณ์ที่ให้มาเด็ดขาด\n"
    "- issues: ประเด็นที่ต้องพิจารณาแต่ละข้อ, analysis: การวิเคราะห์ทางกฎหมายแต่ละประเด็น, conclusion: ความเห็น"
    "เบื้องต้นสรุปรวม 1 รายการ, next_steps: ขั้นตอนที่แนะนำให้ทำต่อ (ข้อความล้วน ไม่ต้องมี chunk_ids)"
)


def _extract_memo_slots(query: str, history: list) -> dict:
    """เรียก Claude (ROUTER_MODEL — งานดึงค่าโครงสร้าง ไม่ต้องเขียนประโยคยาว) ดึง slots จากข้อความ+ประวัติ 6
    ข้อความล่าสุด คืน dict ตาม MEMO_SLOTS_TOOL เสมอ (ทุกช่อง None ถ้า parse/เรียกไม่สำเร็จ)"""
    history_lines = [f"{m['role']}: {(m.get('content') or '')[:500]}" for m in (history or [])[-6:]]
    history_text = ("ประวัติสนทนาล่าสุด:\n" + "\n".join(history_lines) + "\n\n") if history_lines else ""
    prompt = f"{history_text}ข้อความล่าสุดของผู้ใช้: {query}"

    response = call_with_required_tool(
        client, ROUTER_MODEL, 800, "", [{"role": "user", "content": prompt}],
        [MEMO_SLOTS_TOOL], "extract_memo_slots",
    )
    if response is None:
        return {"client_name": None, "legal_area": None, "facts_summary": None, "key_facts": None, "user_questions": None}
    tool_use_block = next(b for b in response.content if b.type == "tool_use" and b.name == "extract_memo_slots")
    return tool_use_block.input


def _format_chunks_for_prompt(chunks: list[dict]) -> str:
    if not chunks:
        return "(ไม่มี chunk ที่เกี่ยวข้องพอ — ใช้ความรู้ทั่วไปทั้งหมด)"
    lines = []
    for c in chunks:
        tags_text = "/".join(c.get("tags") or []) or "-"
        lines.append(f"[id={c['id']}] (tag: {tags_text}) {c['content']}")
    return "\n\n".join(lines)


def _draft_memo_content(slots: dict, grounded_chunks: list[dict]) -> dict:
    """เรียก Claude (ADVISOR_MODEL — ต้องการภาษาไทยที่แม่นยำในคำตอบยาว เหมือน advisor_answer) เรียบเรียงเนื้อหา
    บันทึก คืน dict ตาม MEMO_CONTENT_TOOL เสมอ (list ว่างถ้า parse/เรียกไม่สำเร็จ)"""
    facts_lines = "\n".join(f"- {f}" for f in (slots.get("key_facts") or []))
    questions_lines = "\n".join(f"- {q}" for q in (slots.get("user_questions") or []))
    prompt = (
        f"หมวดกฎหมาย: {slots.get('legal_area') or 'ไม่ระบุ'}\n\n"
        f"สถานการณ์: {slots.get('facts_summary')}\n\n"
        f"ข้อเท็จจริงสำคัญ:\n{facts_lines or '(ไม่มี)'}\n\n"
        f"คำถามของผู้ใช้:\n{questions_lines or '(ไม่มี)'}\n\n"
        f"chunk จากฐานความรู้:\n{_format_chunks_for_prompt(grounded_chunks)}"
    )
    response = call_with_required_tool(
        client, ADVISOR_MODEL, 2500, MEMO_CONTENT_SYSTEM_PROMPT, [{"role": "user", "content": prompt}],
        [MEMO_CONTENT_TOOL], "draft_memo_content",
    )
    if response is None:
        return {"issues": [], "analysis": [], "conclusion": {"text": "", "chunk_ids": []}, "next_steps": []}
    tool_use_block = next(b for b in response.content if b.type == "tool_use" and b.name == "draft_memo_content")
    return tool_use_block.input


def _label_for_chunk_ids(chunk_ids: list[int]) -> str:
    """ป้ายที่โค้ดใส่เอง (ไม่ใช่ LLM เขียน) — มี chunk_ids (ตรวจผ่านว่าอยู่ในชุด grounded จริงแล้ว) -> อ้างอิงฐาน
    ความรู้ ไม่มี -> ความรู้ทั่วไป"""
    if chunk_ids:
        return "📚 ฐานความรู้ " + " ".join(f"#{cid}" for cid in chunk_ids)
    return "💡 ความรู้ทั่วไป"


def _render_paragraph(item: dict, grounded_ids: set[int]) -> dict:
    """กรอง chunk_ids ให้เหลือเฉพาะที่อยู่ในชุด grounded จริง (ตัดที่ LLM อ้างมั่วทิ้ง) แล้วคำนวณป้ายจากโค้ด"""
    valid_ids = [cid for cid in (item.get("chunk_ids") or []) if cid in grounded_ids]
    return {"text": item.get("text") or "", "chunk_ids": valid_ids, "label": _label_for_chunk_ids(valid_ids)}


def _missing_slot_labels(slots: dict) -> list[str]:
    labels = []
    if not slots.get("client_name"):
        labels.append("ชื่อลูกค้า")
    if not slots.get("legal_area"):
        labels.append("หมวดกฎหมาย")
    if not slots.get("key_facts"):
        labels.append("ข้อเท็จจริงสำคัญแบบละเอียด")
    if not slots.get("user_questions"):
        labels.append("คำถามเฉพาะเจาะจงของผู้ใช้")
    return labels


def _thai_date_today() -> str:
    thai_months = (
        "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
        "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม",
    )
    today = datetime.date.today()
    return f"{today.day} {thai_months[today.month - 1]} {today.year + 543}"


def _set_thai_run_font(run, size_pt: int = MEMO_BODY_SIZE_PT, bold: bool = False) -> None:
    """ตั้งฟอนต์ไทยให้ run ครบทั้ง 4 ตำแหน่ง (ascii/hAnsi/eastAsia/cs) — python-docx เดียวไม่พอกับฟอนต์ไทย ต้องลง
    XML (w:rFonts) ตรงๆ ไม่งั้น Word เลือกฟอนต์อื่นแทนที่ eastAsia/cs เงียบๆ"""
    run.font.name = MEMO_FONT_NAME
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), MEMO_FONT_NAME)


def build_memo_docx(slots: dict, rendered: dict, used_chunks: list[dict]) -> bytes:
    """สร้างไฟล์ .docx ด้วย python-docx ล้วนๆ จาก slots (ข้อเท็จจริง) + rendered (ย่อหน้าที่กรอง/ใส่ป้ายแล้ว) +
    used_chunks (chunk ที่ถูกอ้างถึงจริงอย่างน้อย 1 ย่อหน้า สำหรับส่วนแหล่งอ้างอิง)"""
    doc = Document()
    normal_font = doc.styles["Normal"].font
    normal_font.name = MEMO_FONT_NAME
    normal_font.size = Pt(MEMO_BODY_SIZE_PT)

    def add_heading(text: str):
        p = doc.add_paragraph()
        _set_thai_run_font(p.add_run(text), bold=True)
        return p

    def add_body(text: str):
        p = doc.add_paragraph()
        _set_thai_run_font(p.add_run(text))
        return p

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_thai_run_font(title_p.add_run("บันทึกความเห็นทางกฎหมายเบื้องต้น (ฉบับร่าง)"), size_pt=MEMO_HEADING_SIZE_PT, bold=True)

    add_body(f"วันที่: {_thai_date_today()}")
    add_body(f"เรียน: {slots.get('client_name') or _MISSING_PLACEHOLDER}")

    add_heading("1. ข้อเท็จจริง")
    add_body(slots.get("facts_summary") or _MISSING_PLACEHOLDER)
    for fact in (slots.get("key_facts") or []):
        add_body(f"• {fact}")

    add_heading("2. ประเด็นที่ต้องพิจารณา")
    for item in rendered["issues"] or [{"text": _MISSING_PLACEHOLDER, "label": ""}]:
        add_body(f"{item['label']} {item['text']}".strip())

    add_heading("3. ข้อกฎหมายและการวิเคราะห์")
    for item in rendered["analysis"] or [{"text": _MISSING_PLACEHOLDER, "label": ""}]:
        add_body(f"{item['label']} {item['text']}".strip())

    add_heading("4. ความเห็นเบื้องต้น")
    conclusion = rendered["conclusion"]
    add_body(f"{conclusion['label']} {conclusion['text'] or _MISSING_PLACEHOLDER}".strip())

    add_heading("5. ขั้นตอนที่แนะนำ")
    for step in (rendered["next_steps"] or [_MISSING_PLACEHOLDER]):
        add_body(f"• {step}")

    add_heading("6. แหล่งอ้างอิง")
    if used_chunks:
        for chunk in used_chunks:
            tags_text = "/".join(chunk.get("tags") or []) or "-"
            first_line = (chunk["content"].splitlines() or [""])[0][:150]
            add_body(f"#{chunk['id']} ({tags_text}) — {first_line}")
    else:
        add_body("ไม่มีข้อมูลจากฐานความรู้ที่ใช้โดยตรง (ใช้ความรู้ทั่วไปทั้งหมด)")

    footer_p = doc.add_paragraph()
    _set_thai_run_font(footer_p.add_run("ร่างโดย ThaiMeSook เพื่อใช้ประกอบการปรึกษาทนายความ"), size_pt=10)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def dispatch_draft_legal_document(query: str, user_id: int, chat_id: int) -> dict:
    """จุดเดียวที่ main.py เรียกสำหรับ intent "draft_legal_document" — ไม่ผ่าน truth_guard/run_agentic_tool_loop
    เลย (ไม่ใช่การแก้ไฟล์ Excel) คืน {"document_id": None, "answer": ...} เสมอ"""
    history = get_chat_messages(chat_id) if chat_id else []
    slots = _extract_memo_slots(query, history)

    if not slots.get("facts_summary"):
        return {"document_id": None, "answer": "เล่าสถานการณ์ให้ฟังหน่อยครับ จะได้ร่างบันทึกความเห็นให้ถูกต้อง"}

    search_query = slots["facts_summary"] + (
        "\n" + "\n".join(slots.get("user_questions") or []) if slots.get("user_questions") else ""
    )
    results = legal_retrieval.legal_search(search_query, k=5, rerank_top_k=5)
    grounded_chunks = [r for r in results if r["score"] >= legal_retrieval.OUT_OF_KB_SCORE_THRESHOLD]
    grounded_ids = {c["id"] for c in grounded_chunks}
    grounded_by_id = {c["id"]: c for c in grounded_chunks}

    memo_content = _draft_memo_content(slots, grounded_chunks)
    rendered = {
        "issues": [_render_paragraph(item, grounded_ids) for item in (memo_content.get("issues") or [])],
        "analysis": [_render_paragraph(item, grounded_ids) for item in (memo_content.get("analysis") or [])],
        "conclusion": _render_paragraph(memo_content.get("conclusion") or {}, grounded_ids),
        "next_steps": [s for s in (memo_content.get("next_steps") or []) if s],
    }

    used_ids = set()
    for item in rendered["issues"] + rendered["analysis"] + [rendered["conclusion"]]:
        used_ids.update(item["chunk_ids"])
    used_chunks = [grounded_by_id[cid] for cid in sorted(used_ids) if cid in grounded_by_id]

    docx_bytes = build_memo_docx(slots, rendered, used_chunks)
    title = f"บันทึกความเห็นทางกฎหมายเบื้องต้น - {slots.get('client_name') or 'ไม่ระบุชื่อลูกค้า'}"
    memo_id = save_legal_memo(user_id, chat_id, title, slots, docx_bytes)

    grounded_count = sum(1 for item in rendered["issues"] + rendered["analysis"] + [rendered["conclusion"]] if item["chunk_ids"])
    general_count = len(rendered["issues"]) + len(rendered["analysis"]) + 1 - grounded_count
    missing = _missing_slot_labels(slots)

    download_url = f"/api/legal-memo/{memo_id}/download"
    answer_lines = [
        f"ร่างบันทึกความเห็นทางกฎหมายเรียบร้อยครับ [📥 ดาวน์โหลด]({download_url})",
        f"อ้างอิงฐานความรู้ {grounded_count} ย่อหน้า · ความรู้ทั่วไป {general_count} ย่อหน้า",
    ]
    if missing:
        answer_lines.append("ยังไม่ได้ระบุ: " + ", ".join(missing))

    return {"document_id": None, "answer": "\n".join(answer_lines)}
