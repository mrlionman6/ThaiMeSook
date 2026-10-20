"""รูปแบบไฟล์ Word ทางการของบันทึกความเห็นทางกฎหมาย (python-docx ล้วน) ไม่มีสัญลักษณ์/อีโมจิ
อ้างอิงฐานความรู้เป็นเลข [n] ท้ายย่อหน้า ย่อหน้าที่ไม่มีเลขคือความรู้ทั่วไป"""
import io
import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT_NAME = "TH Sarabun New"
BODY_PT = 16
TITLE_PT = 20
NOTE_PT = 12
PLACEHOLDER = "............"
_THAI_MONTHS = (
    "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
    "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม",
)


def thai_date(d: datetime.date | None = None) -> str:
    d = d or datetime.date.today()
    return f"{d.day} {_THAI_MONTHS[d.month - 1]} พ.ศ. {d.year + 543}"


def _set_thai_lang(rpr) -> None:
    """ตั้ง w:lang เป็น th-TH ทุกแกน (val/eastAsia/bidi) — Word ใช้ค่านี้ตัดสินกฎการตัดคำ/ขึ้นบรรทัดใหม่ของไทย
    ค่า default ของ python-docx (en-US/ar-SA) ไม่รู้จักกฎนี้ ทำให้ตัดคำกลางคำได้ (เหตุการณ์จริงที่แก้)"""
    old = rpr.find(qn("w:lang"))
    if old is not None:
        rpr.remove(old)
    lang = OxmlElement("w:lang")
    for attr in ("w:val", "w:eastAsia", "w:bidi"):
        lang.set(qn(attr), "th-TH")
    rpr.append(lang)


def _font(run, size=BODY_PT, bold=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT_NAME)
    for tag in ("w:szCs", "w:bCs"):
        el = rpr.find(qn(tag))
        if el is not None:
            rpr.remove(el)
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(size * 2))
    rpr.append(szcs)
    if bold:
        rpr.append(OxmlElement("w:bCs"))
    _set_thai_lang(rpr)


def _para(doc, text="", *, size=BODY_PT, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT,
          left_cm=0.0, first_cm=0.0, before=0, after=0):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.left_indent = Cm(left_cm)
    pf.first_line_indent = Cm(first_cm)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0
    if text:
        _font(p.add_run(text), size=size, bold=bold)
    return p


def _bottom_border(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "4"), ("w:color", "000000")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    ppr.append(bdr)


def _page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run("- "), size=NOTE_PT)
    run = p.add_run()
    _font(run, size=NOTE_PT)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    _font(p.add_run(" -"), size=NOTE_PT)


def _header_line(doc, label, value):
    p = _para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, left_cm=2.0, first_cm=-2.0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(2.0), WD_TAB_ALIGNMENT.LEFT)
    _font(p.add_run(label), bold=True)
    _font(p.add_run("\t" + value))
    return p


def _ref_suffix(chunk_ids, ref_numbers):
    nums = sorted({ref_numbers[c] for c in chunk_ids if c in ref_numbers})
    return (" [" + ", ".join(str(n) for n in nums) + "]") if nums else ""


def _numbered_items(doc, section_no, items, ref_numbers):
    """ข้อย่อย 1.1, 1.2 แบบ hanging indent — items เป็น dict {text, chunk_ids} หรือ str"""
    for i, item in enumerate(items, start=1):
        text = item if isinstance(item, str) else (item.get("text") or "")
        ids = [] if isinstance(item, str) else (item.get("chunk_ids") or [])
        hang = 1.0 if str(section_no).count(".") == 0 else 1.5
        p = _para(doc, left_cm=1.25 + hang, first_cm=-hang, after=2)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(1.25 + hang), WD_TAB_ALIGNMENT.LEFT)
        _font(p.add_run(f"{section_no}.{i}\t{text}{_ref_suffix(ids, ref_numbers)}"))


def _section_heading(doc, text):
    _para(doc, text, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=10, after=4)


def build_memo_docx(slots: dict, rendered: dict, used_chunks: list[dict]) -> bytes:
    """slots: ข้อเท็จจริงจากผู้ใช้ rendered: {issues, analysis, arguments{our_side, other_side}, conclusion, next_steps}
    used_chunks: chunk ที่ถูกอ้างจริง เรียงตามลำดับอ้างอิง (เลข [n] นับตามลำดับใน list นี้)"""
    ref_numbers = {c["id"]: n for n, c in enumerate(used_chunks, start=1)}
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(BODY_PT)
    _set_thai_lang(normal.element.get_or_add_rPr())
    _page_number_footer(sec)

    _para(doc, "บันทึกความเห็นทางกฎหมายเบื้องต้น", size=TITLE_PT, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, "(ฉบับร่าง)", align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

    subject = slots.get("legal_area") or PLACEHOLDER
    _header_line(doc, "เรื่อง", f"ความเห็นทางกฎหมายเบื้องต้นเกี่ยวกับ{subject}" if slots.get("legal_area") else PLACEHOLDER)
    _header_line(doc, "เรียน", slots.get("client_name") or PLACEHOLDER)
    last = _header_line(doc, "วันที่", thai_date())
    last.paragraph_format.space_after = Pt(6)
    _bottom_border(last)

    _section_heading(doc, "1. ข้อเท็จจริง")
    _para(doc, slots.get("facts_summary") or PLACEHOLDER, first_cm=1.25, after=2)
    if slots.get("key_facts"):
        _para(doc, "ข้อเท็จจริงที่สำคัญ มีดังนี้", first_cm=1.25, after=2)
        _numbered_items(doc, 1, slots["key_facts"], ref_numbers)

    _section_heading(doc, "2. ประเด็นที่ต้องพิจารณา")
    _numbered_items(doc, 2, rendered.get("issues") or [PLACEHOLDER], ref_numbers)

    _section_heading(doc, "3. ข้อกฎหมายและการวิเคราะห์")
    for item in rendered.get("analysis") or [{"text": PLACEHOLDER, "chunk_ids": []}]:
        _para(doc, (item.get("text") or PLACEHOLDER) + _ref_suffix(item.get("chunk_ids") or [], ref_numbers),
              first_cm=1.25, after=4)

    args = rendered.get("arguments") or {}
    _section_heading(doc, "4. แนวทางข้อต่อสู้")
    _para(doc, "4.1 ฝ่ายผู้ขอความเห็น", bold=True, left_cm=1.25, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
    _numbered_items(doc, "4.1", args.get("our_side") or [PLACEHOLDER], ref_numbers)
    _para(doc, "4.2 ฝ่ายคู่กรณี", bold=True, left_cm=1.25, align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=2)
    _numbered_items(doc, "4.2", args.get("other_side") or [PLACEHOLDER], ref_numbers)

    _section_heading(doc, "5. ความเห็นเบื้องต้น")
    concl = rendered.get("conclusion") or {}
    _para(doc, (concl.get("text") or PLACEHOLDER) + _ref_suffix(concl.get("chunk_ids") or [], ref_numbers),
          first_cm=1.25, after=4)

    _section_heading(doc, "6. ข้อเสนอแนะและขั้นตอนดำเนินการ")
    _numbered_items(doc, 6, rendered.get("next_steps") or [PLACEHOLDER], ref_numbers)

    _section_heading(doc, "เอกสารอ้างอิง")
    if used_chunks:
        for n, chunk in enumerate(used_chunks, start=1):
            tags = ", ".join(chunk.get("tags") or [])
            first_line = (chunk.get("content") or "").splitlines()[0][:160] if chunk.get("content") else ""
            tag_part = f" หมวด {tags}" if tags else ""
            p = _para(doc, left_cm=1.0, first_cm=-1.0, after=2)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0), WD_TAB_ALIGNMENT.LEFT)
            _font(p.add_run(f"[{n}]\tฐานข้อมูลกฎหมาย ThaiMeSook รายการที่ {chunk['id']}{tag_part}: {first_line}"))
    else:
        _para(doc, "ไม่มีการอ้างอิงจากฐานข้อมูลกฎหมายของระบบ", first_cm=1.25)
    _para(doc, "ข้อความที่ไม่มีเลขอ้างอิงกำกับ เรียบเรียงจากความรู้ทั่วไปทางกฎหมาย", size=14, first_cm=1.25, before=4)

    _para(doc, "ผู้จัดทำร่าง  ThaiMeSook", align=WD_ALIGN_PARAGRAPH.RIGHT, before=24)
    _para(doc, "เอกสารฉบับนี้จัดทำเพื่อใช้ประกอบการปรึกษาทนายความ", size=NOTE_PT,
          align=WD_ALIGN_PARAGRAPH.RIGHT)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
