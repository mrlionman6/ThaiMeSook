"""ทดสอบ document preview (ดูไฟล์ Excel ในแท็บด้านขวา แบบ read-only ล้วน) ใช้ fixture workbook สังเคราะห์
(ไม่ใช่ไฟล์ลูกค้าจริง) และ fake DB ในหน่วยความจำล้วนๆ ผ่าน tests/_harness.py"""
import io
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

import pytest
import openpyxl

from _harness import build_namespace, inject_formula_cache


@pytest.fixture
def ns():
    return build_namespace()


def _build_quote_workbook():
    """ฟิกซ์เจอร์จำลองใบเสนอราคา — ตัวเลขสัดส่วนเดียวกับไฟล์จริงที่เคยใช้ทดสอบ change_rate ในรอบก่อนหน้า
    (58150/4070.5/62220.5) ไม่ใช่ไฟล์ลูกค้าจริง มี VAT 7% เป็นสูตร + BAHTTEXT อ่านยอดรวมเป็นตัวหนังสือไทย ตั้ง
    number_format เป็น "#,##0.00" ให้คอลัมน์ตัวเลขทุกแถว เพื่อทดสอบการจัดรูปแบบตัวเลขของ preview ไปด้วยในตัว"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws["A26"] = "รวมเงิน"
    ws["G26"] = 58150
    ws["G26"].number_format = "#,##0.00"
    ws["A27"] = "ภาษีมูลค่าเพิ่ม 7%"
    ws["G27"] = "=G26*7/100"
    ws["G27"].number_format = "#,##0.00"
    ws["A28"] = "รวมราคาทั้งสิ้น"
    ws["G28"] = "=SUM(G26:G27)"
    ws["G28"].number_format = "#,##0.00"
    ws["A29"] = "จำนวนเงินเป็นตัวอักษร"
    ws["B29"] = "=BAHTTEXT(G28)"
    buf = io.BytesIO()
    wb.save(buf)
    return inject_formula_cache(buf.getvalue(), {"G27": 4070.5, "G28": 62220.5, "B29": "x"})


def _cell(window, row, col):
    row_entry = next((r for r in window["rows"] if r["row"] == row), None)
    if row_entry is None:
        return None
    return next((c for c in row_entry["cells"] if c["col"] == col), None)


def _cell_text(window, row, col):
    cell = _cell(window, row, col)
    return cell["text"] if cell else None


# ---------- BAHTTEXT ----------

@pytest.mark.parametrize("amount,expected", [
    (0, "ศูนย์บาทถ้วน"),
    (21, "ยี่สิบเอ็ดบาทถ้วน"),
    (100, "หนึ่งร้อยบาทถ้วน"),
    (1_000_000, "หนึ่งล้านบาทถ้วน"),
    (62220.5, "หกหมื่นสองพันสองร้อยยี่สิบบาทห้าสิบสตางค์"),
    (1234567.89, "หนึ่งล้านสองแสนสามหมื่นสี่พันห้าร้อยหกสิบเจ็ดบาทแปดสิบเก้าสตางค์"),
])
def test_bahttext(ns, amount, expected):
    assert ns["_bahttext"](amount) == expected


def test_bahttext_negative(ns):
    assert ns["_bahttext"](-100) == "ลบหนึ่งร้อยบาทถ้วน"


def _compile_to_excel_compiler(ns, raw: bytes):
    """ExcelCompiler ต้องการเปิดจากพาธไฟล์บนดิสก์จริง (ไม่ใช่ BytesIO ในหน่วยความจำ) — เขียนลง temp file ก่อน"""
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.write(raw)
    tmp.close()
    return ns["ExcelCompiler"](tmp.name)


def test_try_evaluate_bahttext_resolves_simple_cell_reference(ns):
    excel = _compile_to_excel_compiler(ns, _build_quote_workbook())
    result = ns["_try_evaluate_bahttext"]("=BAHTTEXT(G28)", "Quotation", excel)
    assert result == "หกหมื่นสองพันสองร้อยยี่สิบบาทห้าสิบสตางค์"


def test_try_evaluate_bahttext_returns_none_for_non_bahttext_formula(ns):
    excel = _compile_to_excel_compiler(ns, _build_quote_workbook())
    assert ns["_try_evaluate_bahttext"]("=SUM(G26:G27)", "Quotation", excel) is None


# ---------- number format display ----------

def test_format_cell_display_text_thousands_and_decimals(ns):
    assert ns["_format_cell_display_text"](58150, "#,##0.00") == "58,150.00"
    assert ns["_format_cell_display_text"](4070.5, "#,##0.00") == "4,070.50"


def test_format_cell_display_text_percent(ns):
    assert ns["_format_cell_display_text"](0.1, "0%") == "10%"


def test_format_cell_display_text_general_passthrough(ns):
    assert ns["_format_cell_display_text"](4070.5, "General") == "4070.5"
    assert ns["_format_cell_display_text"](None, "#,##0.00") == ""


# ---------- render หน้าต่างชีต (ค่าหลัก + number format + BAHTTEXT) ----------

def test_render_sheet_window_computes_formula_and_bahttext_values(ns):
    raw = _build_quote_workbook()
    window = ns["_render_sheet_window"](raw, "Quotation", 1, 100)

    assert _cell_text(window, 26, 7) == "58,150.00"  # G26 รวมเงิน (ค่าคงที่)
    assert _cell_text(window, 27, 7) == "4,070.50"   # G27 VAT 7% (สูตร)
    assert _cell_text(window, 28, 7) == "62,220.50"  # G28 รวมราคาทั้งสิ้น (สูตร)
    assert _cell_text(window, 29, 2) == "หกหมื่นสองพันสองร้อยยี่สิบบาทห้าสิบสตางค์"  # B29 BAHTTEXT


def test_render_sheet_window_marks_uncomputable_formula(ns):
    """สูตรที่ pycel คำนวณไม่ได้จริงๆ (ไม่ใช่ BAHTTEXT แต่เป็นฟังก์ชันที่ pycel ไม่รู้จักเหมือนกัน) ต้องแสดง "—"
    พร้อม not_computable=True และ title แจ้งผู้ใช้ — จงใจไม่ inject ค่า cache ให้เซลล์นี้เลย (ต่างจาก fixture อื่น
    ในไฟล์นี้) เพราะ pycel จะคืนค่า cache เดิมในไฟล์ตรงๆ โดยไม่พยายามคำนวณใหม่เลยถ้ามีอยู่แล้ว (เหมือนสถานการณ์
    จริงที่สูตรนี้เพิ่งถูกเขียนด้วย openpyxl ไม่เคยมีใครคำนวณมาก่อน จึงไม่มีค่า cache ให้เผื่อ)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "=NOTAREALFUNCTION(B1)"
    ws["B1"] = "a"
    buf = io.BytesIO()
    wb.save(buf)
    raw = buf.getvalue()

    window = ns["_render_sheet_window"](raw, "Sheet1", 1, 10)
    cell = _cell(window, 1, 1)
    assert cell["text"] == "—"
    assert cell["not_computable"] is True
    assert cell["title"] == ns["_PREVIEW_NOT_COMPUTABLE_TITLE"]


# ---------- merges / escaping ----------

def test_render_sheet_window_reports_merge_colspan_rowspan(ns):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "หัวตาราง"
    ws.merge_cells("A1:C2")
    buf = io.BytesIO()
    wb.save(buf)
    raw = buf.getvalue()

    window = ns["_render_sheet_window"](raw, "Sheet1", 1, 10)
    assert window["merges"] == [{"start_row": 1, "end_row": 2, "start_col": 1, "end_col": 3}]
    assert _cell_text(window, 1, 1) == "หัวตาราง"


def test_render_sheet_window_escapes_html_in_cell_text(ns):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "<script>alert(1)</script>"
    buf = io.BytesIO()
    wb.save(buf)
    raw = buf.getvalue()

    window = ns["_render_sheet_window"](raw, "Sheet1", 1, 10)
    text = _cell_text(window, 1, 1)
    assert "<script>" not in text
    assert "&lt;script&gt;" in text


# ---------- หลายชีต ----------

def test_render_sheet_window_lists_all_sheets_and_can_select_second(ns):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "หน้าแรก"
    ws1["A1"] = "one"
    ws2 = wb.create_sheet("หน้าสอง")
    ws2["A1"] = "two"
    buf = io.BytesIO()
    wb.save(buf)
    raw = buf.getvalue()

    window = ns["_render_sheet_window"](raw, "หน้าสอง", 1, 10)
    sheet_names = [s["name"] for s in window["sheets"]]
    assert sheet_names == ["หน้าแรก", "หน้าสอง"]
    assert window["sheet"] == "หน้าสอง"
    assert _cell_text(window, 1, 1) == "two"


# ---------- แบ่งหน้า (pagination) ----------

def test_render_sheet_window_paginates_large_sheet(ns):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in range(1, 351):
        ws.cell(row=row, column=1, value=row)
    buf = io.BytesIO()
    wb.save(buf)
    raw = buf.getvalue()

    page1 = ns["_render_sheet_window"](raw, "Sheet1", 1, 100)
    assert page1["used_rows"] == 350
    assert page1["end_row"] == 100
    assert page1["has_more"] is True
    assert len(page1["rows"]) == 100
    assert _cell_text(page1, 1, 1) == "1"
    assert _cell_text(page1, 100, 1) == "100"

    page2 = ns["_render_sheet_window"](raw, "Sheet1", 101, 100)
    assert page2["end_row"] == 200
    assert page2["has_more"] is True

    page3 = ns["_render_sheet_window"](raw, "Sheet1", 201, 100)
    assert page3["end_row"] == 300
    assert page3["has_more"] is True

    page4 = ns["_render_sheet_window"](raw, "Sheet1", 301, 100)
    assert page4["end_row"] == 350
    assert page4["has_more"] is False
    assert _cell_text(page4, 350, 1) == "350"


def test_render_sheet_window_caps_columns_at_50(ns):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for col in range(1, 61):
        ws.cell(row=1, column=col, value=col)
    buf = io.BytesIO()
    wb.save(buf)
    raw = buf.getvalue()

    window = ns["_render_sheet_window"](raw, "Sheet1", 1, 10)
    assert window["used_cols"] == ns["_PREVIEW_MAX_COLUMNS"]
    assert _cell_text(window, 1, 50) == "50"
    assert _cell(window, 1, 51) is None


# ---------- ไฮไลต์เซลล์ที่เปลี่ยน (_merge_preview_highlight) ----------

def test_merge_preview_highlight_marks_changed_cells(ns):
    before = ns["_render_sheet_window"](_build_quote_workbook(), "Quotation", 1, 100)

    after_raw_wb = openpyxl.load_workbook(io.BytesIO(_build_quote_workbook()))
    ws = after_raw_wb["Quotation"]
    ws["G27"] = "=G26*10/100"
    buf = io.BytesIO()
    after_raw_wb.save(buf)
    after_raw = inject_formula_cache(buf.getvalue(), {"G27": 5815.0, "G28": 63965.0})

    after = ns["_render_sheet_window"](after_raw, "Quotation", 1, 100)
    merged = ns["_merge_preview_highlight"](after, before)

    vat_cell = _cell(merged, 27, 7)
    assert vat_cell["text"] == "5,815.00"
    assert vat_cell["changed"] is True
    assert vat_cell["old_text"] == "4,070.50"

    total_cell = _cell(merged, 26, 7)
    assert "changed" not in total_cell


def test_merge_preview_highlight_none_baseline_marks_nothing(ns):
    window = ns["_render_sheet_window"](_build_quote_workbook(), "Quotation", 1, 100)
    merged = ns["_merge_preview_highlight"](window, None)
    for row_entry in merged["rows"]:
        for cell in row_entry["cells"]:
            assert "changed" not in cell


# ---------- resolve mode (current/pending) ----------

def test_resolve_preview_mode_bytes_current_no_version_yet(ns):
    doc = {"id": 1, "original_bytes": b"ORIGINAL"}
    displayed, baseline = ns["_resolve_preview_mode_bytes"](doc, "current", b"ORIGINAL", {}, [])
    assert displayed == b"ORIGINAL"
    assert baseline is None


def test_resolve_preview_mode_bytes_current_compares_previous_version(ns):
    fake_db = ns["_fake_db"]
    fake_db.add_document(1, "quote.xlsx", b"ORIGINAL", {}, chat_id=1, user_id=1)
    fake_db.create_document_version(1, b"V1", {})
    fake_db.create_document_version(1, b"V2", {})
    doc = fake_db.get_editable_document_raw(1)

    displayed, baseline = ns["_resolve_preview_mode_bytes"](doc, "current", b"V2", {}, [])
    assert displayed == b"V2"
    assert baseline == b"V1"


def test_resolve_preview_mode_bytes_current_single_version_compares_original(ns):
    fake_db = ns["_fake_db"]
    fake_db.add_document(1, "quote.xlsx", b"ORIGINAL", {}, chat_id=1, user_id=1)
    fake_db.create_document_version(1, b"V1", {})
    doc = fake_db.get_editable_document_raw(1)

    displayed, baseline = ns["_resolve_preview_mode_bytes"](doc, "current", b"V1", {}, [])
    assert displayed == b"V1"
    assert baseline == b"ORIGINAL"


def test_resolve_preview_mode_bytes_pending_empty_queue_like_current(ns):
    doc = {"id": 1, "original_bytes": b"ORIGINAL"}
    displayed, baseline = ns["_resolve_preview_mode_bytes"](doc, "pending", b"BASE", {}, [])
    assert displayed == b"BASE"
    assert baseline is None


def test_resolve_preview_mode_bytes_pending_with_queue_replays_virtually(ns):
    raw = _build_quote_workbook()
    label_map = ns["_extract_excel_labels"](raw)

    pending_queue = [{
        "id": 1,
        "action": {
            "type": "change_rate", "document_id": 1, "label": "ภาษีมูลค่าเพิ่ม 7%", "new_rate": 10,
            "sheet": label_map["ภาษีมูลค่าเพิ่ม 7%"]["sheet"],
            "row": label_map["ภาษีมูลค่าเพิ่ม 7%"]["row"],
            "col": label_map["ภาษีมูลค่าเพิ่ม 7%"]["col"],
        },
    }]
    doc = {"id": 1, "original_bytes": raw}
    displayed, baseline = ns["_resolve_preview_mode_bytes"](doc, "pending", raw, label_map, pending_queue)
    assert baseline == raw
    assert displayed != raw

    window = ns["_render_sheet_window"](displayed, "Quotation", 1, 100)
    assert _cell_text(window, 27, 7) == "5,815.00"


def test_preview_after_propose_change_rate_pending_vs_current(ns):
    """สถานการณ์เต็ม (เทียบเท่าสิ่งที่ endpoint preview ทำจริงทั้งสอง mode): propose เปลี่ยน VAT เป็น 10% (ยังไม่
    ยืนยัน) -> mode=pending ต้องเห็นค่าใหม่ (5,815.00) พร้อมไฮไลต์ว่าเปลี่ยน / mode=current ต้องยังเป็นค่าเดิม
    (4,070.50) เป๊ะ ไม่ถูกกระทบเลยจนกว่าจะยืนยันจริง"""
    fake_db = ns["_fake_db"]
    raw = _build_quote_workbook()
    label_map = ns["_extract_excel_labels"](raw)
    fake_db.add_document(1, "quote.xlsx", raw, label_map, chat_id=1, user_id=1)

    propose_result = ns["propose_change_rate"](1, 1, 1, "ภาษีมูลค่าเพิ่ม 7%", 10)
    assert propose_result["proposed"] is True

    doc = fake_db.get_editable_document_raw(1)
    base_bytes, base_label_map = ns["get_working_copy"](doc)
    pending_queue = ns["get_chat_pending_actions"](1)

    # mode=current: ไม่แตะคิวเลย ค่ายังเป็นของเดิม ไม่มีอะไรไฮไลต์ (ไม่มีเวอร์ชันก่อนหน้าให้เทียบด้วย)
    current_displayed, current_baseline = ns["_resolve_preview_mode_bytes"](
        doc, "current", base_bytes, base_label_map, pending_queue
    )
    current_window = ns["_render_sheet_window"](current_displayed, "Quotation", 1, 100)
    assert _cell_text(current_window, 27, 7) == "4,070.50"
    assert current_baseline is None

    # mode=pending: เห็นค่าที่ "จะ" เป็นถ้ายืนยันคิวนี้ พร้อมไฮไลต์ว่าเปลี่ยนจากอะไร
    pending_displayed, pending_baseline = ns["_resolve_preview_mode_bytes"](
        doc, "pending", base_bytes, base_label_map, pending_queue
    )
    pending_window = ns["_render_sheet_window"](pending_displayed, "Quotation", 1, 100)
    baseline_window = ns["_render_sheet_window"](pending_baseline, "Quotation", 1, 100)
    merged = ns["_merge_preview_highlight"](pending_window, baseline_window)

    vat_cell = _cell(merged, 27, 7)
    assert vat_cell["text"] == "5,815.00"
    assert vat_cell["changed"] is True
    assert vat_cell["old_text"] == "4,070.50"


# ---------- ownership (เหมือน endpoint download เป๊ะ — เทียบ user_id) ----------

def test_get_editable_document_returns_none_for_other_users_document(ns):
    fake_db = ns["_fake_db"]
    fake_db.add_document(1, "quote.xlsx", b"ORIGINAL", {}, chat_id=1, user_id=1)
    assert ns["get_editable_document"](1, 1) is not None
    assert ns["get_editable_document"](1, 999) is None
