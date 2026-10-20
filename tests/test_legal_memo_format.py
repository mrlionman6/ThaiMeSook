"""ทดสอบรูปแบบ memo .docx หลังแก้ตัวอักษรยืดห่าง/ตัดคำกลางคำใน Word (align LEFT แทน thaiDistribute, w:lang
th-TH ทุก run + style Normal — ดู round2-p0 ข้อ 6) legal_memo_format.py ไม่พึ่ง main.py/DB เลย import ตรงได้"""
import sys
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

import legal_memo_format as memo_fmt

_SLOTS = {"client_name": "ลูกค้า", "legal_area": "ทดสอบ", "facts_summary": "ข้อเท็จจริงทดสอบ", "key_facts": []}
_RENDERED = {
    "issues": [{"text": "ประเด็นที่ 1", "chunk_ids": []}],
    "analysis": [{"text": "วิเคราะห์", "chunk_ids": []}],
    "arguments": {"our_side": [], "other_side": []},
    "conclusion": {"text": "สรุป", "chunk_ids": []},
    "next_steps": ["ขั้นตอนถัดไป"],
}


def _build_sample_doc() -> Document:
    docx_bytes = memo_fmt.build_memo_docx(_SLOTS, _RENDERED, [])
    return Document(BytesIO(docx_bytes))


def test_body_paragraph_is_not_thai_distribute():
    doc = _build_sample_doc()
    body_paragraph = next(p for p in doc.paragraphs if p.text == "ข้อเท็จจริงทดสอบ")
    assert body_paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert body_paragraph.alignment != WD_ALIGN_PARAGRAPH.THAI_JUSTIFY


def test_every_run_has_th_th_lang():
    doc = _build_sample_doc()
    checked = 0
    for p in doc.paragraphs:
        for run in p.runs:
            rpr = run._element.find(qn("w:rPr"))
            assert rpr is not None
            lang = rpr.find(qn("w:lang"))
            assert lang is not None
            assert lang.get(qn("w:val")) == "th-TH"
            assert lang.get(qn("w:bidi")) == "th-TH"
            checked += 1
    assert checked > 0


def test_normal_style_has_th_th_lang():
    doc = _build_sample_doc()
    normal_rpr = doc.styles["Normal"].element.find(qn("w:rPr"))
    assert normal_rpr is not None
    lang = normal_rpr.find(qn("w:lang"))
    assert lang is not None
    assert lang.get(qn("w:val")) == "th-TH"


def test_reference_heading_uses_formal_wording():
    docx_bytes = memo_fmt.build_memo_docx(_SLOTS, _RENDERED, [{"id": 1, "tags": [], "content": "เนื้อหา"}])
    doc = Document(BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "ฐานข้อมูลกฎหมาย ThaiMeSook" in full_text
    assert "ฐานความรู้" not in full_text


def test_no_reference_wording_when_no_chunks_used():
    doc = _build_sample_doc()
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "ไม่มีการอ้างอิงจากฐานข้อมูลกฎหมายของระบบ" in full_text
