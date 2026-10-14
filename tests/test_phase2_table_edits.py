"""ทดสอบเฟส 2 PR2 (แก้โครงสร้าง Excel: เปลี่ยนอัตรา / แก้-ลบ-เพิ่มรายการในตาราง — แก้ทันที ไม่มีขั้นยืนยัน/คิว
อีกต่อไป ดู feat/apply-immediately) ใช้ fixture workbook สังเคราะห์ (ไม่ใช่ไฟล์ลูกค้าจริง) และ fake DB/LLM client
ในหน่วยความจำล้วนๆ (ไม่ต่อ DB จริงเลย) ผ่าน tests/_harness.py"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

import pytest
import openpyxl

from _harness import (
    build_namespace, build_fixture_workbook, inject_formula_cache,
    FIXTURE_TOTAL, FIXTURE_VAT, FIXTURE_GRAND_TOTAL,
    _FakeCitation, _FakeAgenticTextBlock, _FakeAgenticResponse,
    _FakeAgenticOtherBlock, _FakeAgenticToolUseBlock,
)


@pytest.fixture
def ns():
    return build_namespace()


@pytest.fixture
def doc1(ns):
    """สร้างเอกสารทดสอบ (document_id=1, chat_id=1, user_id=1) จาก fixture workbook ไว้ใน fake db"""
    raw = build_fixture_workbook()
    label_map = ns["_extract_excel_labels"](raw)
    ns["_fake_db"].add_document(1, "quote.xlsx", raw, label_map, chat_id=1, user_id=1)
    return ns["_fake_db"].documents[1]


# ---------- ฟิกซ์เจอร์ที่สอง: รายละเอียดสินค้าแบบเหตุการณ์จริงที่ใช้ทดสอบ direct parse (_parse_direct_structural_
# commands) ซึ่งต้องจับคู่ "รายละเอียด" ของรายการในตารางด้วยข้อความจริง ("Box 4x4"/"ไส้ไก่") ไม่ใช่แค่ item_no —
# build_fixture_workbook() ของ _harness.py ใช้ชื่อ "Item A".."Item H" เฉยๆ ไม่พอทดสอบกรณีนี้ ----------
_BOX_ITEMS = [
    ("Box 4x4", 5, "ชิ้น", 50),
    ("Box 6x6", 3, "ชิ้น", 70),
    ("ไส้ไก่", 10, "ชิ้น", 20),
    ("Item D", 2, "ชุด", 300),
    ("Item E", 1, "เครื่อง", 5000),
    ("Item F", 4, "ชิ้น", 25),
    ("Item G", 8, "ชิ้น", 10),
    ("Item H", 6, "ชิ้น", 15),
]
_BOX_TOTAL = sum(qty * price for _, qty, _, price in _BOX_ITEMS)
_BOX_VAT_7 = round(_BOX_TOTAL * 0.07, 2)
_BOX_GRAND_7 = _BOX_TOTAL + _BOX_VAT_7


def _build_box_quote_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws["A15"] = "ลำดับ"
    ws["B15"] = "รายการ"
    ws["D15"] = "จำนวน"
    ws["E15"] = "หน่วย"
    ws["F15"] = "ราคา/หน่วย"
    ws["G15"] = "จำนวนเงิน"

    cache_values = {}
    for i, (desc, qty, unit, price) in enumerate(_BOX_ITEMS):
        row = 16 + i
        ws.cell(row=row, column=1, value=i + 1)
        ws.cell(row=row, column=2, value=desc)
        ws.cell(row=row, column=4, value=qty)
        ws.cell(row=row, column=5, value=unit)
        ws.cell(row=row, column=6, value=price)
        ws.cell(row=row, column=7, value=f"=D{row}*F{row}")
        cache_values[f"G{row}"] = qty * price

    last_item_row = 16 + len(_BOX_ITEMS) - 1
    ws["A26"] = "รวมเงิน"
    ws["G26"] = f"=SUM(G16:G{last_item_row})"
    cache_values["G26"] = _BOX_TOTAL

    ws["A27"] = "ภาษีมูลค่าเพิ่ม 7%"
    ws["G27"] = "=SUM(G26)*7/100"
    cache_values["G27"] = _BOX_VAT_7

    ws["A28"] = "รวมราคาทั้งสิ้น"
    ws["G28"] = "=SUM(G26:G27)"
    cache_values["G28"] = _BOX_GRAND_7

    buf = io.BytesIO()
    wb.save(buf)
    return inject_formula_cache(buf.getvalue(), cache_values)


@pytest.fixture
def box_doc(ns):
    raw = _build_box_quote_workbook()
    label_map = ns["_extract_excel_labels"](raw)
    ns["_fake_db"].add_document(1, "quote.xlsx", raw, label_map, chat_id=1, user_id=1)
    return ns["_fake_db"].documents[1]


def _build_quote_workbook_with_items(items):
    """ทั่วไปกว่า _build_box_quote_workbook — รับรายการ (desc, qty, unit, price) เองได้ ใช้ทดสอบว่า advisor
    ส่งตารางรายการเป็น context ได้กับสินค้า "ทุกประเภท" ไม่ใช่เฉพาะกรณี Box/ไส้ไก่ (ดู feat/advisor-route:
    เกณฑ์เปรียบเทียบ/ผลกระทบ/หน่วยขายต้องเป็นกติกาทั่วไป ไม่ hardcode ตามสินค้า)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws["A15"] = "ลำดับ"
    ws["B15"] = "รายการ"
    ws["D15"] = "จำนวน"
    ws["E15"] = "หน่วย"
    ws["F15"] = "ราคา/หน่วย"
    ws["G15"] = "จำนวนเงิน"

    cache_values = {}
    for i, (desc, qty, unit, price) in enumerate(items):
        row = 16 + i
        ws.cell(row=row, column=1, value=i + 1)
        ws.cell(row=row, column=2, value=desc)
        ws.cell(row=row, column=4, value=qty)
        ws.cell(row=row, column=5, value=unit)
        ws.cell(row=row, column=6, value=price)
        ws.cell(row=row, column=7, value=f"=D{row}*F{row}")
        cache_values[f"G{row}"] = qty * price

    last_item_row = 16 + len(items) - 1
    total = sum(qty * price for _, qty, _, price in items)
    ws["A26"] = "รวมเงิน"
    ws["G26"] = f"=SUM(G16:G{last_item_row})"
    cache_values["G26"] = total

    buf = io.BytesIO()
    wb.save(buf)
    return inject_formula_cache(buf.getvalue(), cache_values)


def _current(ns, document_id=1):
    """อ่านค่าปัจจุบันของเอกสาร (สำเนาทำงานล่าสุด) — ใช้ตรวจผลหลัง apply ทันที"""
    doc = ns["get_editable_document_raw"](document_id)
    return ns["read_document_values"](doc)


# ---------- change_rate (แก้ทันที) ----------

def test_apply_change_rate_to_10_percent(ns, doc1):
    result = ns["apply_change_rate"](1, 1, "ภาษีมูลค่าเพิ่ม 7%", 10)
    assert result["matched"] is True
    assert "✓ เปลี่ยน 'ภาษีมูลค่าเพิ่ม'" in result["message"]
    assert "จาก 7% เป็น 10%" in result["message"]
    assert "พิมพ์ 'ย้อนการแก้'" in result["message"]

    expected_vat = round(FIXTURE_TOTAL * 0.10, 2)
    expected_total = FIXTURE_TOTAL + expected_vat
    # ตัวเลขเงินมีคอมมาคั่นหลักพันเสมอ (ดู _format_money_value) เช่น "8,822" ไม่ใช่ "8822.0"
    assert ns["_format_money_value"](expected_vat) in result["message"]
    assert ns["_format_money_value"](expected_total) in result["message"]

    data = _current(ns)
    label_values = {l["label"]: l["value"] for l in data["labels"]}
    # label เปลี่ยนชื่อตามข้อความจริงในเซลล์ป้ายหลังแก้ (ดู _current_label_display_name) — ไม่ใช่ key เดิม
    # "ภาษีมูลค่าเพิ่ม 7%" ที่แช่แข็งไว้ตั้งแต่ extract ครั้งแรกอีกต่อไป
    assert "ภาษีมูลค่าเพิ่ม 7%" not in label_values
    assert label_values["ภาษีมูลค่าเพิ่ม 10%"] == expected_vat
    assert label_values["รวมราคาทั้งสิ้น"] == expected_total
    assert len(ns["_fake_db"].versions.get(1, [])) == 1  # แก้ทันทีเป็น 1 เวอร์ชัน ไม่มีคิวมาเกี่ยวข้องเลย


def test_apply_change_rate_message_and_values_show_current_label_not_stale_name(ns, doc1):
    """เหตุการณ์จริงที่แก้ (3 รอบ): (1) read_document_values/ยอดล่าสุด เคยโชว์ชื่อเดิม "ภาษีมูลค่าเพิ่ม 7%" ค้างอยู่
    ตลอดไปหลังเปลี่ยนเป็น 25% แล้ว (2) ข้อความผลลัพธ์เคยเขียน "เปลี่ยนอัตราของ 'ภาษีมูลค่าเพิ่ม 25%' เป็น 25%"
    (บอกว่าเปลี่ยนจาก 25% ทั้งที่จริงเปลี่ยนจาก 7%) ตอนนี้ต้องเขียน "เปลี่ยน 'ภาษีมูลค่าเพิ่ม' จาก 7% เป็น 25%"
    ชัดเจน (3) label ที่ถูกแก้ตรงๆ (VAT) ไม่ต้องโผล่ซ้ำใต้หัวข้อ "ยอดล่าสุด:" อีก เพราะมีบรรทัด "เปลี่ยนจาก...เป็น..."
    ของมันเองอยู่แล้วในข้อความหลัก (ดู _document_totals_lines exclude_labels) — read_document_values เอง (ไม่ผ่าน
    exclude_labels) ยังโชว์ "ภาษีมูลค่าเพิ่ม 25%" (ชื่อปัจจุบันจริง) ตามปกติ"""
    result = ns["apply_change_rate"](1, 1, "ภาษีมูลค่าเพิ่ม 7%", 25)
    assert result["matched"] is True
    assert "✓ เปลี่ยน 'ภาษีมูลค่าเพิ่ม' จาก 7% เป็น 25%" in result["message"]
    assert "ภาษีมูลค่าเพิ่ม 7%" not in result["message"]
    assert "ยอดล่าสุด" in result["message"]
    # VAT เองไม่ซ้ำใต้ "ยอดล่าสุด:" อีก (ถูกกันด้วย exclude_labels เพราะมีบรรทัดหลักของมันเองอยู่แล้ว)
    assert "ภาษีมูลค่าเพิ่ม 25%:" not in result["message"]
    # แต่ "รวมเงิน" (ไม่ได้ถูกกล่าวถึงที่ไหนเลย ไม่ใช่ผลข้างเคียงด้วย) ยังโผล่ใต้ "ยอดล่าสุด:" ตามปกติ
    assert "รวมเงิน:" in result["message"]

    data = _current(ns)
    label_values = {l["label"]: l["value"] for l in data["labels"]}
    assert "ภาษีมูลค่าเพิ่ม 7%" not in label_values
    expected_vat = round(FIXTURE_TOTAL * 0.25, 2)
    assert label_values["ภาษีมูลค่าเพิ่ม 25%"] == expected_vat


@pytest.mark.parametrize("value, expected", [
    (2.0, "2"), (2, "2"), (54436, "54,436"), (54436.0, "54,436"),
    (3810.5, "3,810.50"), (58246.52, "58,246.52"), (0, "0"), (None, "None"), ("ชิ้น", "ชิ้น"),
])
def test_format_money_value(ns, value, expected):
    assert ns["_format_money_value"](value) == expected


def test_format_side_effect_lines_has_space_before_verb(ns):
    """เหตุการณ์จริงที่แก้: เดิมติดกัน "ภาษีมูลค่าเพิ่ม 7%เปลี่ยนจาก..." (ไม่มีช่องว่างระหว่างชื่อ label กับ verb)
    อ่านยาก ตอนนี้ต้องมีช่องว่างคั่นเสมอ และตัวเลขต้องมีคอมมาคั่นหลักพัน"""
    side_effects = [{"label": "รวมราคาทั้งสิ้น", "old_value": 8581.4, "new_value": 10025}]
    text = ns["_format_side_effect_lines"](side_effects, "เปลี่ยน")
    assert "รวมราคาทั้งสิ้น เปลี่ยนจาก 8,581.40 เป็น 10,025" in text
    assert "ทั้งสิ้นเปลี่ยนจาก" not in text  # ไม่ติดกันแบบเดิม


def test_apply_change_rate_noop_when_already_at_target_rate(ns, doc1):
    """fixture เริ่มต้นที่ VAT 7% อยู่แล้ว — เสนอเปลี่ยนเป็น 7% (เท่าเดิม) ต้องไม่สร้างเวอร์ชันใหม่เลย"""
    result = ns["apply_change_rate"](1, 1, "ภาษีมูลค่าเพิ่ม 7%", 7)
    assert result["matched"] is False
    assert "อยู่แล้ว" in result["message"]
    assert ns["_fake_db"].versions.get(1, []) == []


def test_apply_change_rate_unsupported_formula_is_rejected(ns, doc1):
    # "รวมเงิน" เป็นสูตร =SUM(...) ไม่ตรงรูปแบบอัตราที่รองรับเลย
    result = ns["apply_change_rate"](1, 1, "รวมเงิน", 10)
    assert result["matched"] is False
    assert "ไม่ใช่รูปแบบอัตราที่รองรับ" in result["message"]
    assert ns["_fake_db"].versions.get(1, []) == []


def test_apply_change_rate_position_based_lookup_matches_real_file_scale(ns):
    """ทวนด้วยตัวเลขจากไฟล์จริง (ป่าไม้เขียว, PR #25): รวมเงิน 58150, VAT 7%->10% ต้องได้ 5815 และรวมราคา
    ทั้งสิ้น 58150+5815=63965 — ยืนยันว่าหา "ค่าใหม่" ถูกต้องด้วยตำแหน่งเซลล์ แม้ label มี % อยู่ในชื่อ"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws["A26"] = "รวมเงิน"
    ws["G26"] = 58150
    ws["A27"] = "ภาษีมูลค่าเพิ่ม 7%"
    ws["G27"] = "=G26*7/100"
    ws["A28"] = "รวมราคาทั้งสิ้น"
    ws["G28"] = "=SUM(G26:G27)"
    buf = io.BytesIO()
    wb.save(buf)
    raw = inject_formula_cache(buf.getvalue(), {"G27": 4070.5, "G28": 62220.5})
    label_map = ns["_extract_excel_labels"](raw)
    ns["_fake_db"].add_document(2, "real_scale.xlsx", raw, label_map, chat_id=1, user_id=1)

    result = ns["apply_change_rate"](2, 1, "ภาษีมูลค่าเพิ่ม 7%", 10)
    assert result["matched"] is True
    assert "5,815" in result["message"]  # คอมมาคั่นหลักพัน (ดู _format_money_value)
    assert "63,965" in result["message"]


# ---------- update_table_item (แก้ทันที) ----------

def test_apply_update_table_item_description(ns, doc1):
    result = ns["apply_update_table_item"](1, 1, 3, {"description": "Item C renamed"})
    assert result["matched"] is True
    assert "Item C renamed" in result["message"]
    assert _current(ns)["tables"][0]["rows"][2]["values"]["description"] == "Item C renamed"


def test_apply_update_table_item_noop_when_value_unchanged(ns, doc1):
    result = ns["apply_update_table_item"](1, 1, 1, {"description": "Item A"})
    assert result["matched"] is False
    assert "อยู่แล้ว" in result["message"]
    assert ns["_fake_db"].versions.get(1, []) == []


def test_apply_update_table_item_uses_thai_field_names(ns, doc1):
    result = ns["apply_update_table_item"](1, 1, 1, {"description": "Box 7x7", "qty": 15})
    assert result["matched"] is True
    assert "รายละเอียด เปลี่ยนเป็น Box 7x7" in result["message"]
    assert "จำนวน เปลี่ยนเป็น 15" in result["message"]
    assert "description" not in result["message"]
    assert "qty" not in result["message"]


# ---------- delete_table_item (แก้ทันที) ----------

def test_apply_delete_middle_item_shifts_and_renumbers_continuously(ns, doc1):
    """ลบรายการที่ 3 จาก 8 รายการ: ลำดับต้องเป็น 1..7 ต่อเนื่อง และรายละเอียดต้องเลื่อนขึ้นถูกแถว (รายการที่ 3
    ใหม่ต้องเป็นของเดิมรายการที่ 4) — รายการที่ 2 ต้องยังอยู่ ไม่ถูกลบผิดแถว"""
    result = ns["apply_delete_table_item"](1, 1, 3)
    assert result["matched"] is True

    filled = _current(ns)["tables"][0]["rows"]
    assert len(filled) == 7
    assert [r["values"]["no"] for r in filled] == [1, 2, 3, 4, 5, 6, 7]
    assert filled[1]["values"]["description"] == "Item B"
    assert filled[2]["values"]["description"] == "Item D"
    assert filled[2]["values"]["qty"] == 2
    assert filled[2]["values"]["unit_price"] == 300


def test_undo_after_apply_delete_restores_original(ns, doc1):
    ns["apply_delete_table_item"](1, 1, 3)

    undo_result = ns["undo_last_edit"](1, 1)
    assert undo_result["matched"] is True

    filled = _current(ns)["tables"][0]["rows"]
    assert len(filled) == 8
    assert filled[2]["values"]["description"] == "Item C"


def test_apply_delete_rejected_when_amount_not_formula(ns):
    raw_wb = openpyxl.Workbook()
    ws = raw_wb.active
    ws.title = "Quotation"
    ws["A15"] = "ลำดับ"
    ws["B15"] = "รายการ"
    ws["D15"] = "จำนวน"
    ws["E15"] = "หน่วย"
    ws["F15"] = "ราคา/หน่วย"
    ws["G15"] = "จำนวนเงิน"
    ws["A16"] = 1
    ws["B16"] = "Item X"
    ws["D16"] = 1
    ws["E16"] = "ชิ้น"
    ws["F16"] = 100
    ws["G16"] = 100  # ค่าคงที่ ไม่ใช่สูตร — ต้องถูกปฏิเสธ
    buf = io.BytesIO()
    raw_wb.save(buf)
    raw = buf.getvalue()
    label_map = ns["_extract_excel_labels"](raw)
    ns["_fake_db"].add_document(3, "bad_table.xlsx", raw, label_map, chat_id=1, user_id=1)

    result = ns["apply_delete_table_item"](3, 1, 1)
    assert result["matched"] is False
    assert result["message"] == ns["_TABLE_NON_FORMULA_AMOUNT_ERROR"]


# ---------- add_table_item (แก้ทันที) ----------

def test_apply_add_table_item_until_full(ns, doc1):
    result = ns["apply_add_table_item"](1, 1, {"description": "Item I", "qty": 2, "unit": "ชิ้น", "unit_price": 20})
    assert result["matched"] is True
    assert "รายละเอียด=Item I" in result["message"]

    result2 = ns["apply_add_table_item"](1, 1, {"description": "Item J", "qty": 1, "unit": "ชิ้น", "unit_price": 5})
    assert result2["matched"] is True

    # ตอนนี้เต็มแล้ว (10/10) — ลองเพิ่มอีกต้องถูกปฏิเสธทันที ไม่สร้างเวอร์ชัน
    versions_before = len(ns["_fake_db"].versions.get(1, []))
    rejected = ns["apply_add_table_item"](1, 1, {"description": "Item K", "qty": 1, "unit": "ชิ้น", "unit_price": 1})
    assert rejected["matched"] is False
    assert "เต็มแล้ว" in rejected["message"]
    assert "10" in rejected["message"]
    assert len(ns["_fake_db"].versions.get(1, [])) == versions_before


def test_apply_add_rejected_when_amount_not_formula(ns):
    raw_wb_obj = openpyxl.Workbook()
    ws = raw_wb_obj.active
    ws.title = "Quotation"
    ws["A15"] = "ลำดับ"
    ws["B15"] = "รายการ"
    ws["D15"] = "จำนวน"
    ws["E15"] = "หน่วย"
    ws["F15"] = "ราคา/หน่วย"
    ws["G15"] = "จำนวนเงิน"
    ws["A16"] = 1
    ws["B16"] = "Item X"
    ws["D16"] = 1
    ws["E16"] = "ชิ้น"
    ws["F16"] = 100
    ws["G16"] = 100  # ค่าคงที่ ไม่ใช่สูตร
    buf = io.BytesIO()
    raw_wb_obj.save(buf)
    raw = buf.getvalue()
    label_map = ns["_extract_excel_labels"](raw)
    ns["_fake_db"].add_document(4, "bad_table2.xlsx", raw, label_map, chat_id=1, user_id=1)

    result = ns["apply_add_table_item"](4, 1, {"description": "New", "qty": 1, "unit": "ชิ้น", "unit_price": 1})
    assert result["matched"] is False
    assert result["message"] == ns["_TABLE_NON_FORMULA_AMOUNT_ERROR"]


# ---------- คำสั่งยืนยัน/ยกเลิกแบบเดิม — ตอนนี้แค่ "ตรวจจับ" ว่าควรตอบคงที่ ไม่ใช่จัดการคิวอีกต่อไป ----------

@pytest.mark.parametrize("text", ["ยืนยันทั้งหมด", "ยืนยันหมด", "ยืนยันทุกรายการ"])
def test_classify_confirm_all(ns, text):
    assert ns["_classify_pending_queue_command"](text) == "confirm_all"


def test_classify_confirm_latest(ns):
    assert ns["_classify_pending_queue_command"]("ยืนยันล่าสุด") == "confirm_latest"


@pytest.mark.parametrize("text", ["ยืนยัน", "ยืนยันครับ", "ตกลง ยืนยัน"])
def test_classify_confirm_single(ns, text):
    assert ns["_classify_pending_queue_command"](text) == "confirm_single"


@pytest.mark.parametrize("text", ["ไม่ยืนยัน", "ยังไม่ยืนยัน"])
def test_classify_confirm_rejects_negated_forms(ns, text):
    assert ns["_classify_pending_queue_command"](text) is None


def test_classify_cancel_latest(ns):
    assert ns["_classify_pending_queue_command"]("ยกเลิกล่าสุด") == "cancel_latest"


@pytest.mark.parametrize("text", ["ยกเลิก", "ยกเลิกทั้งหมด"])
def test_classify_cancel_all(ns, text):
    assert ns["_classify_pending_queue_command"](text) == "cancel_all"


def test_classify_cancel_rejects_negated_form(ns):
    assert ns["_classify_pending_queue_command"]("ไม่ยกเลิก") is None


@pytest.mark.parametrize("text,expected", [
    ("ยืนยันทั้งหมดครับ", "confirm_all"),
    ("ยืนยันทั้งหมดเลย", "confirm_all"),
    ("ยืนยันหมดเลยนะ", "confirm_all"),
    ("ยกเลิกทั้งหมดเลย", "cancel_all"),
    ("ยืนยันด้วย", "confirm_single"),
    ("ยืนยันล่าสุดครับ", "confirm_latest"),
])
def test_classify_tolerates_trailing_particles(ns, text, expected):
    assert ns["_classify_pending_queue_command"](text) == expected


@pytest.mark.parametrize("text", ["ไม่ยืนยันนะ", "อย่าเพิ่งยืนยัน", "ยังไม่ยืนยันครับ"])
def test_classify_still_rejects_negated_forms_with_particles(ns, text):
    assert ns["_classify_pending_queue_command"](text) is None


def test_mentions_pending_queue_intent_catches_unrecognized_confirm_phrasing(ns):
    assert ns["_mentions_pending_queue_intent"]("ยืนยันแบบนั้นแหละ") is True


@pytest.mark.parametrize("text", ["ไม่ยืนยันนะ", "อย่าเพิ่งยืนยัน", "ไม่ยกเลิกนะ"])
def test_mentions_pending_queue_intent_excludes_negated_forms(ns, text):
    assert ns["_mentions_pending_queue_intent"](text) is False


@pytest.mark.parametrize("text", [
    "ต้องยืนยันตัวตนก่อนยื่นภาษีไหม",
    "ยืนยันตัวตนต้องทำอย่างไรบ้าง",
])
def test_mentions_pending_queue_intent_excludes_real_questions(ns, text):
    """กันคำถามกฎหมายทั่วไปที่บังเอิญมีคำว่า "ยืนยัน" ปนอยู่ (เช่น "ต้องยืนยันตัวตนก่อนยื่นภาษีไหม") ถูกเข้าใจผิด
    เป็นความตั้งใจจะยืนยัน/ยกเลิก — ใช้คำเชิงคำถามเดียวกับ _OPEN_COMMAND_QUESTION_WORDS กัน"""
    assert ns["_mentions_pending_queue_intent"](text) is False


def test_mentions_pending_queue_intent_false_when_no_confirm_cancel_word(ns):
    assert ns["_mentions_pending_queue_intent"]("เปลี่ยน Box 4x4 เป็น 7x7") is False


def test_ask_question_would_show_fixed_message_for_confirm_word_with_doc_open(ns, doc1):
    """"ยืนยัน" ขณะมีเอกสารเปิดอยู่ (ไม่มีขั้นยืนยัน/คิวแล้ว) ต้องถูกตัวตรวจจับจับได้ (เพื่อให้ ask_question() ตอบ
    ข้อความคงที่ "ตอนนี้ระบบแก้ทันทีแล้วครับ..." แทนการส่งเข้า LLM — ข้อความคงที่นี้เขียนอยู่ใน ask_question() เอง
    ซึ่งหนักเกินจะ extract ทั้งฟังก์ชันผ่าน harness นี้ (เหมือน ask_question ทั้งฟังก์ชันในรอบก่อนๆ) ทดสอบแค่กลไก
    ตรวจจับที่ ask_question() พึ่งพาอยู่โดยตรงแทน"""
    assert ns["_classify_pending_queue_command"]("ยืนยัน") is not None


# ---------- direct parse: รูปแบบชัดเจนด้วยโค้ดล้วนๆ ไม่ผ่าน LLM (_parse_direct_structural_commands) ----------

def test_find_single_rate_label_finds_the_one_percent_label(ns, doc1):
    base_bytes, label_map = ns["get_working_copy"](doc1)
    assert ns["_find_single_rate_label"](base_bytes, label_map) == "ภาษีมูลค่าเพิ่ม 7%"


@pytest.mark.parametrize("query", ["เปลี่ยน VAT เป็น 10%", "แก้ภาษีเป็น10%", "เปลี่ยนภาษีมูลค่าเพิ่มเป็น 10%"])
def test_parse_direct_structural_commands_detects_rate_change(ns, doc1, query):
    base_bytes, label_map = ns["get_working_copy"](doc1)
    commands = ns["_parse_direct_structural_commands"](query, base_bytes, label_map)
    assert commands == [{"type": "change_rate", "label": "ภาษีมูลค่าเพิ่ม 7%", "new_rate": 10.0}]


def test_parse_direct_structural_commands_detects_delete(ns, doc1):
    base_bytes, label_map = ns["get_working_copy"](doc1)
    commands = ns["_parse_direct_structural_commands"]("ลบรายการที่ 3", base_bytes, label_map)
    assert commands == [{"type": "delete_table_item", "item_no": 3}]


def test_parse_direct_structural_commands_returns_none_for_multiple_rate_labels(ns):
    """ไฟล์มี label ที่มี "%" และเป็นสูตรอัตราที่รู้จักมากกว่า 1 ตัว — กำกวมเกินจะเดา ต้องคืน None (ปล่อยให้ agent/
    LLM ตัดสินใจแทน ไม่ใช่เดาสุ่มตัวใดตัวหนึ่ง)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws["A1"] = "ภาษี A 7%"
    ws["B1"] = "=C1*7/100"
    ws["A2"] = "ภาษี B 7%"
    ws["B2"] = "=C2*7/100"
    buf = io.BytesIO()
    wb.save(buf)
    raw = inject_formula_cache(buf.getvalue(), {"B1": 70, "B2": 70})
    label_map = ns["_extract_excel_labels"](raw)

    commands = ns["_parse_direct_structural_commands"]("เปลี่ยน VAT เป็น 10%", raw, label_map)
    assert commands is None


@pytest.mark.parametrize("query", ["เปลี่ยน Box 4x4 เป็น 7x7", "เปลี่ยน box 4X4 เป็น Box 7x7"])
def test_parse_direct_structural_commands_detects_description_change(ns, box_doc, query):
    """"Box 4x4" -> "7x7" (ส่วนท้ายสั้นๆ) ประกอบเป็น "Box 7x7" / "box 4X4" (ตัวพิมพ์ต่าง) -> "Box 7x7" (ชื่อเต็ม)
    ใช้ตรงๆ — ทั้งสองแบบต้อง parse ได้โดยไม่ต้องผ่าน agent เลย"""
    base_bytes, label_map = ns["get_working_copy"](box_doc)
    commands = ns["_parse_direct_structural_commands"](query, base_bytes, label_map)
    assert commands == [{"type": "update_table_item", "item_no": 1, "fields": {"description": "Box 7x7"}}]


def test_parse_direct_structural_commands_detects_qty_field_update(ns, box_doc):
    base_bytes, label_map = ns["get_working_copy"](box_doc)
    commands = ns["_parse_direct_structural_commands"]("เปลี่ยนจำนวน Box 4x4 เป็น 30", base_bytes, label_map)
    assert commands == [{"type": "update_table_item", "item_no": 1, "fields": {"qty": 30.0}}]


def test_parse_direct_structural_commands_detects_unit_price_field_update_by_item_no(ns, box_doc):
    base_bytes, label_map = ns["get_working_copy"](box_doc)
    commands = ns["_parse_direct_structural_commands"]("เปลี่ยนราคารายการที่ 2 เป็น 80", base_bytes, label_map)
    assert commands == [{"type": "update_table_item", "item_no": 2, "fields": {"unit_price": 80.0}}]


def test_parse_direct_structural_commands_returns_none_when_target_not_found(ns, box_doc):
    base_bytes, label_map = ns["get_working_copy"](box_doc)
    assert ns["_parse_direct_structural_commands"]("เปลี่ยน Box 9x9 เป็น 1x1", base_bytes, label_map) is None


def test_resolve_table_target_returns_none_when_description_matches_multiple_rows(ns):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws["A15"] = "ลำดับ"
    ws["B15"] = "รายการ"
    ws["D15"] = "จำนวน"
    ws["E15"] = "หน่วย"
    ws["F15"] = "ราคา/หน่วย"
    ws["G15"] = "จำนวนเงิน"
    for i, row in enumerate((16, 17)):
        ws.cell(row=row, column=1, value=i + 1)
        ws.cell(row=row, column=2, value="Widget")  # รายละเอียดซ้ำกันทั้งสองแถวโดยตั้งใจ
        ws.cell(row=row, column=4, value=1)
        ws.cell(row=row, column=5, value="ชิ้น")
        ws.cell(row=row, column=6, value=10)
        ws.cell(row=row, column=7, value=f"=D{row}*F{row}")
    buf = io.BytesIO()
    wb.save(buf)
    raw = inject_formula_cache(buf.getvalue(), {"G16": 10, "G17": 10})

    table = ns["_get_single_item_table"](raw)
    assert ns["_resolve_table_target"]("Widget", table) is None


def test_parse_direct_structural_commands_returns_none_for_ambiguous_box_change(ns, doc1):
    """fixture ปกติ ("Item A".."Item H") ไม่มีรายการชื่อ "Box 4x4" เลย — กำกวมเกินจะเดา ต้องปล่อยให้ agent
    ตัดสินใจแทน (คืน None)"""
    base_bytes, label_map = ns["get_working_copy"](doc1)
    assert ns["_parse_direct_structural_commands"]("เปลี่ยน Box 4x4 เป็น 7x7", base_bytes, label_map) is None


def test_parse_direct_structural_commands_handles_two_commands_joined_by_และ(ns, doc1):
    base_bytes, label_map = ns["get_working_copy"](doc1)
    commands = ns["_parse_direct_structural_commands"](
        "เปลี่ยน VAT เป็น 10% และลบรายการที่ 3", base_bytes, label_map
    )
    assert commands == [
        {"type": "change_rate", "label": "ภาษีมูลค่าเพิ่ม 7%", "new_rate": 10.0},
        {"type": "delete_table_item", "item_no": 3},
    ]


def test_parse_direct_structural_commands_none_if_any_clause_ambiguous(ns, doc1):
    """ถ้ามีประโยคย่อยไหนกำกวม/ไม่ตรงรูปแบบเลย แม้ประโยคอื่นจะ parse ได้ ต้องปล่อยทั้งข้อความให้ agent ตัดสินใจ
    (ไม่ทำแค่บางส่วนตามที่เข้าใจไปเอง)"""
    base_bytes, label_map = ns["get_working_copy"](doc1)
    commands = ns["_parse_direct_structural_commands"](
        "เปลี่ยน VAT เป็น 10% และเปลี่ยน Box 9x9 เป็น 1x1", base_bytes, label_map
    )
    assert commands is None


# ---------- สถานการณ์จริงเต็มรูปแบบ: parse + apply ทันที (จำลองสิ่งที่ ask_question() ทำจริงทีละคำสั่ง) ----------

def _apply_query(ns, document_id, query):
    """จำลองสิ่งที่ ask_question() ทำจริงสำหรับคำสั่งแก้โครงสร้างแบบชัดเจน — parse แล้ว apply ทันที คืนข้อความ
    ผลลัพธ์รวม (ใช้ในเทสต์เท่านั้น ไม่ใช่โค้ดจริงใน main.py)"""
    doc = ns["get_editable_document_raw"](document_id)
    base_bytes, label_map = ns["get_working_copy"](doc)
    commands = ns["_parse_direct_structural_commands"](query, base_bytes, label_map)
    assert commands is not None, f"คาดว่า '{query}' ต้อง parse ได้ด้วยโค้ดตรงๆ ไม่ต้องผ่าน agent"
    return ns["_apply_structural_commands"](document_id, 1, commands)


def test_full_scenario_vat_rename_delete_then_undo_three_times(ns, box_doc):
    """เหตุการณ์จริงทั้งชุด: "เปลี่ยน VAT เป็น 10%" -> "เปลี่ยน Box 4x4 เป็น 7x7" -> "ลบรายการที่ 3" (ไส้ไก่) —
    ทุกคำสั่ง parse ได้ด้วยโค้ดตรงๆ ไม่ต้องผ่าน agent แก้ไฟล์ทันทีทีละคำสั่ง (3 เวอร์ชันแยกกัน) แล้ว undo 3 ครั้ง
    ต้องกลับไปเป็นไฟล์เดิมทุกค่าเป๊ะ"""
    _apply_query(ns, 1, "เปลี่ยน VAT เป็น 10%")
    _apply_query(ns, 1, "เปลี่ยน Box 4x4 เป็น 7x7")
    _apply_query(ns, 1, "ลบรายการที่ 3")

    assert len(ns["_fake_db"].versions[1]) == 3

    data = _current(ns)
    filled = data["tables"][0]["rows"]
    assert len(filled) == 7
    assert filled[0]["values"]["description"] == "Box 7x7"
    assert not any(r["values"]["description"] == "ไส้ไก่" for r in filled)
    label_values = {l["label"]: l["value"] for l in data["labels"]}
    # label เปลี่ยนชื่อตามข้อความจริงในเซลล์ป้ายหลังแก้ (ดู _current_label_display_name) — "7%" เดิมไม่มีอีกแล้ว
    assert "ภาษีมูลค่าเพิ่ม 7%" not in label_values
    assert label_values["ภาษีมูลค่าเพิ่ม 10%"] != _BOX_VAT_7

    for _ in range(3):
        undo_result = ns["undo_last_edit"](1, 1)
        assert undo_result["matched"] is True

    assert ns["_fake_db"].versions.get(1, []) == []
    final_data = _current(ns)
    final_filled = final_data["tables"][0]["rows"]
    assert len(final_filled) == 8
    assert final_filled[0]["values"]["description"] == "Box 4x4"
    assert any(r["values"]["description"] == "ไส้ไก่" for r in final_filled)
    final_label_values = {l["label"]: l["value"] for l in final_data["labels"]}
    assert final_label_values["ภาษีมูลค่าเพิ่ม 7%"] == _BOX_VAT_7
    assert final_label_values["รวมราคาทั้งสิ้น"] == _BOX_GRAND_7


def test_full_scenario_field_updates_apply_immediately_with_correct_totals(ns, box_doc):
    _apply_query(ns, 1, "เปลี่ยนจำนวน Box 4x4 เป็น 30")
    _apply_query(ns, 1, "เปลี่ยนราคารายการที่ 2 เป็น 80")

    data = _current(ns)
    filled = data["tables"][0]["rows"]
    assert filled[0]["values"]["qty"] == 30
    assert filled[1]["values"]["unit_price"] == 80

    expected_total = _BOX_TOTAL + (30 - 5) * 50 + (80 - 70) * 3
    label_values = {l["label"]: l["value"] for l in data["labels"]}
    assert label_values["รวมเงิน"] == expected_total


def test_full_scenario_multi_command_single_message_creates_two_versions(ns, doc1):
    answer = _apply_query(ns, 1, "เปลี่ยน VAT เป็น 10% และลบรายการที่ 3")
    assert len(ns["_fake_db"].versions[1]) == 2
    assert answer.count("✓") == 2

    data = _current(ns)
    filled = data["tables"][0]["rows"]
    assert len(filled) == 7
    label_values = {l["label"]: l["value"] for l in data["labels"]}
    assert "ภาษีมูลค่าเพิ่ม 7%" not in label_values
    assert label_values["ภาษีมูลค่าเพิ่ม 10%"] != FIXTURE_VAT


def test_apply_structural_commands_reports_failure_without_stopping_remaining_commands(ns, doc1):
    """คำสั่งหนึ่งล้มเหลว (เช่น label ไม่เจอ) ไม่ควรทำให้คำสั่งอื่นในข้อความเดียวกันไม่ได้ทำไปด้วย — แต่ละคำสั่ง
    เป็นอิสระจากกัน"""
    commands = [
        {"type": "change_rate", "label": "label ไม่มีจริง", "new_rate": 10},
        {"type": "delete_table_item", "item_no": 3},
    ]
    answer = ns["_apply_structural_commands"](1, 1, commands)
    assert "✗" in answer
    assert "✓" in answer
    assert len(ns["_fake_db"].versions.get(1, [])) == 1  # เฉพาะคำสั่งที่สำเร็จเท่านั้นที่สร้างเวอร์ชัน


# ---------- routing ระดับ ask_question(): direct parse ต้องถูกลองก่อน _is_structural_edit_command เสมอ ----------
# (review PR #31 — เดิม direct parse ถูกเรียกเฉพาะเมื่อ _is_structural_edit_command(query) เป็น True เท่านั้น
# "เปลี่ยน Box 4x4 เป็น 7x7"/"เปลี่ยนจำนวน Box 4x4 เป็น 30" ไม่มีคำที่ฟังก์ชันนั้นรู้จักเลย เลยไม่เคยไปถึง
# _parse_direct_structural_commands — ตกไป classifier/agent ทั้งที่ parse ได้ชัดเจนด้วยโค้ดเฉยๆ)

def test_resolve_direct_structural_dispatch_for_description_change_bypasses_is_structural_edit_command(ns, box_doc):
    """"เปลี่ยน Box 4x4 เป็น 7x7" ไม่มีคำไหนที่ _is_structural_edit_command รู้จักเลย (ยืนยันด้วย assert ข้อแรก)
    แต่ _resolve_direct_structural_dispatch ต้องยังจับได้และ apply จริง ไม่ต้องรอผ่าน gate นั้นก่อน"""
    query = "เปลี่ยน Box 4x4 เป็น 7x7"
    assert ns["_is_structural_edit_command"](query) is False

    result = ns["_resolve_direct_structural_dispatch"](query, [box_doc], 1, 1)
    assert result is not None
    assert result["document_id"] == 1
    assert "✓" in result["answer"]

    filled = _current(ns)["tables"][0]["rows"]
    assert filled[0]["values"]["description"] == "Box 7x7"


@pytest.mark.parametrize("query,field,expected_value", [
    ("เปลี่ยนจำนวน Box 4x4 เป็น 30", "qty", 30),
    ("แก้หน่วยของ Box 6x6 เป็น กล่อง", "unit", "กล่อง"),
])
def test_resolve_direct_structural_dispatch_for_field_updates_bypasses_is_structural_edit_command(
    ns, box_doc, query, field, expected_value
):
    assert ns["_is_structural_edit_command"](query) is False

    result = ns["_resolve_direct_structural_dispatch"](query, [box_doc], 1, 1)
    assert result is not None

    filled = _current(ns)["tables"][0]["rows"]
    target_index = 0 if "4x4" in query else 1
    assert filled[target_index]["values"][field] == expected_value


def test_resolve_direct_structural_dispatch_returns_none_for_non_table_label_edit(ns, doc1):
    """"เปลี่ยนชื่อลูกค้าเป็น บริษัท ทดสอบ จำกัด" ไม่ใช่รายการในตาราง ("ชื่อลูกค้า" ไม่ใช่รายละเอียดสินค้าแถวไหน
    เลย) — ต้องคืน None ปล่อยให้ ask_question() ไปทาง classifier (intent="edit") ตามเดิม ไม่ใช่เดา apply ผิดที่"""
    result = ns["_resolve_direct_structural_dispatch"]("เปลี่ยนชื่อลูกค้าเป็น บริษัท ทดสอบ จำกัด", [doc1], 1, 1)
    assert result is None
    assert ns["_fake_db"].versions.get(1, []) == []


def test_resolve_direct_structural_dispatch_returns_none_for_document_reading_question(ns, doc1):
    """"มีรายการอะไรบ้าง" ไม่มีกริยาแก้ไฟล์เลย (ไม่มี "เปลี่ยน"/"แก้"/"ลบ") — parse ไม่ได้ตามธรรมชาติอยู่แล้ว คืน
    None เสมอ ไม่ต้องพึ่ง gate อะไรเพิ่ม (ใน ask_question() จริง คำถามแบบนี้ถูกดักไปตอบด้วย
    read_document_values ตั้งแต่ก่อนถึงจุดนี้อยู่แล้วเช่นกัน — ดู reading_question/document_content_question)"""
    assert ns["_resolve_direct_structural_dispatch"]("มีรายการอะไรบ้าง", [doc1], 1, 1) is None


def test_resolve_direct_structural_dispatch_returns_none_when_no_target_document(ns):
    assert ns["_resolve_direct_structural_dispatch"]("เปลี่ยน VAT เป็น 10%", [], 1, 1) is None


# ---------- routing: โครงสร้างต้องไปทาง agent ไม่ใช่ _match_and_apply_excel_edit ----------

@pytest.mark.parametrize("query", [
    "เปลี่ยน VAT เป็น 10%",
    "เปลี่ยนอัตราภาษีเป็น 10 เปอร์เซ็นต์",
])
def test_structural_command_detects_rate_change_wording(ns, query):
    assert ns["_is_structural_edit_command"](query) is True


@pytest.mark.parametrize("query", [
    "ลบรายการที่ 3",
    "เพิ่มรายการ Box 2x2 5 ชิ้น 30 บาท",
])
def test_structural_command_detects_table_item_wording(ns, query):
    assert ns["_is_structural_edit_command"](query) is True


def test_edit_command_with_question_word_is_not_reading_question(ns):
    assert ns["_is_document_reading_question"]("แก้ชื่อลูกค้าเป็น บริษัท ดูดี จำกัด ได้ไหม") is False


def test_match_and_apply_excel_edit_refuses_percent_on_formula_cell(ns, doc1):
    """safety net: เซลล์เป้าหมายเป็นสูตร + ค่าใหม่มี "%" -> ห้ามเขียนทับ ต้องบอกให้ใช้ change_rate แทน"""
    ns["_fake_client"].response_text = '{"matched": true, "label": "ภาษีมูลค่าเพิ่ม 7%", "new_value": "10%"}'
    result = ns["_match_and_apply_excel_edit"](1, 1, "เปลี่ยน VAT เป็น 10%")
    assert result["matched"] is False
    assert "change_rate" in result["message"] or "เปลี่ยนอัตรา" in result["message"]

    assert ns["_fake_db"].versions.get(1, []) == []
    vat_info = doc1["label_map"]["ภาษีมูลค่าเพิ่ม 7%"]
    wb_raw = openpyxl.load_workbook(io.BytesIO(doc1["original_bytes"]), data_only=False)
    cell_value = wb_raw[vat_info["sheet"]].cell(row=vat_info["row"], column=vat_info["col"]).value
    assert isinstance(cell_value, str) and cell_value.startswith("=")


def test_match_and_apply_excel_edit_still_works_for_plain_values(ns, doc1):
    ns["_fake_client"].response_text = '{"matched": true, "label": "รวมเงิน", "new_value": "99999"}'
    result = ns["_match_and_apply_excel_edit"](1, 1, "เปลี่ยนรวมเงินเป็น 99999")
    assert result["matched"] is True


def test_document_reading_question_fallback_when_client_fails(ns, doc1):
    ns["_fake_client"].raise_on_create = RuntimeError("simulated API failure")
    active_docs = [doc1]
    answer, used_doc_id = ns["_answer_document_reading_question"]("มีรายการอะไรบ้าง", active_docs, 1, 1)
    assert used_doc_id == 1
    expected = ns["_format_document_values_text"](ns["read_document_values"](doc1))
    assert answer == expected


def test_open_library_file_sets_focus_and_appears_in_active_docs(ns):
    raw = build_fixture_workbook()
    ns["_fake_db"].add_library_file(file_id=42, filename="quote.xlsx", storage_key="lib/quote.xlsx")
    ns["_fake_db"].add_storage_blob("lib/quote.xlsx", raw)

    result = ns["execute_open_library_file"]({"file_id": 42}, user_id=1, chat_id=1)
    assert result.get("opened") is True
    document_id = result["document_id"]

    assert ns["_fake_db"].focus.get(1) == document_id
    active_docs = ns["list_active_editable_documents_by_chat"](1, 1)
    assert any(d["id"] == document_id for d in active_docs)


def test_reading_question_without_legal_keyword_routes_to_document_values(ns):
    assert ns["_is_document_reading_question"]("มีรายการอะไรบ้าง") is True
    assert ns["_has_legal_keyword"]("มีรายการอะไรบ้าง") is False


def test_reading_question_with_legal_keyword_routes_to_rag(ns):
    """"ธุรกิจบัญชีสามมีอะไรบ้าง" ไม่มีคำอ้างอิงไฟล์เลย (ดู _DOCUMENT_REFERENCE_KEYWORDS หลัง tighten) —
    _is_document_reading_question จึงเป็น False อยู่แล้วด้วยตัวมันเอง แต่ที่สำคัญกว่าคือ _has_legal_keyword
    เป็น True ("บัญชี") ซึ่งทำให้ _classify_legal_or_advisor_route ตัดสินเป็น legal ก่อนเช็คอะไรอื่นเสมอ ไม่ว่า
    _is_document_reading_question จะตอบอย่างไรก็ตาม"""
    query = "ธุรกิจบัญชีสามมีอะไรบ้าง"
    assert ns["_is_document_reading_question"](query) is False
    assert ns["_has_legal_keyword"](query) is True
    document_content_question = ns["_is_document_reading_question"](query) and not ns["_has_legal_keyword"](query)
    assert document_content_question is False


def test_answer_document_reading_question_uses_read_document_values_not_rag(ns, doc1):
    active_docs = [doc1]
    answer, used_doc_id = ns["_answer_document_reading_question"]("มีรายการอะไรบ้าง", active_docs, 1, 1)
    assert used_doc_id == 1
    sent_messages = ns["_fake_client"].last_call_kwargs["messages"]
    sent_content = sent_messages[0]["content"]
    assert "ตารางรายการ" in sent_content
    assert "รวมเงิน" in sent_content
    assert "ภาษีมูลค่าเพิ่ม 7%" in sent_content
    assert answer == ns["_fake_client"].response_text


# ---------- ส่วน 2: ห้าม AI แต่งผลการแก้เอง (ทดสอบ guard logic ของ rag_answer โดยตรง ด้วย fake agentic loop) ----------

def test_rag_answer_blocks_structural_route_with_no_tool_called(ns):
    """structural_to_agent แต่ไม่มี structural tool ถูกเรียกเลย (LLM mock เขียน "Preview...พิมพ์ยืนยัน" เอง) ->
    ผู้ใช้ต้องได้ข้อความคงที่ ไม่ใช่ข้อความที่ LLM แต่งขึ้น"""
    ns["_fake_agentic_loop"].set_response("Preview: จะเปลี่ยนราคาเป็น 70 บาท พิมพ์ยืนยันเพื่อดำเนินการ", [], False)
    answer, sources = ns["rag_answer"]("เปลี่ยน Box 4x4 เป็น 7x7", user_id=1, chat_id=1, structural_route=True)
    assert answer == ns["_STRUCTURAL_NO_TOOL_FALLBACK_MESSAGE"]
    assert "Preview" not in answer


def test_rag_answer_allows_structural_route_when_tool_called(ns):
    ns["_fake_agentic_loop"].set_response("จะเปลี่ยนราคาเป็น 70 บาท\n\nพิมพ์ 'ยืนยัน' เพื่อดำเนินการ หรือ 'ยกเลิก'", [], True)
    answer, sources = ns["rag_answer"]("เปลี่ยน Box 4x4 เป็น 7x7", user_id=1, chat_id=1, structural_route=True)
    assert "ยืนยัน" in answer  # ข้อความจริงจาก tool (จำลองผ่าน fake loop) ไม่ถูก block


def test_rag_answer_does_not_override_default_max_tokens(ns):
    """เส้นทางอื่น (ไม่ใช่ advisor) ต้องไม่ส่ง max_tokens เข้า run_agentic_tool_loop เลย — ปล่อยให้ใช้ค่าเริ่มต้น
    2500 เหมือนเดิม (ดู ADVISOR_MAX_TOKENS ที่เฉพาะ advisor_answer เท่านั้นที่ส่งมา)"""
    ns["_fake_agentic_loop"].set_response("คำตอบปกติ", [], False)
    ns["rag_answer"]("มาตรา 10 พ.ร.บ.การประกอบธุรกิจของคนต่างด้าวว่าอย่างไร", user_id=1, chat_id=1)
    assert ns["_fake_agentic_loop"].last_call_args["max_tokens"] is None


def test_rag_answer_blocks_fabricated_edit_claim_without_receipts(ns, doc1):
    """LLM ตอบ "✓ แก้แล้ว" โดยไม่มี receipt จาก tool เลย (และมีเอกสารเปิดอยู่ในแชทจริง) -> ต้องถูก block"""
    ns["_fake_agentic_loop"].set_response("✓ แก้รายการเรียบร้อยแล้วครับ", [], False)
    answer, sources = ns["rag_answer"]("เปลี่ยน Box 4x4 เป็น 7x7", user_id=1, chat_id=1, structural_route=False)
    assert answer == ns["_STRUCTURAL_NO_TOOL_FALLBACK_MESSAGE"]


def test_rag_answer_blocks_claim_with_filler_words_between_verb_and_done(ns, doc1):
    """"ลบรายการที่ 3 แล้วครับ" (มีข้อความคั่นกลางระหว่าง "ลบ" กับ "แล้ว" ไม่ใช่ "ลบแล้ว" ติดกันเป๊ะ) ต้องยังถูก
    จับได้เหมือนเดิม แม้ structural_route=False — ไม่ได้พึ่งคำว่า "ยืนยัน"/"preview"/"✓" เลย — นี่คือเหตุการณ์จริง
    ที่โจทย์ขอให้ยืนยันซ้ำ: agent (mock) อ้างว่าแก้แล้วโดยไม่มี receipt ต้องถูกบล็อก"""
    ns["_fake_agentic_loop"].set_response("ลบรายการที่ 3 แล้วครับ", [], False)
    answer, sources = ns["rag_answer"]("ลบรายการที่ 3", user_id=1, chat_id=1, structural_route=False)
    assert answer == ns["_STRUCTURAL_NO_TOOL_FALLBACK_MESSAGE"]


def test_rag_answer_does_not_block_legal_confirm_word_even_with_active_docs(ns, doc1):
    """คำตอบกฎหมายปกติที่มีคำว่า "ยืนยัน" ปนอยู่ (เช่น "ต้องยืนยันตัวตนก่อนยื่นภาษี") ต้องไม่ถูก block แม้มีเอกสาร
    เปิดอยู่ในแชทจริงก็ตาม — "ยืนยัน"/"preview"/"✓" เช็คเฉพาะตอน structural_route=True เท่านั้น"""
    ns["_fake_agentic_loop"].set_response("ต้องยืนยันตัวตนก่อนยื่นภาษีครับ", [], False)
    answer, sources = ns["rag_answer"]("ต้องทำอะไรก่อนยื่นภาษี", user_id=1, chat_id=1, structural_route=False)
    assert answer != ns["_STRUCTURAL_NO_TOOL_FALLBACK_MESSAGE"]
    assert "ยืนยันตัวตน" in answer


@pytest.mark.parametrize("query,answer_text", [
    ("เพิ่มทุนจดทะเบียนต้องทำอย่างไร", "ถ้าต้องการเพิ่มทุนจดทะเบียน ต้องยื่นคำขอแล้วรออนุมัติ"),
    ("แก้ไขหนังสือบริคณห์สนธิต้องทำอย่างไร", "หากแก้ไขหนังสือบริคณห์สนธิแล้ว ต้องแจ้งกรมพัฒน์ภายใน 14 วัน"),
    ("ผู้ถือหุ้นต่างด้าวเปลี่ยนสัญชาติต้องตรวจอะไรบ้าง", "บริษัทที่เปลี่ยนสัญชาติผู้ถือหุ้นแล้วต้องตรวจ FBA"),
])
def test_rag_answer_does_not_block_normal_legal_sentences_with_active_docs(ns, doc1, query, answer_text):
    """ประโยคกฎหมายปกติที่มีกริยา (เพิ่ม/แก้ไข/เปลี่ยน) + "แล้ว" ปนอยู่ในประโยคเดียวกัน แต่กริยาไม่ได้ติดกับคำนาม
    ของสิ่งที่ถูกแก้ในเอกสารเลย (เช่น "เพิ่มทุน" ไม่ใช่ "เพิ่มรายการ") ต้องไม่ถูก guard บล็อก"""
    ns["_fake_agentic_loop"].set_response(answer_text, [], False)
    answer, sources = ns["rag_answer"](query, user_id=1, chat_id=1, structural_route=False)
    assert answer != ns["_STRUCTURAL_NO_TOOL_FALLBACK_MESSAGE"]


@pytest.mark.parametrize("answer_text", [
    "เปลี่ยน VAT เป็น 10% ให้แล้วครับ",
    "แก้ไขเรียบร้อยครับ",
    "ไฟล์ถูกแก้ไขแล้วครับ",
    "แก้ชื่อลูกค้าเป็น บริษัท ทดสอบ แล้วครับ",
])
def test_rag_answer_blocks_additional_fabricated_claim_phrasings(ns, doc1, answer_text):
    ns["_fake_agentic_loop"].set_response(answer_text, [], False)
    answer, sources = ns["rag_answer"]("แก้ไฟล์ให้หน่อย", user_id=1, chat_id=1, structural_route=False)
    assert answer == ns["_STRUCTURAL_NO_TOOL_FALLBACK_MESSAGE"]


def test_rag_answer_blocks_fabricated_claim_even_with_nonempty_top_chunks(ns, doc1):
    ns["_fake_rag_context"].top_chunks = ["chunk 1 เนื้อหาอะไรสักอย่างจากฐานความรู้", "chunk 2"]
    ns["_fake_agentic_loop"].set_response("ลบรายการที่ 3 แล้วครับ", [], False)
    answer, sources = ns["rag_answer"]("ลบรายการที่ 3", user_id=1, chat_id=1, structural_route=False)
    assert answer == ns["_STRUCTURAL_NO_TOOL_FALLBACK_MESSAGE"]


@pytest.mark.parametrize("query,answer_text", [
    ("เพิ่มทุนจดทะเบียนต้องทำอย่างไร", "ถ้าต้องการเพิ่มทุนจดทะเบียน ต้องยื่นคำขอแล้วรออนุมัติ"),
    ("แก้ไขหนังสือบริคณห์สนธิต้องทำอย่างไร", "หากแก้ไขหนังสือบริคณห์สนธิแล้ว ต้องแจ้งกรมพัฒน์ภายใน 14 วัน"),
    ("ผู้ถือหุ้นต่างด้าวเปลี่ยนสัญชาติต้องตรวจอะไรบ้าง", "บริษัทที่เปลี่ยนสัญชาติผู้ถือหุ้นแล้วต้องตรวจ FBA"),
])
def test_rag_answer_does_not_block_legal_sentences_with_nonempty_top_chunks(ns, doc1, query, answer_text):
    ns["_fake_rag_context"].top_chunks = ["chunk 1 เนื้อหาอะไรสักอย่างจากฐานความรู้", "chunk 2"]
    ns["_fake_agentic_loop"].set_response(answer_text, [], False)
    answer, sources = ns["rag_answer"](query, user_id=1, chat_id=1, structural_route=False)
    assert answer != ns["_STRUCTURAL_NO_TOOL_FALLBACK_MESSAGE"]


def test_rag_answer_does_not_block_normal_legal_answer_without_active_docs(ns):
    ns["_fake_agentic_loop"].set_response("ต้องยืนยันตัวตนก่อนยื่นภาษีครับ", [], False)
    answer, sources = ns["rag_answer"]("ต้องทำอะไรก่อนยื่นภาษี", user_id=None, chat_id=None, structural_route=False)
    assert answer != ns["_STRUCTURAL_NO_TOOL_FALLBACK_MESSAGE"]
    assert "ยืนยันตัวตน" in answer


# ---------- feat/advisor-route: แยกคำถามไฟล์/กฎหมาย/advisor ออกจากกัน + advisor_answer ----------
# เหตุการณ์จริงที่แก้: คำถามแนะนำสินค้า (เช่น "แนะนำ Rack case server ที่คุ้มกว่า 9U") ขณะมีไฟล์เปิดอยู่ ถูก
# _is_document_reading_question เดิม (เช็คแค่คำเชิงคำถามทั่วไปอย่าง "บ้าง"/"ไหม") ตีความผิดเป็นคำถามเกี่ยวกับไฟล์
# ตอบ "ไฟล์ไม่มีข้อมูลรุ่นให้เลือก" หรือหลุดไปทาง rag_answer เดิม (มี context ฐานความรู้กฎหมายเสมอ) ได้คำตอบไม่
# เกี่ยวข้องเรื่อง FBA/BOI แทนที่จะแนะนำสินค้าจริง

@pytest.mark.parametrize("query", ["มีรายการอะไรบ้าง", "ยอดรวมเท่าไร"])
def test_document_reading_keywords_still_match_quotation_questions(ns, query):
    assert ns["_is_document_reading_question"](query) is True


@pytest.mark.parametrize("query", ["มีรุ่นไหนให้เลือกบ้าง พร้อมบอกราคามาเลย", "แนะนำ Rack case server 12U"])
def test_advisor_style_questions_are_not_document_reading_questions(ns, doc1, query):
    """คำถามแนะนำสินค้าทั่วไปขณะมีไฟล์เปิดอยู่ ต้องไม่ถูกตีความเป็นคำถามเกี่ยวกับไฟล์เลย ทั้งทางคำเชิงคำถามทั่วไป
    (ไม่มีคำอ้างอิงไฟล์ใน _DOCUMENT_REFERENCE_KEYWORDS) และทางชื่อรายการ/label ในไฟล์ (fixture ไม่มี "server"/
    "รุ่น" อยู่ในตารางหรือ label เลย)"""
    assert ns["_is_document_reading_question"](query) is False
    assert ns["_query_mentions_table_item_or_label"](query, doc1) is False


def test_classify_question_route_legal_keyword_short_circuits_without_llm_call(ns):
    """"ธุรกิจบัญชีสามมีอะไรบ้าง" มีคำเชิงความรู้กฎหมาย ("บัญชี") ปนมา -> ต้องตัดสินเป็น legal ทันทีด้วยโค้ด
    (_has_legal_keyword) ไม่ต้องเรียก Haiku classifier เลยสักครั้ง (ประหยัด API call)"""
    result = ns["_classify_question_route"]("ธุรกิจบัญชีสามมีอะไรบ้าง")
    assert result == "legal"
    assert ns["_fake_client"].last_call_kwargs is None


@pytest.mark.parametrize("query", ["มีรุ่นไหนให้เลือกบ้าง พร้อมบอกราคามาเลย", "แนะนำ Rack case server 12U"])
def test_classify_question_route_falls_back_to_llm_classifier_for_advisor_questions(ns, query):
    """ไม่มีคำเชิงความรู้กฎหมายปนมาเลย -> ต้องถาม Haiku classifier (ดู _classify_legal_or_advisor_route) ซึ่ง
    ในเทสต์นี้ตั้งค่าให้ตอบ "advisor" กลับมา"""
    ns["_fake_client"].response_text = '{"category": "advisor"}'
    result = ns["_classify_question_route"](query)
    assert result == "advisor"
    assert ns["_fake_client"].last_call_kwargs is not None


def test_classify_legal_or_advisor_route_defaults_to_advisor_on_api_failure(ns):
    ns["_fake_client"].raise_on_create = RuntimeError("simulated API failure")
    assert ns["_classify_legal_or_advisor_route"]("คำถามอะไรสักอย่างที่กำกวม") == "advisor"


# ---------- fix/library-route: เส้นทาง "library" (คลังเอกสาร/ลูกค้า/โปรเจกต์/เทมเพลต ไม่มีไฟล์เปิดอยู่เลย) ----------
# เหตุการณ์จริงที่แก้: "SSF-02"/"ในคลังมีเอกสารอะไรบ้าง"/"คลังใบเสนอราคา"/"ของบริษัทป่าไม้เขียวจำกัด" (ไม่มีไฟล์
# เปิดอยู่) -> [Route] branch=advisor -> ตอบว่า "ไม่มีคลังเอกสารเชื่อมอยู่" เพราะ advisor ใช้ ADVISOR_TOOLS
# (web_search + present_recommendations) ไม่มี LIBRARY_TOOLS เลย — เครื่องมือคลังมีเฉพาะในเส้นทาง rag_answer

@pytest.mark.parametrize("query", ["SSF-02", "ในคลังมีเอกสารอะไรบ้าง", "คลังใบเสนอราคา", "เปิดใบ SSF-02"])
def test_classify_question_route_detects_library_without_llm_call(ns, query):
    """ทั้ง 4 ข้อความต้องถูกตัดสินเป็น library ด้วยโค้ดล้วนๆ (คำสำคัญ/รูปแบบเลขที่เอกสาร) ไม่ต้องเรียก Haiku
    classifier เลยสักครั้ง (ประหยัด API call เหมือน _has_legal_keyword)"""
    result = ns["_classify_question_route"](query)
    assert result == "library"
    assert ns["_fake_client"].last_call_kwargs is None


@pytest.mark.parametrize("query", [
    "แนะนำสาย CAT-6 ที่ดีที่สุด", "หัว RJ-45 ยี่ห้อไหนดี", "กล้อง IPC-2125 ดีไหม",
])
def test_looks_like_library_query_does_not_match_product_model_codes(ns, query):
    """เหตุการณ์จริงที่แก้: _LIBRARY_DOCUMENT_CODE_RE (รูปแบบ "ตัวอักษร-ตัวเลข") แมตช์รหัสรุ่นสินค้าด้วย ทำให้
    คำถามแนะนำสินค้าถูกตัดสินเป็น library ผิดๆ ด้วยโค้ด (ไม่ได้เรียก present_recommendations เลย) — ต้องคืน False
    ปล่อยให้ Haiku classifier ตัดสินแทน (มีตัวอย่างรหัสรุ่นสินค้าสอนไว้แล้วว่าเป็น advisor)"""
    assert ns["_looks_like_library_query"](query) is False


@pytest.mark.parametrize("query", [
    "SSF-02", "SSF-02 ครับ", "เปิดใบ SSF-02", "เอกสารเลขที่ ASA-02",
])
def test_looks_like_library_query_matches_real_document_codes(ns, query):
    """รหัสเอกสารที่ "เกือบทั้งหมด" คือตัวข้อความเอง (เหลือแค่คำลงท้ายสุภาพ/ว่างเปล่า) หรือมีคำอ้างอิงเอกสารปนมา
    ด้วย (ใบ/เอกสาร/เลขที่/ไฟล์/เปิด) ต้องยังคืน True เหมือนเดิม — ไม่ใช่รหัสรุ่นสินค้าที่ปนมาเฉยๆ ในคำถามอื่น"""
    assert ns["_looks_like_library_query"](query) is True


def test_looks_like_library_query_single_word_customer_is_not_enough(ns):
    """"ลูกค้า" เดี่ยวๆ ถูกตัดออกจาก _LIBRARY_ROUTE_KEYWORDS โดยตั้งใจ (กว้างเกินไป) — "แนะนำโปรแกรมจัดการลูกค้า"
    ต้องไม่ถูกตัดสินเป็น library ด้วยโค้ดอีกต่อไป (ปล่อยให้ Haiku classifier ตัดสินแทน) แต่ "รายชื่อลูกค้าทั้งหมด"
    (คำเฉพาะเจาะจงกว่ามาก) ยังต้องเป็น library เหมือนเดิม"""
    assert ns["_looks_like_library_query"]("แนะนำโปรแกรมจัดการลูกค้า") is False
    assert ns["_looks_like_library_query"]("รายชื่อลูกค้าทั้งหมด") is True


def test_classify_question_route_uses_history_for_continuation_message(ns):
    """"ของบริษัทป่าไม้เขียวจำกัด" เดี่ยวๆ ไม่มีคำเชิงคลังเอกสารปนมาเองเลยสักคำ (ไม่ตรง _looks_like_library_query
    ด้วยโค้ด) ต้องตกไปถาม Haiku classifier พร้อมส่งประวัติสนทนา (มี "คลังใบเสนอราคา" อยู่) ไปด้วย — fake classifier
    ในเทสต์นี้ตั้งค่าให้ตอบ "library" กลับมา (จำลองว่า Haiku ตัดสินถูกจากบริบทในประวัติ)"""
    ns["_fake_client"].response_text = '{"category": "library"}'
    history = [
        {"role": "user", "content": "คลังใบเสนอราคา"},
        {"role": "assistant", "content": "ต้องการดูของบริษัทไหนครับ"},
    ]
    result = ns["_classify_question_route"]("ของบริษัทป่าไม้เขียวจำกัด", history=history)
    assert result == "library"
    sent_content = ns["_fake_client"].last_call_kwargs["messages"][0]["content"]
    assert "คลังใบเสนอราคา" in sent_content


@pytest.mark.parametrize("query,expected", [
    ("แนะนำตู้ Rack 9U ที่ดีที่สุด", "advisor"),
    ("ธุรกิจบัญชีสามมีอะไรบ้าง", "legal"),
])
def test_classify_question_route_unaffected_for_non_library_questions(ns, query, expected):
    """คำถามแนะนำสินค้าทั่วไปยังเป็น advisor และคำถามกฎหมายยังเป็น legal เหมือนเดิมทุกประการ ไม่ถูกเหมาเข่งเป็น
    library ผิดๆ (ไม่ตั้งค่า fake client ไว้เลย — ถ้าหลุดไปเรียก Haiku classifier โดยไม่ได้ตั้งใจ จะได้ "advisor"
    เสมอ ซึ่งพอดีตรงกับ query แรกอยู่แล้ว ต้องเช็ค legal แยกด้วยเพื่อยืนยันว่าไม่ได้บังเอิญผ่านเพราะ fallback)"""
    assert ns["_classify_question_route"](query) == expected


def test_resolve_library_route_access_requires_login(ns):
    assert ns["_resolve_library_route_access"]("library", None) == ns["_LIBRARY_LOGIN_REQUIRED_MESSAGE"]
    assert ns["_resolve_library_route_access"]("library", 1) is None
    assert ns["_resolve_library_route_access"]("advisor", None) is None
    assert ns["_resolve_library_route_access"]("legal", None) is None


def test_rag_answer_skips_legal_context_for_library_route(ns):
    """rag_answer(skip_legal_context=True) ต้องส่งต่อ skip_legal_context=True ให้ _prepare_rag_context ตรงๆ
    (ใช้โดย ask_question() สำหรับเส้นทาง library เท่านั้น — ไม่กระทบเส้นทาง legal ที่ไม่ส่ง parameter นี้มาเลย)"""
    ns["_fake_agentic_loop"].set_response("มีเอกสาร 3 รายการครับ", [], False)
    ns["rag_answer"]("ในคลังมีเอกสารอะไรบ้าง", user_id=1, chat_id=1, skip_legal_context=True)
    assert ns["_fake_rag_context"].last_call_kwargs["skip_legal_context"] is True


def test_advisor_answer_sends_chat_history_and_has_no_legal_kb_text_in_system_prompt(ns):
    """advisor_answer ต้องส่งประวัติสนทนาเข้าไปด้วย (ให้ AI เข้าใจว่าผู้ใช้เพิ่งตอบคำถามที่เคยถูกถามกลับไปแล้ว)
    และ system prompt ต้องไม่มีคำเชิงความรู้ฐาน KB กฎหมายเฉพาะทาง (BOI/FBA/พ.ร.บ.ต่างด้าว) ปนเลยสักคำ — ต่างจาก
    system prompt ของ rag_answer (_prepare_rag_context) ที่มีคำพวกนี้อยู่เสมอ"""
    ns["_fake_agentic_loop"].set_response("แนะนำรุ่น A, B, C ครับ", [], False)
    history = [
        {"role": "user", "content": "ฉันกำลังหา Rack case server ที่น่าจะคุ้มกว่า 9U แนะนำหน่อย"},
        {"role": "assistant", "content": "งบประมาณเท่าไรครับ"},
        {"role": "user", "content": "งบไม่จำกัด ส่วนอย่างอื่นเอาแบบดีที่สุด"},
    ]
    answer, sources = ns["advisor_answer"](
        "มีรุ่นไหนให้เลือกบ้าง พร้อมบอกราคามาเลย", history=history, user_id=1, chat_id=1, active_docs=[],
    )
    assert sources == []
    call_args = ns["_fake_agentic_loop"].last_call_args
    # ล็อกอินอยู่ (user_id/chat_id ทั้งคู่) -> ต้องมี SEARCH_LIBRARY_TOOL เพิ่มเข้ามาด้วยเสมอ (ดู
    # test_advisor_answer_includes_search_library_tool_only_when_logged_in สำหรับเคส guest)
    assert call_args["tools_override"] == ns["ADVISOR_TOOLS"] + [ns["SEARCH_LIBRARY_TOOL"]]
    # advisor ต้องส่ง max_tokens=ADVISOR_MAX_TOKENS (8000) เข้า run_agentic_tool_loop เสมอ — เดิมใช้ 2500 เท่า
    # เส้นทางอื่น แต่ present_recommendations เขียนรายละเอียดสินค้าหลายรุ่นยาวจนชน max_tokens กลางคันมาแล้วจริง
    assert call_args["max_tokens"] == ns["ADVISOR_MAX_TOKENS"]
    sent_messages = call_args["messages"]
    assert len(sent_messages) == len(history) + 1
    assert sent_messages[-1]["content"] == "มีรุ่นไหนให้เลือกบ้าง พร้อมบอกราคามาเลย"
    system_prompt = call_args["system_prompt"]
    for forbidden in ("BOI", "FBA", "พ.ร.บ.การประกอบธุรกิจของคนต่างด้าว"):
        assert forbidden not in system_prompt


def test_advisor_answer_offers_replacement_when_asked_item_exists_in_open_quotation_table(ns, box_doc):
    """มีใบเสนอราคาเปิดอยู่ และคำถามเอ่ยถึงรายการที่มีอยู่จริงในตารางตรงๆ ("Box 4x4") -> ต่อท้ายคำตอบด้วยข้อความ
    เสนอ "แทนรายการที่ N ด้วยรุ่นที่เลือก" (การแทนจริงยังไม่ต่อสายใน PR นี้ — แค่ข้อความเสนอ)"""
    ns["_fake_agentic_loop"].set_response("Box 4x4 รุ่นที่คุ้มกว่าคือ ... ครับ", [], False)
    answer, _ = ns["advisor_answer"](
        "Box 4x4 มีรุ่นที่คุ้มกว่านี้ไหม", history=[], user_id=1, chat_id=1, active_docs=[box_doc],
    )
    assert ns["_ADVISOR_REPLACEMENT_OFFER_LINE"] in answer


def test_advisor_answer_does_not_offer_replacement_when_question_unrelated_to_open_table(ns, box_doc):
    ns["_fake_agentic_loop"].set_response("แนะนำรุ่น A, B, C ครับ", [], False)
    answer, _ = ns["advisor_answer"](
        "แนะนำ Rack case server 12U", history=[], user_id=1, chat_id=1, active_docs=[box_doc],
    )
    assert ns["_ADVISOR_REPLACEMENT_OFFER_LINE"] not in answer


# ---------- feat/advisor-route (review fix): web_search แยกคำตอบเป็นหลาย text block กลางประโยค + pause_turn ----------
# เหตุการณ์จริงที่แก้: web_search ทำให้ Claude แบ่งคำตอบเป็นหลาย text block ตรงจุดอ้างอิงแหล่งที่มา (มักกลาง
# ประโยค) เดิมต่อ text block ด้วย "\n\n" ทำให้ประโยคแตกเป็นย่อหน้าใหม่กลางคัน — กระทบเส้นทาง advisor ที่ค้นเว็บบ่อย
# ที่สุด แต่ใช้ฟังก์ชันร่วมกับเส้นทางกฎหมายเดิมด้วย (ดู _build_final_answer_with_citations) และ stop_reason
# "pause_turn" (server tool อย่าง web_search ใช้เวลานาน) เดิมถูกตีความเป็น "จบเทิร์นแล้ว" (break) ทั้งที่ต้องส่ง
# messages เดิมกลับไปเรียกต่อในรอบถัดไป (ดู run_agentic_tool_loop)

def test_build_final_answer_with_citations_joins_text_blocks_without_breaking_sentences(ns):
    blocks = [
        _FakeAgenticTextBlock("กล้องรุ่น A ราคาประมาณ"),
        _FakeAgenticTextBlock(" 1,500 บาท", citations=[_FakeCitation("https://a.com", "ร้าน A")]),
        _FakeAgenticTextBlock(" ตามที่ประกาศไว้", citations=[_FakeCitation("https://a.com", "ร้าน A")]),
        _FakeAgenticTextBlock(" และยังมีอีกร้านที่เสนอราคา"),
        _FakeAgenticTextBlock(" 1,450 บาท", citations=[_FakeCitation("https://b.com", "ร้าน B")]),
    ]
    text = ns["_build_final_answer_with_citations"](blocks)

    # ต่อกันเป็นประโยคเดียวต่อเนื่อง ไม่มี "\n\n" คั่นกลางประโยคจาก citation เลย ("\n\n" ก่อนหัวข้อ "แหล่งที่มา:"
    # เป็นตัวคั่นที่ตั้งใจเติมเอง ไม่ใช่สิ่งที่ต้องเช็ค — ตัดส่วนท้ายนั้นออกก่อน)
    body = text.split("\n\nแหล่งที่มา:")[0]
    assert body == "กล้องรุ่น A ราคาประมาณ 1,500 บาท ตามที่ประกาศไว้ และยังมีอีกร้านที่เสนอราคา 1,450 บาท"
    assert "\n\n" not in body

    # citation ตัดซ้ำตาม url (https://a.com เจอ 2 ครั้งแต่นับครั้งเดียว) ไม่เกิน 5 แหล่ง รูปแบบ "N. <title> —
    # [ลิงก์](<url>)" (markdown link ไม่ใช่ URL ดิบ — ดู _build_final_answer_with_citations)
    assert text.count("แหล่งที่มา:") == 1
    assert "1. ร้าน A — [ลิงก์](https://a.com)" in text
    assert "2. ร้าน B — [ลิงก์](https://b.com)" in text
    assert text.count("ร้าน A —") == 1


def test_build_final_answer_with_citations_caps_at_five_sources(ns):
    blocks = [
        _FakeAgenticTextBlock(f"ข้อความที่ {i}", citations=[_FakeCitation(f"https://site{i}.com", f"แหล่ง {i}")])
        for i in range(8)
    ]
    text = ns["_build_final_answer_with_citations"](blocks)
    assert text.count("[ลิงก์](https://site") == 5
    assert "https://site5.com" not in text  # index 0-4 เท่านั้น (5 แหล่งแรก)


def test_build_final_answer_with_citations_no_source_heading_without_any_citation(ns):
    blocks = [_FakeAgenticTextBlock("คำตอบธรรมดา"), _FakeAgenticTextBlock(" ไม่มีการอ้างอิงเว็บเลย")]
    text = ns["_build_final_answer_with_citations"](blocks)
    assert text == "คำตอบธรรมดา ไม่มีการอ้างอิงเว็บเลย"
    assert "แหล่งที่มา" not in text


def test_run_agentic_tool_loop_continues_after_pause_turn_then_ends_on_second_call(ns):
    """stop_reason="pause_turn" (server tool เช่น web_search ยังทำงานไม่เสร็จ) ต้องไม่ break ทันที ต้องเรียก API
    ต่อในรอบถัดไปจนกว่าจะได้ stop_reason อื่น (ในเทสต์นี้คือ end_turn ในรอบที่สอง) และข้อความ/citation ของรอบ
    pause_turn (ส่วนแรกของคำตอบ) ต้องไม่หายไป — ต้องถูกสะสมไว้แล้วต่อกับรอบสุดท้ายตอนประกอบคำตอบ (เหตุการณ์จริงที่
    แก้: เดิมใช้แค่ response.content ของรอบสุดท้ายรอบเดียว ทำให้ข้อความ/citation ของรอบ pause_turn ก่อนหน้าหายไป)"""
    real_loop = ns["_real_run_agentic_tool_loop"]
    fake_client = ns["_fake_client"]
    fake_client.response_queue = [
        _FakeAgenticResponse(
            [_FakeAgenticTextBlock("ส่วนแรก ", citations=[_FakeCitation("https://a.com", "แหล่ง A")])],
            stop_reason="pause_turn",
        ),
        _FakeAgenticResponse(
            [_FakeAgenticTextBlock("ส่วนหลัง", citations=[_FakeCitation("https://b.com", "แหล่ง B")])],
            stop_reason="end_turn",
        ),
    ]
    answer, receipts, structural = real_loop(
        "system prompt ทดสอบ", [{"role": "user", "content": "แนะนำกล้องหน่อย"}],
        user_id=None, chat_id=None, tools_override=[],
    )
    assert answer.startswith("ส่วนแรก ส่วนหลัง")
    assert "แหล่ง A — [ลิงก์](https://a.com)" in answer
    assert "แหล่ง B — [ลิงก์](https://b.com)" in answer
    assert fake_client.call_count == 2
    assert receipts == []
    assert structural is False


def test_build_final_answer_with_citations_drops_text_written_before_search(ns):
    """ข้อความที่ Claude เขียนไว้ "ก่อน" เรียก web_search (เช่น "ขอค้นเพิ่มเติม...") ต้องไม่หลุดเข้าคำตอบสุดท้าย —
    ใช้เฉพาะ text block ที่อยู่ "หลัง" server_tool_use/web_search_tool_result block สุดท้ายเท่านั้น"""
    blocks = [
        _FakeAgenticTextBlock("ขอค้นเพิ่มเติมเพื่อหาราคาไทยที่ชัดเจน:"),
        _FakeAgenticOtherBlock("server_tool_use"),
        _FakeAgenticOtherBlock("web_search_tool_result", content=[{"url": "https://a.com", "title": "A"}]),
        _FakeAgenticTextBlock("คำตอบ"),
    ]
    text = ns["_build_final_answer_with_citations"](blocks)
    assert text == "คำตอบ"
    assert "ขอค้นเพิ่มเติม" not in text


def test_run_agentic_tool_loop_present_recommendations_filters_unverified_source_urls(ns):
    """Claude เรียก present_recommendations 3 รุ่น (url 2 อันอยู่ในผลค้นจริง, 1 อันไม่อยู่เลย — แต่งขึ้นเอง) ->
    ต้องเหลือแค่ 2 รุ่นที่ยืนยันได้ เลข 1-2 มี [ลิงก์] รุ่นที่ 3 (แต่ง url) ต้องไม่โผล่มาในคำตอบเลย และบันทึกแค่
    รายการที่เหลือ (2 รุ่น) ไว้ที่ chat_last_recommendations"""
    real_loop = ns["_real_run_agentic_tool_loop"]
    fake_client = ns["_fake_client"]
    tool_input = {
        "criteria": ["ความจุ", "ความเร็ว", "ราคา"],
        "items": [
            {
                "name": "รุ่น A", "key_specs": "สเปก A", "price_thb": 1000, "price_note": None, "unit": "ชิ้น",
                "source_url": "https://a.com", "compatible": True, "compatibility_note": None,
            },
            {
                "name": "รุ่น B", "key_specs": "สเปก B", "price_thb": 2000, "price_note": None, "unit": "ชิ้น",
                "source_url": "https://b.com", "compatible": False, "compatibility_note": "แรงดันไฟไม่ตรงกัน",
            },
            {
                "name": "รุ่น C (แต่งขึ้นเอง)", "key_specs": "สเปก C", "price_thb": 3000, "price_note": None,
                "unit": "ชิ้น", "source_url": "https://ไม่มีจริง.com", "compatible": True,
                "compatibility_note": None,
            },
        ],
        "impacts": ["ต้องเปลี่ยนสายไฟให้รองรับด้วย"],
        "summary": "แนะนำรุ่น A และ B ครับ"
    }
    search_result_block = _FakeAgenticOtherBlock("web_search_tool_result", content=[
        {"url": "https://a.com", "title": "เว็บ A"}, {"url": "https://b.com", "title": "เว็บ B"},
    ])
    tool_use_block = _FakeAgenticToolUseBlock("tool_1", "present_recommendations", tool_input)
    fake_client.response_queue = [
        _FakeAgenticResponse([search_result_block, tool_use_block], stop_reason="tool_use"),
        _FakeAgenticResponse([_FakeAgenticTextBlock("จบ")], stop_reason="end_turn"),
    ]

    answer, receipts, structural = real_loop(
        "system prompt ทดสอบ", [{"role": "user", "content": "แนะนำกล้องหน่อย"}],
        user_id=1, chat_id=1, tools_override=ns["ADVISOR_TOOLS"],
    )

    assert "1. รุ่น A" in answer
    assert "2. รุ่น B" in answer
    assert "[ลิงก์](https://a.com)" in answer
    assert "[ลิงก์](https://b.com)" in answer
    assert "รุ่น C" not in answer
    assert "⚠ แรงดันไฟไม่ตรงกัน" in answer
    assert "ผลกระทบต่อรายการอื่น:" in answer
    assert "ต้องเปลี่ยนสายไฟให้รองรับด้วย" in answer

    saved = ns["_fake_db"].last_recommendations.get(1)
    assert saved is not None
    assert len(saved) == 2
    assert saved[0]["name"] == "รุ่น A"


# ---------- fix/advisor-robust: advisor ค้นเว็บแล้วไม่เรียก present_recommendations จนจบด้วยข้อความว่าง ----------
# เหตุการณ์จริงที่แก้: log มี [Route] branch=advisor แต่ไม่มี [ToolCall] present_recommendations เลย — Claude ค้น
# เว็บเสร็จแล้วแต่ไม่ยอมเรียก present_recommendations เอง จนจบเทิร์น (ข้อความว่าง) หรือชน max_iterations ไปเลย
_FORCED_RECOMMENDATION_INPUT = {
    "criteria": ["ความจุ", "ราคา"],
    "items": [{
        "name": "รุ่น D (บังคับเรียก)", "key_specs": "สเปก D", "price_thb": 4500, "price_note": None,
        "unit": "ชิ้น", "source_url": "https://shop.com", "compatible": True, "compatibility_note": None,
    }],
    "impacts": [], "summary": "แนะนำรุ่น D ครับ",
}


def test_run_agentic_tool_loop_forces_present_recommendations_when_ends_empty_after_search(ns):
    """ค้นเว็บ 3 ครั้ง (จำลองรวมอยู่ในรอบเดียว — Anthropic resolve web_search ที่ฝั่ง server เอง) แล้วจบด้วย
    stop_reason=end_turn โดยไม่มี text block เลย (คำตอบว่างเปล่า) และไม่เรียก present_recommendations เลย ->
    ต้องบังคับเรียกอีกครั้งด้วย tool_choice={"type":"tool","name":"present_recommendations"} เพราะมี
    verified_source_urls จากผลค้นจริงอยู่แล้ว"""
    real_loop = ns["_real_run_agentic_tool_loop"]
    fake_client = ns["_fake_client"]
    search_rounds = []
    search_urls = ["https://shop.com", "https://shop1.com", "https://shop2.com"]  # shop.com ตรงกับ source_url
    # ของ _FORCED_RECOMMENDATION_INPUT ให้รอดการตรวจ verified_source_urls ใน execute_present_recommendations
    for i, url in enumerate(search_urls):
        search_rounds.append(_FakeAgenticOtherBlock("server_tool_use"))
        search_rounds.append(_FakeAgenticOtherBlock(
            "web_search_tool_result", content=[{"url": url, "title": f"ร้าน {i}"}],
        ))
    fake_client.response_queue = [
        _FakeAgenticResponse(search_rounds, stop_reason="end_turn"),
        _FakeAgenticResponse(
            [_FakeAgenticToolUseBlock("forced_1", "present_recommendations", _FORCED_RECOMMENDATION_INPUT)],
            stop_reason="tool_use",
        ),
    ]

    answer, receipts, structural = real_loop(
        "system prompt ทดสอบ", [{"role": "user", "content": "แนะนำ Rack case หน่อย"}],
        user_id=1, chat_id=1, tools_override=ns["ADVISOR_TOOLS"],
    )

    assert fake_client.call_count == 2
    assert fake_client.last_call_kwargs["tool_choice"] == {"type": "tool", "name": "present_recommendations"}
    # tools ของการเรียกบังคับต้องเป็น tools_for_this_call เดิมทั้งชุด (มีทั้ง web_search และ
    # present_recommendations) ไม่ใช่ตัดเหลือ present_recommendations ตัวเดียว — ประวัติ messages มี
    # server_tool_use/web_search_tool_result ติดมาแล้ว ตัด web_search ออกจาก tools เสี่ยง API จริง 400
    sent_tool_names = {t.get("name") for t in fake_client.last_call_kwargs["tools"] if isinstance(t, dict)}
    assert "web_search" in sent_tool_names
    assert "present_recommendations" in sent_tool_names
    # จบด้วย end_turn (assistant content ล่าสุดไม่มี text เลย) — ต้องเติม user turn สั้นๆ ขอสรุปก่อนเรียกบังคับ
    # เสมอ (กัน assistant content ว่างๆ ชนกับ tool_choice) — messages ที่ส่งไปต้องจบด้วย role "user" จริงๆ
    sent_messages = fake_client.last_call_kwargs["messages"]
    assert sent_messages[-1] == {"role": "user", "content": "สรุปผลจากการค้นหาข้างต้นให้ผู้ใช้ตอนนี้เลย"}
    assert "1. รุ่น D (บังคับเรียก)" in answer
    assert receipts == []
    assert structural is False


def test_run_agentic_tool_loop_forces_present_recommendations_when_hitting_max_iterations(ns):
    """ชน max_iterations (ตั้งไว้ต่ำๆ ในเทสต์นี้) โดยยังไม่เคยเรียก present_recommendations เลย (เรียก
    open_library_file ซ้ำๆ แทน — ใช้เป็นแค่ "ตัวเติม" ให้วนหลายรอบ ไม่ได้ทดสอบ open_library_file เอง) ->
    ต้องบังคับเรียก present_recommendations เช่นกันหลังชนเพดาน"""
    real_loop = ns["_real_run_agentic_tool_loop"]
    fake_client = ns["_fake_client"]
    filler_round = _FakeAgenticResponse(
        [
            _FakeAgenticOtherBlock("web_search_tool_result", content=[{"url": "https://shop.com", "title": "ร้าน"}]),
            _FakeAgenticToolUseBlock("filler", "open_library_file", {}),
        ],
        stop_reason="tool_use",
    )
    fake_client.response_queue = [
        filler_round,
        filler_round,
        _FakeAgenticResponse(
            [_FakeAgenticToolUseBlock("forced_1", "present_recommendations", _FORCED_RECOMMENDATION_INPUT)],
            stop_reason="tool_use",
        ),
    ]

    answer, receipts, structural = real_loop(
        "system prompt ทดสอบ", [{"role": "user", "content": "แนะนำ Rack case หน่อย"}],
        user_id=1, chat_id=1, tools_override=ns["ADVISOR_TOOLS"], max_iterations=2,
    )

    assert fake_client.call_count == 3
    assert fake_client.last_call_kwargs["tool_choice"] == {"type": "tool", "name": "present_recommendations"}
    sent_tool_names = {t.get("name") for t in fake_client.last_call_kwargs["tools"] if isinstance(t, dict)}
    assert "web_search" in sent_tool_names
    assert "present_recommendations" in sent_tool_names
    assert "1. รุ่น D (บังคับเรียก)" in answer


def test_run_agentic_tool_loop_falls_back_to_plain_summary_when_no_verified_url(ns):
    """ไม่มี verified_source_urls เลย (ไม่เคยค้นเว็บได้ผลจริง) และจบด้วยข้อความว่าง ไม่เรียก present_recommendations
    -> ไม่บังคับเรียก present_recommendations ด้วย tool_choice (ไม่มี url ให้ยืนยันเลย) แต่เรียกอีกครั้งด้วย
    tool_choice={"type":"none"} (ยังต้องส่ง tools=tools_for_this_call เดิมเสมอ — ห้ามส่งแบบไม่มี tools เลย
    เพราะประวัติอาจมี tool block ติดมาแล้ว) ให้สรุปเป็นข้อความธรรมดาแทน"""
    real_loop = ns["_real_run_agentic_tool_loop"]
    fake_client = ns["_fake_client"]
    fake_client.response_queue = [
        _FakeAgenticResponse([], stop_reason="end_turn"),
        _FakeAgenticResponse([_FakeAgenticTextBlock("สรุปจากข้อมูลที่มีครับ ยังไม่มีราคาที่ยืนยันได้ชัดเจน")], stop_reason="end_turn"),
    ]

    answer, receipts, structural = real_loop(
        "system prompt ทดสอบ", [{"role": "user", "content": "แนะนำ Rack case หน่อย"}],
        user_id=1, chat_id=1, tools_override=ns["ADVISOR_TOOLS"],
    )

    assert fake_client.call_count == 2
    assert fake_client.last_call_kwargs["tools"] == ns["ADVISOR_TOOLS"]
    assert fake_client.last_call_kwargs["tool_choice"] == {"type": "none"}
    assert answer == "สรุปจากข้อมูลที่มีครับ ยังไม่มีราคาที่ยืนยันได้ชัดเจน"


def test_run_agentic_tool_loop_forced_calls_survive_exception_and_fall_back_gracefully(ns):
    """การเรียกพิเศษหลัง loop (ทั้ง forced tool_choice call และ summary call) ต้อง "กลืน" exception ของตัวเองไว้
    เสมอ ไม่ปล่อยหลุดออกไปถึงผู้ใช้เป็น 500 — จำลอง API ล้มทั้งคู่ (raise_on_create ติดตั้งไว้ตั้งแต่ก่อนเรียก แต่
    response_queue ยังมีรอบหลักรอบแรกให้ผ่านไปตามปกติก่อน — ดู _FakeAnthropicMessages.create: คิวมาก่อน
    raise_on_create เสมอ) -> ต้องได้ข้อความ fallback เดิม ไม่ crash เลย"""
    real_loop = ns["_real_run_agentic_tool_loop"]
    fake_client = ns["_fake_client"]
    fake_client.response_queue = [
        _FakeAgenticResponse(
            [_FakeAgenticOtherBlock("web_search_tool_result", content=[{"url": "https://shop.com", "title": "ร้าน"}])],
            stop_reason="end_turn",
        ),
    ]
    fake_client.raise_on_create = RuntimeError("simulated API failure")

    answer, receipts, structural = real_loop(
        "system prompt ทดสอบ", [{"role": "user", "content": "แนะนำ Rack case หน่อย"}],
        user_id=1, chat_id=1, tools_override=ns["ADVISOR_TOOLS"],
    )

    assert fake_client.call_count == 3  # รอบหลัก 1 + forced (ล้ม) 1 + summary (ล้ม) 1
    assert answer == "งานนี้มีหลายขั้นเกินไป ลองแบ่งเป็นคำสั่งสั้นลงครับ"
    assert receipts == []
    assert structural is False


def test_run_agentic_tool_loop_forces_present_recommendations_when_hitting_max_tokens(ns):
    """เหตุการณ์จริงที่แก้: [AgentLoop] iter=1 stop_reason=max_tokens tools=present_recommendations text_len=451
    — Claude เขียนข้อความเกณฑ์การเลือกไปก่อนแล้วเริ่มเรียก present_recommendations แต่ input ภาษาไทยยาวเกินจน
    ชน max_tokens กลางคัน (JSON ไม่สมบูรณ์) ต้องถือเป็นเงื่อนไขบังคับเรียกซ้ำเหมือนข้อความว่าง/ชน iteration limit
    และต้องตัด tool_use block ที่ไม่สมบูรณ์ทิ้งจาก messages ก่อน ไม่งั้น API จริงจะ 400 (tool_use ไม่มี tool_result
    คู่กัน)"""
    real_loop = ns["_real_run_agentic_tool_loop"]
    fake_client = ns["_fake_client"]
    incomplete_input = {"criteria": ["ความจุ"]}  # จงใจไม่ครบ (ไม่มี items/summary) จำลอง JSON ที่ถูกตัดกลางคัน
    fake_client.response_queue = [
        _FakeAgenticResponse(
            [
                _FakeAgenticOtherBlock("web_search_tool_result", content=[{"url": "https://shop.com", "title": "ร้าน"}]),
                _FakeAgenticTextBlock("เกณฑ์ที่ใช้พิจารณา: ความจุ, ราคา"),
                _FakeAgenticToolUseBlock("incomplete_1", "present_recommendations", incomplete_input),
            ],
            stop_reason="max_tokens",
        ),
        _FakeAgenticResponse(
            [_FakeAgenticToolUseBlock("forced_1", "present_recommendations", _FORCED_RECOMMENDATION_INPUT)],
            stop_reason="tool_use",
        ),
    ]

    answer, receipts, structural = real_loop(
        "system prompt ทดสอบ", [{"role": "user", "content": "แนะนำ Rack case หน่อย"}],
        user_id=1, chat_id=1, tools_override=ns["ADVISOR_TOOLS"],
    )

    assert fake_client.call_count == 2  # รอบหลัก 1 (ชน max_tokens) + forced 1 ไม่ต้องลองรอบหลักซ้ำ
    assert fake_client.last_call_kwargs["tool_choice"] == {"type": "tool", "name": "present_recommendations"}
    sent_messages = fake_client.last_call_kwargs["messages"]
    # ไม่มี tool_use block ค้างอยู่ใน messages ที่ส่งไปเลยสักที่ (ตัวที่ input ไม่สมบูรณ์ต้องถูกตัดทิ้งไปแล้ว)
    for message in sent_messages:
        content = message["content"]
        if isinstance(content, list):
            assert all(block.get("type") != "tool_use" for block in content)
    # ข้อความเกณฑ์ (text block) ที่เขียนไว้ก่อนชน max_tokens ต้องยังอยู่ ไม่ถูกตัดทิ้งไปด้วย
    assistant_messages = [m for m in sent_messages if m["role"] == "assistant"]
    assert len(assistant_messages) == 1
    assert any(
        block.get("type") == "text" and "เกณฑ์" in block.get("text", "")
        for block in assistant_messages[0]["content"]
    )
    assert sent_messages[-1] == {"role": "user", "content": "สรุปผลจากการค้นหาข้างต้นให้ผู้ใช้ตอนนี้เลย"}
    assert "1. รุ่น D (บังคับเรียก)" in answer
    assert receipts == []
    assert structural is False


def test_run_agentic_tool_loop_max_tokens_drops_whole_assistant_message_when_empty_after_strip(ns):
    """ถ้า assistant message รอบที่ชน max_tokens มีแค่ tool_use block ที่ไม่สมบูรณ์ล้วนๆ (ไม่มี text/block อื่นเลย)
    ตัด tool_use ทิ้งแล้วเหลือ content ว่างเปล่า -> ต้องลบทั้ง message ทิ้งไปเลย ไม่ใช่ปล่อยให้มี assistant message
    ที่ content เป็น list ว่างเปล่าค้างอยู่ (ไม่ใช่รูปแบบที่ Anthropic API ยอมรับ) — รอบแรกเรียก open_library_file
    เป็น "ตัวเติม" ให้มี verified_source_urls อยู่แล้วก่อนชน max_tokens (ตัว tool_use ของรอบแรกจับคู่กับ
    tool_result ปกติแล้ว ไม่ใช่ตัวที่ไม่สมบูรณ์ — ต้องไม่ถูกตัดทิ้งไปด้วย)"""
    real_loop = ns["_real_run_agentic_tool_loop"]
    fake_client = ns["_fake_client"]
    fake_client.response_queue = [
        _FakeAgenticResponse(
            [
                _FakeAgenticOtherBlock("web_search_tool_result", content=[{"url": "https://shop.com", "title": "ร้าน"}]),
                _FakeAgenticToolUseBlock("filler", "open_library_file", {}),
            ],
            stop_reason="tool_use",
        ),
        _FakeAgenticResponse(
            [_FakeAgenticToolUseBlock("incomplete_1", "present_recommendations", {})],
            stop_reason="max_tokens",
        ),
        _FakeAgenticResponse(
            [_FakeAgenticToolUseBlock("forced_1", "present_recommendations", _FORCED_RECOMMENDATION_INPUT)],
            stop_reason="tool_use",
        ),
    ]

    answer, receipts, structural = real_loop(
        "system prompt ทดสอบ", [{"role": "user", "content": "แนะนำ Rack case หน่อย"}],
        user_id=1, chat_id=1, tools_override=ns["ADVISOR_TOOLS"],
    )

    sent_messages = fake_client.last_call_kwargs["messages"]
    assistant_messages = [m for m in sent_messages if m["role"] == "assistant"]
    # เหลือแค่ assistant message ของรอบแรก (filler) — รอบที่ชน max_tokens (content ว่างเปล่าหลังตัด tool_use)
    # ถูกลบทั้ง message ทิ้งไปแล้ว
    assert len(assistant_messages) == 1
    assert "1. รุ่น D (บังคับเรียก)" in answer


def test_run_agentic_tool_loop_pause_turn_three_times_does_not_consume_iteration_quota(ns):
    """pause_turn ติดกัน 3 ครั้ง (ต่ำกว่า MAX_PAUSE_TURNS=5) ก่อนจบด้วย end_turn จริง — ต้องไม่กินโควตา
    max_iterations เลย (ตั้ง max_iterations=1 ในเทสต์นี้โดยตั้งใจ — ถ้า pause_turn กินโควตาจะชนเพดานทันทีตั้งแต่
    รอบที่ 2 และตอบ fallback แทนคำตอบจริง)"""
    real_loop = ns["_real_run_agentic_tool_loop"]
    fake_client = ns["_fake_client"]
    fake_client.response_queue = [
        _FakeAgenticResponse([_FakeAgenticTextBlock("ส่วนที่ 1 ")], stop_reason="pause_turn"),
        _FakeAgenticResponse([_FakeAgenticTextBlock("ส่วนที่ 2 ")], stop_reason="pause_turn"),
        _FakeAgenticResponse([_FakeAgenticTextBlock("ส่วนที่ 3 ")], stop_reason="pause_turn"),
        _FakeAgenticResponse([_FakeAgenticTextBlock("ส่วนสุดท้าย")], stop_reason="end_turn"),
    ]

    answer, receipts, structural = real_loop(
        "system prompt ทดสอบ", [{"role": "user", "content": "แนะนำ Rack case หน่อย"}],
        user_id=None, chat_id=None, tools_override=[], max_iterations=1,
    )

    assert fake_client.call_count == 4
    assert answer == "ส่วนที่ 1 ส่วนที่ 2 ส่วนที่ 3 ส่วนสุดท้าย"
    assert answer != "งานนี้มีหลายขั้นเกินไป ลองแบ่งเป็นคำสั่งสั้นลงครับ"


def test_run_agentic_tool_loop_gives_up_after_too_many_pause_turns(ns):
    """pause_turn เกิน MAX_PAUSE_TURNS (5) ติดกัน — ต้องเลิกรอ (ไม่วนไม่รู้จบ) แล้วคืนข้อความที่สะสมได้จริงจากทุก
    รอบ pause_turn ก่อนหน้า (ดีกว่าทิ้งข้อความที่มีอยู่แล้วไปเฉยๆ) ไม่นับซ้ำสองของรอบสุดท้าย"""
    real_loop = ns["_real_run_agentic_tool_loop"]
    fake_client = ns["_fake_client"]
    max_pause = ns["MAX_PAUSE_TURNS"]
    fake_client.response_queue = [
        _FakeAgenticResponse([_FakeAgenticTextBlock(f"ส่วนที่ {i} ")], stop_reason="pause_turn")
        for i in range(max_pause + 1)
    ]

    answer, receipts, structural = real_loop(
        "system prompt ทดสอบ", [{"role": "user", "content": "แนะนำ Rack case หน่อย"}],
        user_id=None, chat_id=None, tools_override=[],
    )

    assert fake_client.call_count == max_pause + 1
    expected = "".join(f"ส่วนที่ {i} " for i in range(max_pause + 1)).strip()
    assert answer == expected
    assert answer.count("ส่วนที่ 0") == 1  # ไม่นับซ้ำสองของเนื้อหารอบใดเลย (ดู response=None ตอนเลิกรอ)


def test_render_recommendation_item_line_percent_encodes_unsafe_url_characters(ns):
    """source_url ที่มีช่องว่าง/อักขระไทย/วงเล็บ ต้อง percent-encode ก่อนใส่ใน markdown link เสมอ กันลิงก์ขาดกลาง
    คันหรือวงเล็บของ url ไปปิด markdown link ก่อนเวลา (ทำให้รุ่นถัดไปหายไปจากคำตอบ) — ตรวจว่าทุกรุ่นยังอยู่ครบ
    เรียงลำดับถูกต้อง ไม่มีรุ่นไหนขาดหายไปเพราะ markdown พัง"""
    url_space_thai = "https://shop.com/สินค้า ราคาถูก"
    url_paren = "https://shop.com/product(v2)"
    url_plain = "https://shop.com/plain"
    url_unsafe_scheme = "javascript:alert(1)"
    items = [
        {
            "name": "รุ่นช่องว่างไทย", "key_specs": "A", "price_thb": 100, "price_note": None, "unit": "ชิ้น",
            "source_url": url_space_thai, "compatible": True, "compatibility_note": None,
        },
        {
            "name": "รุ่นมีวงเล็บ", "key_specs": "B", "price_thb": 200, "price_note": None, "unit": "ชิ้น",
            "source_url": url_paren, "compatible": True, "compatibility_note": None,
        },
        {
            "name": "รุ่นปกติ", "key_specs": "C", "price_thb": 300, "price_note": None, "unit": "ชิ้น",
            "source_url": url_plain, "compatible": True, "compatibility_note": None,
        },
        {
            "name": "รุ่น URL ไม่ปลอดภัย", "key_specs": "D", "price_thb": 400, "price_note": None, "unit": "ชิ้น",
            "source_url": url_unsafe_scheme, "compatible": True, "compatibility_note": None,
        },
    ]
    answer = ns["_render_recommendations_answer"](
        {"criteria": ["ก"], "items": items, "impacts": [], "summary": ""}, False,
    )
    lines = [l for l in answer.split("\n") if l]
    item_lines = [l for l in lines if l[:1].isdigit()]
    assert len(item_lines) == 4  # ครบทั้ง 4 รุ่น ไม่มีรุ่นไหนหายไปเพราะ markdown พัง
    assert item_lines[0].startswith("1. รุ่นช่องว่างไทย")
    assert item_lines[1].startswith("2. รุ่นมีวงเล็บ")
    assert item_lines[2].startswith("3. รุ่นปกติ")
    assert item_lines[3].startswith("4. รุ่น URL ไม่ปลอดภัย")

    for line, original_url in ((item_lines[0], url_space_thai), (item_lines[1], url_paren), (item_lines[2], url_plain)):
        start = line.index("[ลิงก์](") + len("[ลิงก์](")
        end = line.rindex(")")
        encoded = line[start:end]
        assert " " not in encoded
        assert "(" not in encoded and ")" not in encoded
        assert ns["urllib"].parse.unquote(encoded) == original_url

    # URL ที่ไม่ใช่ http(s) ต้องไม่ถูกทำเป็นลิงก์เลย (ตัดส่วน "— [ลิงก์](...)" ทิ้งทั้งหมด)
    assert "[ลิงก์]" not in item_lines[3]


def test_render_recommendations_answer_returns_fixed_message_when_no_items_survive_filtering(ns):
    answer = ns["_render_recommendations_answer"]({"criteria": [], "items": [], "impacts": [], "summary": ""}, False)
    assert answer == ns["_NO_VERIFIED_RECOMMENDATIONS_MESSAGE"]


def test_render_recommendations_answer_appends_offer_line_only_when_document_open(ns):
    result = {
        "criteria": ["ก"], "items": [{
            "name": "รุ่น A", "key_specs": "สเปก", "price_thb": 100, "price_note": None, "unit": "ชิ้น",
            "source_url": "https://a.com", "compatible": True, "compatibility_note": None,
        }],
        "impacts": [], "summary": "",
    }
    with_doc = ns["_render_recommendations_answer"](result, True)
    without_doc = ns["_render_recommendations_answer"](result, False)
    assert ns["_ADVISOR_REPLACEMENT_OFFER_LINE"] in with_doc
    assert ns["_ADVISOR_REPLACEMENT_OFFER_LINE"] not in without_doc


# ---------- feat/recommend-to-document: เชื่อมรายการแนะนำสินค้าเข้าใบเสนอราคา ----------
# box_doc (ดู _BOX_ITEMS ด้านบน) มี 8 รายการ: รายการที่ 5 = "Item E" (1 เครื่อง ราคา 5000)

_SAMPLE_RECOMMENDATIONS = [
    {
        "name": "Rack case 9U รุ่นประหยัด", "key_specs": "เหล็กหนา 1.2mm", "price_thb": 4405.0,
        "price_note": None, "unit": "ตู้", "source_url": "https://shop-a.com/product", "compatible": True,
        "compatibility_note": None,
    },
    {
        "name": "Rack case 9U รุ่นพรีเมียม", "key_specs": "เหล็กหนา 1.5mm กันสนิม", "price_thb": None,
        "price_note": "≈ $423 ราคาต่างประเทศ", "unit": "ตู้", "source_url": "https://shop-b.com/product",
        "compatible": True, "compatibility_note": None,
    },
]


def test_recommendation_replace_command_changes_description_and_price_keeps_qty(ns, box_doc):
    ns["set_chat_last_recommendations"](1, _SAMPLE_RECOMMENDATIONS)
    result = ns["_resolve_recommendation_command_dispatch"]("แทนรายการที่ 5 ด้วยรุ่นที่ 1", [box_doc], 1, 1)
    assert result is not None
    answer = result["answer"]
    assert "✓ แทนรายการที่ 5 (Item E) ด้วย Rack case 9U รุ่นประหยัด" in answer
    assert "ราคา/หน่วย 5,000 → 4,405" in answer
    assert "(ราคาอ้างอิงจาก shop-a.com)" in answer
    assert "พิมพ์ 'ย้อนการแก้'" in answer

    row = _current(ns)["tables"][0]["rows"][4]["values"]
    assert row["description"] == "Rack case 9U รุ่นประหยัด"
    assert row["unit_price"] == 4405.0
    assert row["qty"] == 1  # จำนวนคงเดิม

    undo_result = ns["undo_last_edit"](1, 1)
    assert undo_result["matched"] is True
    restored_row = _current(ns)["tables"][0]["rows"][4]["values"]
    assert restored_row["description"] == "Item E"
    assert restored_row["unit_price"] == 5000


def test_recommendation_replace_command_with_explicit_price(ns, box_doc):
    """ผู้ใช้ระบุราคาเอง ("ราคา X") ถือว่าราคาเปลี่ยนจริงเสมอ (แม้บังเอิญตรงกับราคาเดิมเป๊ะ) ต้องมี "→" และ
    วงเล็บ "(ราคาที่กำหนดเอง)" ไม่ใช่ "(ราคาอ้างอิงจาก ...)" เพราะไม่ได้มาจากเว็บ"""
    ns["set_chat_last_recommendations"](1, _SAMPLE_RECOMMENDATIONS)
    result = ns["_resolve_recommendation_command_dispatch"]("แทนรายการที่ 5 ด้วยรุ่นที่ 1 ราคา 5000", [box_doc], 1, 1)
    row = _current(ns)["tables"][0]["rows"][4]["values"]
    assert row["unit_price"] == 5000
    assert "ราคา/หน่วย 5,000 → 5,000 (ราคาที่กำหนดเอง)" in result["answer"]
    assert "ราคาอ้างอิงจาก" not in result["answer"]


def test_recommendation_replace_command_with_original_price_shows_no_change_wording(ns, box_doc):
    """เหตุการณ์จริงที่แก้: "ราคาเดิม" เคยโชว์ "ราคา/หน่วย 5,000 → 5,000 (ราคาอ้างอิงจาก ...)" ซึ่งชวนเข้าใจผิดว่า
    ราคาเปลี่ยนไปแล้วทั้งที่ไม่ได้อ้างอิงเว็บอะไรเลย ตอนนี้ต้องโชว์ "ราคาเดิม 5,000 บาท" เฉยๆ ไม่มี "→"/วงเล็บเลย"""
    ns["set_chat_last_recommendations"](1, _SAMPLE_RECOMMENDATIONS)
    result = ns["_resolve_recommendation_command_dispatch"]("แทนรายการที่ 5 ด้วยรุ่นที่ 1 ราคาเดิม", [box_doc], 1, 1)
    row = _current(ns)["tables"][0]["rows"][4]["values"]
    assert row["description"] == "Rack case 9U รุ่นประหยัด"
    assert row["unit_price"] == 5000  # ราคาเดิมไม่เปลี่ยน
    assert "ราคาเดิม 5,000 บาท" in result["answer"]
    assert "→" not in result["answer"].splitlines()[0]
    assert "ราคาอ้างอิงจาก" not in result["answer"]
    assert "ยังไม่ได้เปลี่ยนราคา" not in result["answer"]  # ผู้ใช้ตั้งใจเลือกไม่เปลี่ยนเอง ไม่ใช่ไม่มีราคาให้ใช้


def test_recommendation_replace_command_with_markup_percent(ns, box_doc):
    ns["set_chat_last_recommendations"](1, _SAMPLE_RECOMMENDATIONS)
    result = ns["_resolve_recommendation_command_dispatch"](
        "แทนรายการที่ 5 ด้วยรุ่นที่ 1 บวกกำไร 20%", [box_doc], 1, 1,
    )
    row = _current(ns)["tables"][0]["rows"][4]["values"]
    assert row["unit_price"] == 5286.00
    assert "ราคา/หน่วย 5,000 → 5,286 (ราคาอ้างอิงจาก shop-a.com)" in result["answer"]


def test_recommendation_replace_command_with_null_price_only_changes_name_and_notifies(ns, box_doc):
    """รุ่นที่ 2 (รุ่นพรีเมียม) price_thb เป็น null และผู้ใช้ไม่ได้ระบุราคาต่อท้ายคำสั่งเลย -> แก้แค่ชื่อ ราคาเดิม
    คงอยู่ พร้อมข้อความแจ้งให้พิมพ์ราคาเอง — receipt ต้องโชว์ "ราคาเดิม X บาท" เฉยๆ ไม่มี "→"/วงเล็บที่มาเลย"""
    ns["set_chat_last_recommendations"](1, _SAMPLE_RECOMMENDATIONS)
    result = ns["_resolve_recommendation_command_dispatch"]("แทนรายการที่ 5 ด้วยรุ่นที่ 2", [box_doc], 1, 1)
    row = _current(ns)["tables"][0]["rows"][4]["values"]
    assert row["description"] == "Rack case 9U รุ่นพรีเมียม"
    assert row["unit_price"] == 5000  # ราคาเดิมยังไม่เปลี่ยน
    assert "ราคาเดิม 5,000 บาท" in result["answer"]
    assert "→" not in result["answer"].splitlines()[0]
    assert "ราคาอ้างอิงจาก" not in result["answer"]
    assert "ยังไม่ได้เปลี่ยนราคา" in result["answer"]
    assert "≈ $423 ราคาต่างประเทศ" in result["answer"]


def test_recommendation_add_command_with_quantity(ns, doc1):
    """doc1 (build_fixture_workbook) มีความจุ 10 แถว เต็ม 8 รายการ เหลือว่าง 2 แถว (ดู test_apply_add_table_item_
    until_full) — box_doc เต็มพอดี 8/8 ไม่มีที่ว่างให้เพิ่มเลย จึงใช้ doc1 แทนสำหรับเทสต์ "เพิ่มสำเร็จ" นี้"""
    ns["set_chat_last_recommendations"](1, _SAMPLE_RECOMMENDATIONS)
    result = ns["_resolve_recommendation_command_dispatch"]("เพิ่มรุ่นที่ 2 จำนวน 2", [doc1], 1, 1)
    assert result is not None
    # qty parse มาเป็น float (2.0) เสมอ — ข้อความต้องโชว์ "จำนวน=2" ไม่ใช่ "จำนวน=2.0" (ดู _format_money_value)
    assert "จำนวน=2" in result["answer"]
    assert "จำนวน=2.0" not in result["answer"]
    data = _current(ns)
    filled = data["tables"][0]["rows"]
    assert len(filled) == 9
    new_row = filled[-1]["values"]
    assert new_row["description"] == "Rack case 9U รุ่นพรีเมียม"
    assert new_row["qty"] == 2


def test_recommendation_add_command_reports_full_table(ns, doc1):
    """เติมที่ว่างทั้ง 2 แถวที่เหลือของ doc1 จนเต็ม (10/10) แล้วลองเพิ่มอีกต้องถูกปฏิเสธพร้อมแจ้งตรงๆ"""
    ns["set_chat_last_recommendations"](1, _SAMPLE_RECOMMENDATIONS)
    ns["_resolve_recommendation_command_dispatch"]("เพิ่มรุ่นที่ 1", [doc1], 1, 1)
    ns["_resolve_recommendation_command_dispatch"]("เพิ่มรุ่นที่ 1", [doc1], 1, 1)
    result = ns["_resolve_recommendation_command_dispatch"]("เพิ่มรุ่นที่ 1", [doc1], 1, 1)
    assert result is not None
    assert "เต็มแล้ว" in result["answer"]


def test_recommendation_command_without_any_recommendations_yet(ns, box_doc):
    result = ns["_resolve_recommendation_command_dispatch"]("แทนรายการที่ 5 ด้วยรุ่นที่ 1", [box_doc], 1, 1)
    assert result is not None
    assert "ยังไม่มีรายการแนะนำในแชทนี้ ลองถามให้แนะนำสินค้าก่อนครับ" in result["answer"]


def test_recommendation_command_with_out_of_range_model_index(ns, box_doc):
    ns["set_chat_last_recommendations"](1, _SAMPLE_RECOMMENDATIONS)
    result = ns["_resolve_recommendation_command_dispatch"]("แทนรายการที่ 5 ด้วยรุ่นที่ 9", [box_doc], 1, 1)
    assert result is not None
    assert "มีรุ่นแนะนำแค่ 2 รุ่น" in result["answer"]


def test_recommendation_command_dispatch_returns_none_for_unrelated_query(ns, box_doc):
    ns["set_chat_last_recommendations"](1, _SAMPLE_RECOMMENDATIONS)
    assert ns["_resolve_recommendation_command_dispatch"]("แนะนำ Rack case server 12U", [box_doc], 1, 1) is None


# ---------- feat/contextual-recommendation-edit: ตีความคำสั่งแทน/เพิ่มรุ่นแนะนำจากบริบทด้วย LLM เมื่อ regex
# จับไม่ได้ (box_doc รายการที่ 6 = "Item F" จำนวน 4 ราคา 25 — ดู _BOX_ITEMS ด้านบน) ----------

def test_recommendation_intent_llm_resolves_replace_with_quantity_from_chat_history(ns, box_doc):
    """เหตุการณ์จริงที่แก้: แนะนำ HDD -> "อยากได้ 16tb 2 ตัว" -> "เอารุ่นนี้แหละ" (recs เหลือ 1 รุ่น) ->
    "เเทน ลงในรายการเลย" (สระเอซ้ำ+เว้นวรรค) เคยตอบว่าคำสั่งไม่ชัดเจน/ไม่มีข้อมูลการสนทนาที่ผ่านมา เพราะ regex รับ
    แค่ "แทน...ด้วยรุ่นที่ N" เท่านั้น -> ต้องตกไปให้ LLM ตีความจากบริบท (ตาราง+รายการแนะนำ+ประวัติแชท) แทน"""
    ns["set_chat_last_recommendations"](1, [{
        "name": "Seagate SkyHawk AI 16TB", "key_specs": "7200RPM เขียนอ่านต่อเนื่อง", "price_thb": 19500,
        "price_note": None, "unit": "ตัว", "source_url": "https://shop.example/skyhawk", "compatible": True,
        "compatibility_note": None,
    }])
    ns["_fake_db"].chat_messages[1] = [
        {"role": "user", "content": "แนะนำฮาร์ดดิสก์กล้องวงจรปิดให้หน่อย"},
        {"role": "assistant", "content": "ต้องการความจุเท่าไรครับ"},
        {"role": "user", "content": "อยากได้ 16tb 2 ตัว"},
        {"role": "assistant", "content": "แนะนำ Seagate SkyHawk AI 16TB ราคา 19,500 บาท/ตัวครับ"},
    ]
    ns["_fake_client"].response_queue = [
        _FakeAgenticResponse(
            [_FakeAgenticToolUseBlock("call_1", "resolve_recommendation_action", {
                "action": "replace", "target_item_index": 6, "recommendation_index": 1, "quantity": 2,
                "price_mode": "recommended", "explicit_price": None, "markup_percent": None, "question": None,
            })],
            stop_reason="tool_use",
        ),
    ]

    result = ns["_resolve_recommendation_command_dispatch"]("เเทน ลงในรายการเลย", [box_doc], 1, 1)
    assert result is not None
    answer = result["answer"]
    assert "✓ แทนรายการที่ 6 (Item F) ด้วย Seagate SkyHawk AI 16TB" in answer
    assert "ราคา/หน่วย 25 → 19,500 (ราคาอ้างอิงจาก shop.example)" in answer
    assert "จำนวน 4 → 2" in answer

    row = _current(ns)["tables"][0]["rows"][5]["values"]
    assert row["description"] == "Seagate SkyHawk AI 16TB"
    assert row["unit_price"] == 19500
    assert row["qty"] == 2

    new_amount = 2 * 19500
    expected_total = _BOX_TOTAL - 100 + new_amount  # 100 = amount เดิมของ Item F (4*25) ก่อนแก้
    expected_vat = round(expected_total * 0.07, 2)
    expected_grand = expected_total + expected_vat
    totals_by_label = {l["label"]: l["value"] for l in _current(ns)["labels"]}
    assert totals_by_label["รวมเงิน"] == expected_total
    assert totals_by_label["ภาษีมูลค่าเพิ่ม 7%"] == expected_vat
    assert totals_by_label["รวมราคาทั้งสิ้น"] == expected_grand

    assert ns["_fake_client"].call_count == 1
    assert ns["_fake_client"].last_call_kwargs["tool_choice"] == {"type": "tool", "name": "resolve_recommendation_action"}
    assert ns["_fake_client"].last_call_kwargs["model"] == ns["ADVISOR_MODEL"]


def test_recommendation_intent_llm_out_of_range_target_index_does_not_edit(ns, box_doc):
    """index เกินช่วง -> ไม่แก้ไฟล์ แต่ยังต้องได้คำตอบชี้แจงจากโค้ด ไม่ใช่ปล่อยให้ตกไป routing อื่นเงียบๆ (action
    "ask" เท่านั้นที่เป็นกรณีกำกวมจริงๆ — นี่คือ LLM ตั้งใจแก้แต่ส่งเลขมาไม่ตรงช่วง)"""
    ns["set_chat_last_recommendations"](1, _SAMPLE_RECOMMENDATIONS)
    ns["_fake_client"].response_queue = [
        _FakeAgenticResponse(
            [_FakeAgenticToolUseBlock("call_1", "resolve_recommendation_action", {
                "action": "replace", "target_item_index": 99, "recommendation_index": 1, "quantity": None,
                "price_mode": "recommended", "explicit_price": None, "markup_percent": None, "question": None,
            })],
            stop_reason="tool_use",
        ),
    ]
    result = ns["_resolve_recommendation_command_dispatch"]("แทน ลงในรายการเลย", [box_doc], 1, 1)
    assert result is not None
    assert "ไม่พบรายการที่ระบุในตาราง" in result["answer"]
    row = _current(ns)["tables"][0]["rows"][5]["values"]
    assert row["description"] == "Item F"  # ไม่ถูกแก้เลย


def test_recommendation_intent_llm_action_none_returns_none(ns, box_doc):
    ns["set_chat_last_recommendations"](1, _SAMPLE_RECOMMENDATIONS)
    ns["_fake_client"].response_queue = [
        _FakeAgenticResponse(
            [_FakeAgenticToolUseBlock("call_1", "resolve_recommendation_action", {
                "action": "none", "target_item_index": None, "recommendation_index": None, "quantity": None,
                "price_mode": "recommended", "explicit_price": None, "markup_percent": None, "question": None,
            })],
            stop_reason="tool_use",
        ),
    ]
    assert ns["_resolve_recommendation_command_dispatch"]("ใช้งานยังไงเหรอ", [box_doc], 1, 1) is None


def test_recommendation_intent_llm_action_ask_returns_question_without_editing(ns, box_doc):
    ns["set_chat_last_recommendations"](1, _SAMPLE_RECOMMENDATIONS)
    ns["_fake_client"].response_queue = [
        _FakeAgenticResponse(
            [_FakeAgenticToolUseBlock("call_1", "resolve_recommendation_action", {
                "action": "ask", "target_item_index": None, "recommendation_index": None, "quantity": None,
                "price_mode": "recommended", "explicit_price": None, "markup_percent": None,
                "question": "มีรายการที่อาจเป็นไปได้หลายรายการ ต้องการแทนรายการไหนครับ",
            })],
            stop_reason="tool_use",
        ),
    ]
    result = ns["_resolve_recommendation_command_dispatch"]("เอารุ่นนี้แทนไปเลย", [box_doc], 1, 1)
    assert result is not None
    assert result["answer"] == "มีรายการที่อาจเป็นไปได้หลายรายการ ต้องการแทนรายการไหนครับ"
    row = _current(ns)["tables"][0]["rows"][5]["values"]
    assert row["description"] == "Item F"  # ไม่ถูกแก้เลย


def test_recommendation_regex_match_does_not_call_llm(ns, box_doc):
    """คำสั่งที่ตรงรูปแบบ regex เดิมเป๊ะ ("แทนรายการที่ N ด้วยรุ่นที่ K ...") ต้องทำงานโดยไม่เรียก LLM เลย"""
    ns["set_chat_last_recommendations"](1, _SAMPLE_RECOMMENDATIONS)
    result = ns["_resolve_recommendation_command_dispatch"](
        "แทนรายการที่ 5 ด้วยรุ่นที่ 1 บวกกำไร 20%", [box_doc], 1, 1,
    )
    assert result is not None
    assert ns["_fake_client"].call_count == 0


def test_recommendation_intent_llm_not_called_without_recommendations(ns, box_doc):
    """ไม่มีรายการแนะนำในแชทนี้เลย (get_chat_last_recommendations ว่างเปล่า) ต้องไม่เรียก LLM แม้ข้อความจะมีคำกริยา
    ที่น่าจะเป็นเจตนาเชื่อมรายการแนะนำก็ตาม"""
    assert ns["_resolve_recommendation_command_dispatch"]("แทน ลงในรายการเลย", [box_doc], 1, 1) is None
    assert ns["_fake_client"].call_count == 0


def test_sanitize_recommendation_text_strips_cite_tags_quotes_and_trailing_artifacts(ns):
    sanitize = ns["_sanitize_recommendation_text"]
    assert sanitize('<cite index="3-4">ข้อความ</cite>') == "ข้อความ"
    assert sanitize('ราคา 19,500 บาท"}') == "ราคา 19,500 บาท"
    assert sanitize('สเปก 7200RPM">') == "สเปก 7200RPM"
    assert sanitize("มี   ช่องว่าง\n\nซ้ำ") == "มี ช่องว่าง ซ้ำ"
    assert sanitize(None) is None
    assert sanitize(123) == 123


def test_execute_present_recommendations_sanitizes_fields_and_reports_unverified_count(ns):
    verified_urls = {"https://shop.example/a"}
    tool_input = {
        "criteria": ["ความจุ", "ราคา"],
        "items": [
            {
                "name": '<cite index="1-1">SkyHawk 16TB</cite>', "key_specs": "7200RPM", "price_thb": 19500,
                "price_note": None, "unit": "ตัว", "source_url": "https://shop.example/a", "compatible": True,
                "compatibility_note": None,
            },
            {
                "name": "รุ่นที่ยืนยันแหล่งที่มาไม่ได้", "key_specs": "สเปก", "price_thb": 1000, "price_note": None,
                "unit": "ชิ้น", "source_url": "https://unverified.example", "compatible": True,
                "compatibility_note": None,
            },
        ],
        "impacts": [], "summary": "สรุป",
    }
    result = ns["execute_present_recommendations"](tool_input, chat_id=1, verified_urls=verified_urls)
    assert len(result["items"]) == 1
    assert result["items"][0]["name"] == "SkyHawk 16TB"
    assert result["unverified_count"] == 1

    answer = ns["_render_recommendations_answer"](result, False)
    assert "(ตัดออก 1 รุ่นที่ยืนยันแหล่งที่มาจากผลค้นหาไม่ได้)" in answer


def test_answer_document_reading_question_sends_chat_history(ns, doc1):
    ns["_fake_db"].chat_messages[1] = [
        {"role": "user", "content": "สวัสดีครับ"},
        {"role": "assistant", "content": "สวัสดีครับ มีอะไรให้ช่วยไหมครับ"},
    ]
    ns["_answer_document_reading_question"]("มีรายการอะไรบ้าง", [doc1], 1, 1)
    sent_messages = ns["_fake_client"].last_call_kwargs["messages"]
    assert sent_messages[0] == {"role": "user", "content": "สวัสดีครับ"}
    assert sent_messages[1] == {"role": "assistant", "content": "สวัสดีครับ มีอะไรให้ช่วยไหมครับ"}
    assert sent_messages[-1]["role"] == "user"
    assert "ตารางรายการ" in sent_messages[-1]["content"]


def test_is_document_reading_question_treats_replace_command_as_edit_not_reading(ns):
    """เหตุการณ์จริงที่แก้: "เเทน ลงในรายการเลย" (สระเอซ้ำ) มีคำว่า "รายการ" (_DOCUMENT_REFERENCE_KEYWORDS) แต่เดิม
    _DOCUMENT_EDIT_COMMAND_WORDS ไม่มี "แทน" และไม่ normalize "เเ"→"แ" ก่อนเช็ค ทำให้ถูกตีความเป็นคำถามเชิงอ่านไฟล์
    ผิดๆ ทั้งที่เป็นคำสั่งแก้ไฟล์จริงๆ (ทำให้ไปไม่ถึง _resolve_recommendation_command_dispatch เลยสักครั้ง)"""
    assert ns["_is_document_reading_question"]("เเทน ลงในรายการเลย") is False
    assert ns["_is_document_reading_question"]("แทนลงในรายการเลย") is False
    # ยังต้องไม่กระทบคำถามเชิงอ่านไฟล์จริงๆ เลย
    assert ns["_is_document_reading_question"]("มีรายการอะไรบ้าง") is True
    assert ns["_is_document_reading_question"]("ยอดรวมทั้งสิ้นเท่าไร") is True


def test_recommendation_dispatch_checked_before_document_reading_in_routing_order(ns, box_doc):
    """จำลองลำดับการตัดสินใจจริงใน ask_question() (ย้าย _resolve_recommendation_command_dispatch ไปเช็คก่อนบล็อก
    reading_question/document_content_question แล้ว) ด้วยการเรียกฟังก์ชันตัดสินใจทั้งสองตัวตรงๆ ตามลำดับเดียวกัน
    โดยไม่ต้องพึ่ง FastAPI/Request จริงเลย (เหมือนการทดสอบ _resolve_direct_structural_dispatch): "เเทน ลงในรายการ
    เลย" ต้องถูก recommendation dispatch จัดการจบก่อนเลย (ask_question() จะ return ทันทีที่จุดนี้ ไม่มีทางไปถึง
    _is_document_reading_question อีกต่อไป) ส่วน "มีรายการอะไรบ้าง" ต้องตกไปที่ document_reading เหมือนเดิมทุก
    ประการ (dispatch คืน None เพราะไม่มีคำกริยาที่เกี่ยวข้องเลย)"""
    ns["set_chat_last_recommendations"](1, [{
        "name": "Seagate SkyHawk AI 16TB", "key_specs": "7200RPM", "price_thb": 19500, "price_note": None,
        "unit": "ตัว", "source_url": "https://shop.example/skyhawk", "compatible": True,
        "compatibility_note": None,
    }])
    ns["_fake_db"].chat_messages[1] = [{"role": "user", "content": "อยากได้ 16tb 2 ตัว"}]
    ns["_fake_client"].response_queue = [
        _FakeAgenticResponse(
            [_FakeAgenticToolUseBlock("call_1", "resolve_recommendation_action", {
                "action": "replace", "target_item_index": 6, "recommendation_index": 1, "quantity": 2,
                "price_mode": "recommended", "explicit_price": None, "markup_percent": None, "question": None,
            })],
            stop_reason="tool_use",
        ),
    ]

    # ----- ลำดับที่ 1 ใน ask_question(): recommendation dispatch เช็คก่อนเสมอตอนนี้ -----
    replace_dispatch = ns["_resolve_recommendation_command_dispatch"]("เเทน ลงในรายการเลย", [box_doc], 1, 1)
    assert replace_dispatch is not None  # จัดการจบแล้ว — ask_question() return ตรงนี้เลย ไม่เช็ค reading_question ต่อ

    # ----- คำถามอ่านข้อมูลธรรมดา ต้องยังตกไป document_reading เหมือนเดิม (dispatch คืน None ก่อนเสมอ) -----
    reading_dispatch = ns["_resolve_recommendation_command_dispatch"]("มีรายการอะไรบ้าง", [box_doc], 1, 1)
    assert reading_dispatch is None
    assert ns["_is_document_reading_question"]("มีรายการอะไรบ้าง") is True


# ---------- feat/advisor-route: ส่งตารางรายการเป็น context + เกณฑ์เปรียบเทียบ/ผลกระทบ/หน่วยขาย ----------
# ต้องเป็นกติกาทั่วไปที่ใช้ได้กับสินค้า "ทุกประเภท" ไม่ hardcode ตามสินค้าใดสินค้าหนึ่ง — ทดสอบด้วย 3 ประเภทสินค้า
# ที่ต่างกันมาก (กล้องวงจรปิด/ฮาร์ดดิสก์/สายสัญญาณ) ยืนยันว่า system prompt มีตารางรายการจริง + คำสั่งเรื่องเกณฑ์
# เปรียบเทียบและผลกระทบต่อรายการอื่นเหมือนกันทุกประเภท (ไม่ใช่ข้อความเฉพาะกิจของสินค้าใดสินค้าหนึ่ง)

_CAMERA_ITEMS = [("กล้องวงจรปิด 4MP รุ่น A", 8, "ตัว", 1200), ("Switch PoE 8 พอร์ต", 1, "เครื่อง", 2500)]
_HDD_ITEMS = [("ฮาร์ดดิสก์ 4TB สำหรับ NVR", 2, "ลูก", 2800), ("NVR 8 ช่อง", 1, "เครื่อง", 4500)]
_CABLE_ITEMS = [("สาย UTP Cat6 กล่อง 305 ม.", 3, "กล่อง", 1800), ("หัวต่อ RJ45", 100, "ชิ้น", 3)]


@pytest.fixture
def camera_doc(ns):
    raw = _build_quote_workbook_with_items(_CAMERA_ITEMS)
    label_map = ns["_extract_excel_labels"](raw)
    ns["_fake_db"].add_document(1, "quote.xlsx", raw, label_map, chat_id=1, user_id=1)
    return ns["_fake_db"].documents[1]


@pytest.fixture
def hdd_doc(ns):
    raw = _build_quote_workbook_with_items(_HDD_ITEMS)
    label_map = ns["_extract_excel_labels"](raw)
    ns["_fake_db"].add_document(1, "quote.xlsx", raw, label_map, chat_id=1, user_id=1)
    return ns["_fake_db"].documents[1]


@pytest.fixture
def cable_doc(ns):
    raw = _build_quote_workbook_with_items(_CABLE_ITEMS)
    label_map = ns["_extract_excel_labels"](raw)
    ns["_fake_db"].add_document(1, "quote.xlsx", raw, label_map, chat_id=1, user_id=1)
    return ns["_fake_db"].documents[1]


@pytest.mark.parametrize("doc_fixture_name, query, item_text", [
    ("camera_doc", "มีกล้องวงจรปิด 4MP รุ่นอื่นที่คุ้มกว่านี้ไหม", "กล้องวงจรปิด 4MP รุ่น A"),
    ("hdd_doc", "ฮาร์ดดิสก์สำหรับ NVR รุ่นนี้มีตัวเลือกอื่นไหม", "ฮาร์ดดิสก์ 4TB สำหรับ NVR"),
    ("cable_doc", "สาย UTP แบบนี้มีเกรดอื่นให้เลือกไหม", "สาย UTP Cat6 กล่อง 305 ม."),
])
def test_advisor_system_prompt_includes_quotation_table_for_any_product_category(
    ns, request, doc_fixture_name, query, item_text,
):
    """ไม่ว่าจะเป็นสินค้าประเภทไหน (กล้อง/HDD/สาย) advisor_answer ต้องส่งตารางรายการจริงของใบเสนอราคาที่เปิดอยู่
    เป็น context เสมอ (ดู _advisor_document_context) และ system prompt ต้องมีกติกาทั่วไป 3 ข้อ: เกณฑ์เปรียบเทียบ
    3-5 ข้อ, กติกาเรื่อง impacts (ผลกระทบต่อรายการอื่น/อุปกรณ์จำเป็นที่ขาดไป), และหน่วยขายของรุ่นที่แนะนำ —
    ข้อความเดียวกันทุกประเภทสินค้า ไม่ใช่ข้อความเฉพาะกิจของสินค้าใดสินค้าหนึ่ง (หัวข้อ "ผลกระทบต่อรายการอื่น:" ที่
    ผู้ใช้เห็นจริงๆ ตอนนี้มาจากโค้ด _render_recommendations_answer ไม่ใช่ system prompt อีกต่อไป — AI กรอกแค่
    ข้อมูลดิบผ่าน present_recommendations)"""
    doc = request.getfixturevalue(doc_fixture_name)
    ns["_fake_agentic_loop"].set_response("แนะนำรุ่น A, B, C ครับ", [], False)
    ns["advisor_answer"](query, history=[], user_id=1, chat_id=1, active_docs=[doc])

    system_prompt = ns["_fake_agentic_loop"].last_call_args["system_prompt"]
    assert item_text in system_prompt
    assert "เกณฑ์เปรียบเทียบ 3-5 ข้อ" in system_prompt
    assert "impacts" in system_prompt
    assert "NVR/DVR" in system_prompt
    assert "หน่วยขาย" in system_prompt
    # ห้ามบอกว่าแก้ไฟล์ไปแล้ว/กำลังจะแก้ — แจ้งผลกระทบเท่านั้น
    assert "ยังไม่มีการแก้ไฟล์ใดๆ เกิดขึ้นจริง" in system_prompt


def test_advisor_system_prompt_has_no_document_context_section_without_open_document(ns):
    ns["_fake_agentic_loop"].set_response("แนะนำรุ่น A, B, C ครับ", [], False)
    ns["advisor_answer"]("แนะนำ Rack case server 12U", history=[], user_id=1, chat_id=1, active_docs=[])

    system_prompt = ns["_fake_agentic_loop"].last_call_args["system_prompt"]
    assert "เกณฑ์เปรียบเทียบ 3-5 ข้อ" not in system_prompt
    assert "ผลกระทบต่อรายการอื่น" not in system_prompt


# ---------- feat/advisor-route: export บทสนทนาเป็นไฟล์ .txt ----------

def test_is_chat_export_command_matches_known_phrases(ns):
    for phrase in ["ขอบทสนทนา", "export แชท", "บันทึกแชทเป็นไฟล์", "ดาวน์โหลดบทสนทนา", "ช่วยขอบทสนทนานี้หน่อย"]:
        assert ns["_is_chat_export_command"](phrase) is True
    assert ns["_is_chat_export_command"]("ลบรายการที่ 3") is False


def test_build_chat_export_answer_is_pure_text_without_calling_llm(ns):
    """คำสั่ง export ต้องตอบลิงก์ดาวน์โหลดตรงๆ โดยไม่เรียก LLM เลยสักครั้ง (ดู ask_question — ดักก่อนเรียก AI
    ใดๆ เสมอ)"""
    answer = ns["_build_chat_export_answer"](42)
    assert "/api/chat/42/export.txt" in answer
    assert ns["_fake_client"].last_call_kwargs is None


def test_build_chat_export_text_contains_all_messages_in_chronological_order_with_thai_role_labels(ns):
    messages = [
        {"role": "user", "content": "สวัสดีครับ", "created_at": "2026-10-12T12:30:00"},
        {"role": "assistant", "content": "สวัสดีครับ มีอะไรให้ช่วยไหมครับ", "created_at": "2026-10-12T12:30:05"},
        {"role": "user", "content": "ขอบทสนทนานี้หน่อย", "created_at": "2026-10-12T12:35:00"},
    ]
    text = ns["_build_chat_export_text"]("แชททดสอบ", messages)

    assert "แชททดสอบ" in text
    # UTC naive -> เวลาไทย (+7) เสมอ (ดู _format_chat_export_timestamp)
    assert "[2026-10-12 19:30] ผู้ใช้: สวัสดีครับ" in text
    assert "[2026-10-12 19:30] ThaiMeSook: สวัสดีครับ มีอะไรให้ช่วยไหมครับ" in text
    assert "[2026-10-12 19:35] ผู้ใช้: ขอบทสนทนานี้หน่อย" in text

    # เรียงตามเวลาจริง (ลำดับเดียวกับที่ get_chat_messages คืนมา ไม่ใช่สลับเอง)
    first_idx = text.index("สวัสดีครับ")
    second_idx = text.index("มีอะไรให้ช่วยไหมครับ")
    third_idx = text.index("ขอบทสนทนานี้หน่อย")
    assert first_idx < second_idx < third_idx


def test_build_chat_export_text_handles_missing_or_unparsable_timestamp(ns):
    messages = [{"role": "user", "content": "ข้อความไม่มีเวลา", "created_at": None}]
    text = ns["_build_chat_export_text"]("แชททดสอบ", messages)
    assert "[-] ผู้ใช้: ข้อความไม่มีเวลา" in text


def test_strip_markdown_for_export_removes_bold_and_rewrites_links(ns):
    assert ns["_strip_markdown_for_export"]("**ตัวหนา**") == "ตัวหนา"
    assert ns["_strip_markdown_for_export"]("__ตัวหนา__") == "ตัวหนา"
    assert ns["_strip_markdown_for_export"]("[ลิงก์](https://x.com)") == "ลิงก์ (https://x.com)"
    assert ns["_strip_markdown_for_export"]("# หัวข้อใหญ่") == "หัวข้อใหญ่"


def test_build_chat_export_text_strips_markdown_from_message_content(ns):
    messages = [{
        "role": "assistant",
        "content": "แนะนำ **รุ่น A** — [ลิงก์](https://x.com)",
        "created_at": "2026-10-12T12:00:00",
    }]
    text = ns["_build_chat_export_text"]("แชททดสอบ", messages)
    assert "แนะนำ รุ่น A — ลิงก์ (https://x.com)" in text
    assert "**" not in text
    assert "[ลิงก์]" not in text


# ---------- feat/library-search-and-api-errors ส่วน A: API error ต้องไม่เป็น 500 + "คลัง" ตอนมีไฟล์เปิดอยู่ ----------
# เหตุการณ์จริงที่แก้ (1): เครดิต API หมด -> anthropic.BadRequestError 400 "credit balance is too low" หลุดออก
# จาก ask_question -> FastAPI คืน 500 เป็น text/plain -> หน้าเว็บ parse JSON ไม่ได้ พัง
# เหตุการณ์จริงที่แก้ (2): เปิด SSF-02 อยู่ -> "ในคลังมีเอกสารอะไรบ้าง" -> branch=document_reading (เพราะ
# "เอกสาร" อยู่ใน _DOCUMENT_REFERENCE_KEYWORDS) -> แสดงค่าในไฟล์แทนรายการในคลังทั้งหมด

def test_categorize_anthropic_api_error_detects_credit_balance_exhausted(ns):
    error = ns["anthropic"].BadRequestError("Your credit balance is too low to access the Anthropic API", status_code=400)
    message = ns["_categorize_anthropic_api_error"](error)
    assert "โควตา" in message
    assert "แจ้งผู้ดูแลระบบ" in message


def test_categorize_anthropic_api_error_detects_rate_limit(ns):
    error = ns["anthropic"].RateLimitError("rate limited", status_code=429)
    message = ns["_categorize_anthropic_api_error"](error)
    assert "หนาแน่น" in message


def test_categorize_anthropic_api_error_detects_overloaded_by_status_code(ns):
    """Anthropic "overloaded" ไม่มี subclass เฉพาะใน SDK — มาเป็น status_code=529 บน InternalServerError
    (SDK แมป >=500 ทั้งหมดเป็นคลาสนี้) ต้องยังได้ข้อความ "หนาแน่น" เหมือน RateLimitError"""
    error = ns["anthropic"].InternalServerError("Overloaded", status_code=529)
    message = ns["_categorize_anthropic_api_error"](error)
    assert "หนาแน่น" in message


def test_categorize_anthropic_api_error_falls_back_to_generic_message(ns):
    error = ns["anthropic"].InternalServerError("some other internal error", status_code=500)
    message = ns["_categorize_anthropic_api_error"](error)
    assert "ขัดข้องชั่วคราว" in message
    assert "หนาแน่น" not in message
    assert "โควตา" not in message


def test_ask_question_api_error_response_returns_normal_json_shape(ns):
    """ต้องคืน dict รูปแบบเดียวกับคำตอบปกติ ({"answer","sources","chat_id"}) ให้ ask_question() return ได้ตรงๆ —
    sources ว่างเสมอ ไม่ raise ออกไปให้ FastAPI คืน 500 เด็ดขาด"""
    error = ns["anthropic"].RateLimitError("rate limited", status_code=429)
    result = ns["_ask_question_api_error_response"](error, chat_id=42)
    assert result == {"answer": ns["_categorize_anthropic_api_error"](error), "sources": [], "chat_id": 42}


def test_is_document_reading_question_excludes_library_keyword(ns):
    """มีคำว่า "คลัง" ปนมา -> ไม่ใช่ reading question เด็ดขาด แม้มีคำอ้างอิงไฟล์ (เช่น "เอกสาร") ปนมาด้วยก็ตาม —
    เหตุการณ์จริงที่แก้: เปิด SSF-02 อยู่ -> "ในคลังมีเอกสารอะไรบ้าง" เคยถูกตีความเป็นคำถามเกี่ยวกับไฟล์ที่เปิดอยู่
    (มีคำว่า "เอกสาร") ตอบค่าในไฟล์นั้นแทนที่จะเป็นรายการในคลังทั้งหมด"""
    assert ns["_is_document_reading_question"]("ในคลังมีเอกสารอะไรบ้าง") is False
    assert ns["_is_document_reading_question"]("มีรายการอะไรบ้าง") is True


def test_library_query_with_active_document_still_routes_to_library(ns, box_doc):
    """"ในคลังมีเอกสารอะไรบ้าง" ขณะมีไฟล์เปิดอยู่ (box_doc) ต้องไปถึงเส้นทาง library ได้จริง — ไม่ใช่แค่ไม่ถูกจับ
    เป็น reading question เฉยๆ แต่ _classify_question_route (ซึ่งเป็นตัวตัดสินสุดท้ายใน ask_question หลังข้าม
    ตัวแยก intent ของไฟล์ที่เปิดอยู่ไปแล้ว — ดู library_route_query ใน ask_question) ก็ต้องตอบ library ด้วย"""
    assert ns["_is_document_reading_question"]("ในคลังมีเอกสารอะไรบ้าง") is False
    assert ns["_looks_like_library_query"]("ในคลังมีเอกสารอะไรบ้าง") is True


# ---------- feat/library-search-and-api-errors ส่วน B: ค้นคลังเอกสารจากทุกข้อมูล (search_library) ----------
# เหตุการณ์จริงที่แก้: "SSF-02" -> AI ตอบ "ในระบบไม่พบเอกสารที่มีรหัส SSF-02" ทั้งที่มีจริง เพราะมีเครื่องมือคลัง
# เดิม 5 ตัวที่รู้แค่ข้อมูลระบุตัวตนคงที่ ต้องเดาว่าใช้ตัวไหน — ไฟล์จำลอง 3 ไฟล์: ใบเสนอราคาป่าไม้เขียว (document_no
# "SSF-02", มีรายการ "Rack case server 9U", ยอดรวม 62,220.50), ใบประเมินป่าไม้เขียว (document_no "Kj-01",
# customer_code เดียวกับไฟล์แรก "ASA-02" — ลูกค้าเดียวกันมี 2 ไฟล์), ห้างทองร่ำรวย (document_no "LA-89")

@pytest.fixture
def library_fixture(ns):
    fake_db = ns["_fake_db"]
    fake_db.add_file_category(1, "ใบเสนอราคา")
    fake_db.add_file_category(2, "ใบประเมิน")

    fake_db.add_library_file(1, "ใบเสนอราคาป่าไม้เขียว.xlsx", "lib/1.xlsx", category_id=1)
    fake_db.add_library_file_meta(1, customer_code="ASA-02", customer_name="ป่าไม้เขียว", document_no="SSF-02")
    fake_db.add_library_search_card(
        1,
        "ชื่อไฟล์: ใบเสนอราคาป่าไม้เขียว.xlsx\nหมวดหมู่: ใบเสนอราคา\nชื่อลูกค้า: ป่าไม้เขียว\nรหัสลูกค้า: ASA-02\n"
        "เลขที่เอกสาร: SSF-02\nรายการ: Rack case server 9U จำนวน 1 ตู้ ราคา/หน่วย 62,220.50\n"
        "รวมราคาทั้งสิ้น: 62,220.50",
        total_amount=62220.50, issue_date="2560-01-15",
    )

    fake_db.add_library_file(2, "ใบประเมินป่าไม้เขียว.xlsx", "lib/2.xlsx", category_id=2)
    fake_db.add_library_file_meta(2, customer_code="ASA-02", customer_name="ป่าไม้เขียว", document_no="Kj-01")
    fake_db.add_library_search_card(
        2,
        "ชื่อไฟล์: ใบประเมินป่าไม้เขียว.xlsx\nหมวดหมู่: ใบประเมิน\nชื่อลูกค้า: ป่าไม้เขียว\nรหัสลูกค้า: ASA-02\n"
        "เลขที่เอกสาร: Kj-01",
    )

    fake_db.add_library_file(3, "ห้างทองร่ำรวย.xlsx", "lib/3.xlsx", category_id=1)
    fake_db.add_library_file_meta(3, customer_code="RRR-01", customer_name="ห้างทองร่ำรวย", document_no="LA-89")
    fake_db.add_library_search_card(
        3,
        "ชื่อไฟล์: ห้างทองร่ำรวย.xlsx\nหมวดหมู่: ใบเสนอราคา\nชื่อลูกค้า: ห้างทองร่ำรวย\nรหัสลูกค้า: RRR-01\n"
        "เลขที่เอกสาร: LA-89",
    )
    return fake_db


@pytest.mark.parametrize("query", ["SSF-02", "ssf-02", "เลขเอกสาร SSF-02"])
def test_search_library_exact_document_no_matches_single_file(ns, library_fixture, query):
    results = ns["search_library"](query, user_id=1)
    assert len(results) == 1
    assert results[0]["file_id"] == 1
    assert results[0]["exact"] is True


def test_search_library_exact_customer_code_matches_both_files(ns, library_fixture):
    """"ASA-02" เป็นรหัสลูกค้าของทั้งใบเสนอราคาและใบประเมินป่าไม้เขียว -> ต้องได้ผลตรงเป๊ะ 2 ไฟล์ (ไม่ใช่ไฟล์เดียว)
    ให้ ask_question() แสดงรายการให้เลือกแทนการเปิดไฟล์ไหนไฟล์หนึ่งไปเลย"""
    results = ns["search_library"]("ASA-02", user_id=1)
    exact_file_ids = {r["file_id"] for r in results if r["exact"]}
    assert exact_file_ids == {1, 2}


def test_search_library_ranks_item_match_first_by_bm25(ns, library_fixture):
    """"ใบที่มี Rack 9U" ไม่ตรงเป๊ะกับฟิลด์ระบุตัวตนไหนเลย แต่ต้องเจอใบเสนอราคาป่าไม้เขียวเป็นอันดับ 1 ผ่าน BM25
    บน card_text (มีบรรทัด "รายการ: Rack case server 9U ...") matched_on ต้องโยงไปที่บรรทัดรายการนั้นด้วย"""
    results = ns["search_library"]("ใบที่มี Rack 9U", user_id=1)
    assert results
    assert results[0]["file_id"] == 1
    assert results[0]["exact"] is False
    assert "Rack" in results[0]["matched_on"]


def test_search_library_min_total_filter_excludes_lower_amount_files(ns, library_fixture):
    """filters min_total=60000 ต้องเหลือเฉพาะไฟล์ที่ total_amount >= 60,000 (ใบเสนอราคาป่าไม้เขียว 62,220.50
    เท่านั้น — อีก 2 ไฟล์ไม่มี total_amount เลย ถือว่าไม่ผ่านเงื่อนไขนี้)"""
    results = ns["search_library"]("ป่าไม้เขียว", user_id=1, filters={"min_total": 60000})
    assert {r["file_id"] for r in results} == {1}


def test_search_library_no_match_returns_empty(ns, library_fixture):
    """"แนะนำสาย CAT-6" ไม่มีอยู่ในคลังเลย (ไม่ตรงเป๊ะ ไม่มีคำไหนปนกับ card_text ไฟล์ไหนเลย) ต้องได้ผลว่างเปล่า —
    ปล่อยให้ routing เดิมตัดสินเป็น advisor ต่อไป (ดู test_looks_like_library_query_does_not_match_product_model_
    codes ที่ยืนยันแล้วว่า _looks_like_library_query เองก็คืน False สำหรับคำถามแบบนี้อยู่แล้วตั้งแต่ก่อนถึงจุดนี้)"""
    results = ns["search_library"]("แนะนำสาย CAT-6", user_id=1)
    assert results == []


def test_build_library_search_card_survives_uncomputable_values(ns, doc1):
    """ไฟล์ที่ label เป็นสูตรคำนวณไม่ได้ (เช่น BAHTTEXT ที่ pycel ไม่รองรับ — ดู fixture doc1/build_fixture_
    workbook) ต้องยังสร้างบัตรได้สำเร็จไม่ crash (ใช้เท่าที่มี — ข้าม label ที่ value เป็น None)"""
    fake_db = ns["_fake_db"]
    doc = fake_db.documents[1]
    fake_db.add_library_file(99, doc["filename"], "lib/99.xlsx", label_map=doc["label_map"])
    fake_db.add_storage_blob("lib/99.xlsx", doc["original_bytes"])

    card = ns["build_library_search_card"](99)
    assert card is not None
    assert "ชื่อไฟล์" in card["card_text"]


def test_build_library_search_card_returns_none_for_missing_file(ns):
    assert ns["build_library_search_card"](9999) is None


def test_library_route_system_prompt_forbids_reference_data_wording(ns):
    """system prompt เส้นทาง library ต้องไม่มีกติกาแนว "ใช้ข้อมูลอ้างอิงเป็นหลัก"/เนื้อหากฎหมาย พ.ร.บ.ต่างด้าว
    (ของเส้นทาง legal) ปนมาเลย แต่ต้องสั่งห้ามใช้คำว่า "ข้อมูลอ้างอิง" กับผู้ใช้แทน (คำนี้ปรากฏได้แค่ในประโยคห้าม
    เท่านั้น ไม่ใช่ใช้เป็นกติกาเชิงบวกแบบเส้นทาง legal) — เหตุการณ์จริงที่แก้: "เลขเอกสาร SSF-02" เคยได้คำตอบ
    "ข้อมูลอ้างอิงที่ให้มาไม่มี..." เพราะ system prompt เดิมของเส้นทางกฎหมายถูกใช้ร่วมกันทุกเส้นทาง"""
    prompt = ns["LIBRARY_ROUTE_SYSTEM_PROMPT"]
    assert "ห้ามใช้คำว่า \"ข้อมูลอ้างอิง\"" in prompt
    assert "พ.ร.บ.การประกอบธุรกิจของคนต่างด้าว" not in prompt
    assert "ใช้ข้อมูลนั้นเป็นหลัก" not in prompt


def test_advisor_answer_includes_search_library_tool_only_when_logged_in(ns):
    """SEARCH_LIBRARY_TOOL เข้าถึงข้อมูลส่วนตัว (เหมือน LIBRARY_TOOLS ทุกตัว) ต้องไม่ส่งให้ guest เห็นเลย"""
    ns["_fake_agentic_loop"].set_response("คำตอบ", [], False)
    ns["advisor_answer"]("แนะนำตู้ Rack 9U ที่ดีที่สุด", user_id=None, chat_id=None, active_docs=[])
    call_args = ns["_fake_agentic_loop"].last_call_args
    assert ns["SEARCH_LIBRARY_TOOL"] not in call_args["tools_override"]


def test_looks_like_library_query_false_for_advisor_product_question(ns):
    """"แนะนำตู้ Rack 9U ที่ดีที่สุด" ต้องไม่ถูกตัดสินเป็น library ด้วยโค้ดเลย (ไม่มีรหัสเอกสาร/คำเชิงคลังปนมา) —
    ห้ามเปิดไฟล์อัตโนมัติจากคะแนน BM25/embedding อย่างเดียว (กัน "แนะนำตู้ Rack 9U" ไปเปิดใบป่าไม้เขียวที่มี
    "Rack case server 9U" อยู่ในรายการ โดยไม่ได้ตั้งใจ — ดู B4 ใน ask_question เอง ที่เช็คเงื่อนไขนี้ก่อนเรียก
    search_library เสมอ)"""
    assert ns["_looks_like_library_query"]("แนะนำตู้ Rack 9U ที่ดีที่สุด") is False


def test_resolve_library_exact_match_dispatch_opens_single_exact_match(ns, library_fixture):
    result = ns["_resolve_library_exact_match_dispatch"]("SSF-02", user_id=1, chat_id=1)
    assert result is not None
    assert result["document_id"] is not None
    assert "เปิดไฟล์" in result["answer"]
    active_docs = ns["list_active_editable_documents_by_chat"](1, 1)
    assert any(d["id"] == result["document_id"] for d in active_docs)


def test_resolve_library_exact_match_dispatch_reports_already_open(ns, library_fixture):
    """เรียกซ้ำสองครั้งด้วย query เดิม ("SSF-02") — รอบสองต้องไม่สร้าง EditableDocument ใหม่ซ้ำ ตอบว่าเปิดอยู่แล้ว
    แทน (พฤติกรรมเดิมของ execute_open_library_file — ใช้ตรงๆ ไม่เขียน logic ใหม่)"""
    first = ns["_resolve_library_exact_match_dispatch"]("SSF-02", user_id=1, chat_id=1)
    second = ns["_resolve_library_exact_match_dispatch"]("SSF-02", user_id=1, chat_id=1)
    assert second["document_id"] == first["document_id"]
    assert "เปิดอยู่ในแชทนี้แล้ว" in second["answer"]


def test_resolve_library_exact_match_dispatch_lists_multiple_exact_matches(ns, library_fixture):
    """"ASA-02" ตรงเป๊ะ 2 ไฟล์ (รหัสลูกค้าเดียวกัน) -> ต้องได้รายการให้เลือก ไม่เปิดไฟล์ไหนไปเลย"""
    result = ns["_resolve_library_exact_match_dispatch"]("ASA-02", user_id=1, chat_id=1)
    assert result is not None
    assert result["document_id"] is None
    assert "ต้องการเปิดข้อไหนครับ" in result["answer"]
    listing = ns["_fake_db"].get_chat_last_listing(1)
    assert listing is not None
    assert listing["kind"] == "documents"
    assert len(listing["items"]) == 2
    assert ns["list_active_editable_documents_by_chat"](1, 1) == []


def test_resolve_library_exact_match_dispatch_none_for_advisor_product_question(ns, library_fixture):
    """"แนะนำตู้ Rack 9U ที่ดีที่สุด" ต้องคืน None (ไม่ผ่านด่าน _looks_like_library_query เลย) ไม่เปิดไฟล์ป่าไม้
    เขียวอัตโนมัติทั้งที่มันมี "Rack case server 9U" อยู่ในรายการ (คะแนน BM25/embedding สูงแค่ไหนก็ไม่เปิดเอง)"""
    result = ns["_resolve_library_exact_match_dispatch"]("แนะนำตู้ Rack 9U ที่ดีที่สุด", user_id=1, chat_id=1)
    assert result is None
    assert ns["list_active_editable_documents_by_chat"](1, 1) == []


def test_resolve_library_exact_match_dispatch_none_when_no_exact_match(ns, library_fixture):
    assert ns["_resolve_library_exact_match_dispatch"]("แนะนำสาย CAT-6", user_id=1, chat_id=1) is None


# ---------- feat/library-search-and-api-errors ส่วนต่อ: รหัสเอกสารติดอักษรไทย (ไม่เว้นวรรค) ----------
# บั๊ก: \b ใน Python regex นับอักษรไทยเป็น \w ด้วย -> ไม่มีขอบเขต \w ระหว่างรหัสอังกฤษ/ตัวเลขกับอักษรไทยที่ติดกัน
# "เปิดใบSSF-02"/"เลขเอกสารSSF-02"/"SSF-02ครับ" (คนไทยพิมพ์ติดกันบ่อยมาก) จึงไม่เจอทั้งที่แบบเว้นวรรคเจอปกติ —
# แก้ด้วย lookaround ที่เช็คเฉพาะอักษรอังกฤษ/ตัวเลข (_LIBRARY_DOCUMENT_CODE_RE, _library_search_exact_match_fields)

def test_query_matches_document_code_context_true_for_code_touching_polite_word(ns):
    """"SSF-02ครับ" (รหัสติดคำลงท้ายสุภาพ ไม่มีช่องว่างคั่น) -> ตัดรหัสออกแล้วเหลือ "ครับ" ล้วนๆ -> True เหมือนกับ
    "SSF-02 ครับ" ที่เว้นวรรค"""
    assert ns["_query_matches_document_code_context"]("SSF-02ครับ") is True


@pytest.mark.parametrize("query", ["เปิดใบSSF-02", "เลขเอกสารSSF-02", "SSF-02ครับ", "ssf-02"])
def test_search_library_exact_match_works_when_code_touches_thai_text(ns, library_fixture, query):
    results = ns["search_library"](query, user_id=1)
    exact_file_ids = {r["file_id"] for r in results if r["exact"]}
    assert exact_file_ids == {1}


@pytest.mark.parametrize("query", ["เปิดใบSSF-02", "เลขเอกสารSSF-02", "SSF-02ครับ", "ssf-02"])
def test_resolve_library_exact_match_dispatch_opens_when_code_touches_thai_text(ns, library_fixture, query):
    result = ns["_resolve_library_exact_match_dispatch"](query, user_id=1, chat_id=1)
    assert result is not None
    assert result["document_id"] is not None


def test_resolve_library_exact_match_dispatch_lists_multiple_when_code_touches_thai_text(ns, library_fixture):
    """"ASA-02ครับ" (รหัสลูกค้าติดคำลงท้ายสุภาพ ไม่เว้นวรรค) ต้องยังตรงเป๊ะ 2 ไฟล์เหมือน "ASA-02" เดี่ยวๆ"""
    result = ns["_resolve_library_exact_match_dispatch"]("ASA-02ครับ", user_id=1, chat_id=1)
    assert result is not None
    assert result["document_id"] is None
    assert "ต้องการเปิดข้อไหนครับ" in result["answer"]


def test_resolve_library_exact_match_dispatch_none_for_product_code_touching_thai_text(ns, library_fixture):
    """"แนะนำสายCAT-6ที่ดีที่สุด" (รหัสรุ่นสินค้าติดอักษรไทย ไม่เว้นวรรค) ต้องยังไม่ถูกตัดสินเป็น library เลย
    เหมือนเวอร์ชันเว้นวรรค ไม่เปิดไฟล์อัตโนมัติจากคะแนน BM25/embedding"""
    assert ns["_looks_like_library_query"]("แนะนำสายCAT-6ที่ดีที่สุด") is False
    result = ns["_resolve_library_exact_match_dispatch"]("แนะนำสายCAT-6ที่ดีที่สุด", user_id=1, chat_id=1)
    assert result is None


@pytest.mark.parametrize("query", ["RSSF-02", "SSF-021"])
def test_search_library_exact_match_does_not_match_longer_adjacent_codes(ns, library_fixture, query):
    """"RSSF-02"/"SSF-021" มีตัวอักษร/ตัวเลขภาษาอังกฤษติดกับ "SSF-02" จริงๆ (ไม่ใช่อักษรไทย) -> ต้องไม่ถือว่าตรงเป๊ะ
    กับ SSF-02 (lookaround กันเฉพาะอักษรไทยที่ติดกัน ต้องไม่กันหลวมจนจับรหัสอื่นที่ยาว/สั้นกว่าผิด)"""
    results = ns["search_library"](query, user_id=1)
    exact_file_ids = {r["file_id"] for r in results if r["exact"]}
    assert 1 not in exact_file_ids


# ---------- fix/library-search-precision ----------
# ผลทดสอบจริงหลัง #38: "ใบที่มี Rack 9U" ได้ผล 4 ไฟล์ทั้งที่มี Rack 9U จริงแค่ไฟล์เดียว (ทุกไฟล์ในคลังเป็นงาน
# CCTV คะแนน embedding ใกล้กันหมด) และ "ในคลังมีเอกสารอะไรบ้าง" ได้รายชื่อบริษัทแทนรายชื่อเอกสาร

def test_search_library_relative_cutoff_drops_lookalikes_with_only_embedding_signal(ns, library_fixture):
    """เพิ่มไฟล์ CCTV อีก 2 ไฟล์ที่ไม่มีคำว่า Rack/9U เลยสักคำ (card_text ไม่ทับกับคำค้นเลย — bm25=0) แต่จำลองให้
    embedding score ใกล้เคียงไฟล์ SSF-02 มาก (เหมือนเหตุการณ์จริง: ทุกไฟล์เป็นงานประเภทเดียวกันหมด) — override
    get_library_search_vector_scores_for_all ชั่วคราวใน ns (namespace สดใหม่ต่อเทสเสมอ ไม่กระทบเทสอื่น) เพราะ
    _FakeEmbedModel คืนเวกเตอร์คงที่เสมอ ไม่สะท้อนความคล้ายจริงได้ ต้องเหลือเฉพาะไฟล์ที่มี Rack ในบรรทัดรายการจริง
    (ตัดสัมพัทธ์ _LIBRARY_SEARCH_RELATIVE_CUTOFF ตัดไฟล์ CCTV ทิ้งแม้ผ่าน threshold สัมบูรณ์ได้สบายๆ)"""
    fake_db = ns["_fake_db"]
    fake_db.add_library_file(10, "ใบเสนอราคากล้องวงจรปิดเอ.xlsx", "lib/10.xlsx")
    fake_db.add_library_file_meta(10, customer_code="CCTV-01", customer_name="กล้องวงจรปิดเอ")
    fake_db.add_library_search_card(
        10,
        "ชื่อไฟล์: ใบเสนอราคากล้องวงจรปิดเอ.xlsx\nหมวดหมู่: ใบเสนอราคา\nชื่อลูกค้า: กล้องวงจรปิดเอ\n"
        "รหัสลูกค้า: CCTV-01\nรายการ: กล้องวงจรปิด 4MP จำนวน 8 ตัว",
    )
    fake_db.add_library_file(11, "ใบเสนอราคากล้องวงจรปิดบี.xlsx", "lib/11.xlsx")
    fake_db.add_library_file_meta(11, customer_code="CCTV-02", customer_name="กล้องวงจรปิดบี")
    fake_db.add_library_search_card(
        11,
        "ชื่อไฟล์: ใบเสนอราคากล้องวงจรปิดบี.xlsx\nหมวดหมู่: ใบเสนอราคา\nชื่อลูกค้า: กล้องวงจรปิดบี\n"
        "รหัสลูกค้า: CCTV-02\nรายการ: เครื่องบันทึก NVR 16 ช่อง",
    )

    ns["get_library_search_vector_scores_for_all"] = (
        lambda query_embedding: {1: 0.90, 2: 0.85, 3: 0.80, 10: 0.88, 11: 0.87}
    )

    results = ns["search_library"]("ใบที่มี Rack 9U", user_id=1)
    assert {r["file_id"] for r in results} == {1}


def test_search_library_contains_item_filter_only_matches_item_line(ns, library_fixture):
    """filters contains_item="Rack 9U" ต้องเหลือเฉพาะ SSF-02 (ไฟล์เดียวที่มี "Rack"/"9U" อยู่ในบรรทัดรายการจริง —
    ไม่ใช่แค่บังเอิญมีคำไปตรงส่วนอื่นของบัตร เช่น ชื่อไฟล์/สรุป)"""
    results = ns["search_library"]("ป่าไม้เขียว", user_id=1, filters={"contains_item": "Rack 9U"})
    assert {r["file_id"] for r in results} == {1}


def test_score_library_search_candidates_assigns_semantic_match_type_without_keyword_overlap(ns):
    """ผลที่คะแนนมาจาก embedding ล้วนๆ (ไม่มีคำจาก query ทับกับ card_text เลยสักคำ — keyword_overlap=False) ต้อง
    ได้ match_type="semantic" ไม่ใช่ "keyword" (ทดสอบฟังก์ชันล้วนๆ ไม่ต้องพึ่ง BM25Okapi/embed_model จริง)"""
    candidates = [{
        "file_id": 1, "filename": "a.xlsx", "customer_name": "ลูกค้า1", "document_no": None,
        "customer_code": None, "project_code": None, "card_text": "ชื่อไฟล์: a.xlsx",
    }]
    results = ns["_score_library_search_candidates"](
        "คำถามทั่วไปที่ไม่มีคำตรงกันเลย", candidates,
        bm25_scores={1: 0.0}, embedding_scores={1: 1.0}, keyword_overlap={1: False},
    )
    assert len(results) == 1
    assert results[0]["match_type"] == "semantic"


def test_score_library_search_candidates_keyword_top_result_drops_pure_semantic_results(ns):
    """ผลอันดับ 1 เป็น match_type="keyword" -> ผลอื่นที่เป็น "semantic" ล้วนๆ ต้องถูกตัดทิ้งทั้งหมด (มีผลตรงคำจริง
    แล้ว ไม่ต้องเติมผลที่แค่หัวข้อคล้ายกัน)"""
    candidates = [
        {
            "file_id": 1, "filename": "a.xlsx", "customer_name": "ลูกค้า1", "document_no": None,
            "customer_code": None, "project_code": None, "card_text": "รายการ: Rack case server 9U",
        },
        {
            "file_id": 2, "filename": "b.xlsx", "customer_name": "ลูกค้า2", "document_no": None,
            "customer_code": None, "project_code": None, "card_text": "ชื่อไฟล์: b.xlsx",
        },
    ]
    results = ns["_score_library_search_candidates"](
        "ใบที่มี Rack 9U", candidates,
        bm25_scores={1: 1.0, 2: 0.0}, embedding_scores={1: 0.9, 2: 0.85},
        keyword_overlap={1: True, 2: False},
    )
    file_ids = {r["file_id"] for r in results}
    assert file_ids == {1}


def test_resolve_library_full_listing_dispatch_lists_documents_grouped_by_customer(ns, library_fixture):
    """"ในคลังมีเอกสารอะไรบ้าง" ต้องได้ชื่อเอกสารทุกไฟล์ (ไม่ใช่รายชื่อบริษัท) บันทึก listing แบบ kind="documents"
    ให้ตอบเลขข้อเปิดได้ต่อ (เหมือน find_customers/list_customer_documents ฯลฯ)"""
    result = ns["_resolve_library_full_listing_dispatch"]("ในคลังมีเอกสารอะไรบ้าง", chat_id=1)
    assert result is not None
    assert result["document_id"] is None
    for name in ("ใบเสนอราคาป่าไม้เขียว", "ใบประเมินป่าไม้เขียว", "ห้างทองร่ำรวย"):
        assert name in result["answer"]

    listing = ns["_fake_db"].get_chat_last_listing(1)
    assert listing is not None
    assert listing["kind"] == "documents"
    assert len(listing["items"]) == 3


def test_resolve_library_full_listing_dispatch_separates_templates(ns, library_fixture):
    """ไฟล์แม่แบบ (is_template=True) ต้องไม่ปนอยู่ในรายการเอกสารหลัก ไม่นับรวมในเลขข้อที่บันทึกไว้ให้เลือก แต่ยัง
    แสดงชื่อแยกต่างหากท้ายข้อความในหัวข้อของตัวเอง"""
    fake_db = ns["_fake_db"]
    fake_db.add_library_file(50, "แม่แบบใบเสนอราคา.xlsx", "lib/50.xlsx")
    fake_db.add_library_file_meta(50, is_template=True)

    result = ns["_resolve_library_full_listing_dispatch"]("ในคลังมีเอกสารอะไรบ้าง", chat_id=1)
    assert result is not None
    assert "แม่แบบใบเสนอราคา" in result["answer"]
    assert "**ไฟล์แม่แบบ**" in result["answer"]

    listing = ns["_fake_db"].get_chat_last_listing(1)
    assert len(listing["items"]) == 3
    assert all(item["file_id"] != 50 for item in listing["items"])


def test_resolve_library_full_listing_dispatch_shows_customer_counts_when_library_is_large(ns):
    """คลังมีเอกสาร (ไม่รวมแม่แบบ) เกิน _LIBRARY_FULL_LISTING_SIZE_CUTOFF (30) ไฟล์ -> ต้องแสดงจำนวนเอกสารแยกตาม
    บริษัทแทนการแจกแจงทีละไฟล์ พร้อมขอให้ระบุชื่อบริษัท บันทึก listing แบบ kind="customers" แทน"""
    fake_db = ns["_fake_db"]
    for i in range(31):
        file_id = 200 + i
        fake_db.add_library_file(file_id, f"เอกสาร{i}.xlsx", f"lib/{file_id}.xlsx")
        fake_db.add_library_file_meta(file_id, customer_code=f"CUS-{i:02d}", customer_name=f"ลูกค้า{i}")

    result = ns["_resolve_library_full_listing_dispatch"]("ในคลังมีเอกสารอะไรบ้าง", chat_id=1)
    assert result is not None
    assert result["document_id"] is None
    assert "31 ไฟล์" in result["answer"]
    assert "ระบุชื่อบริษัท" in result["answer"]

    listing = fake_db.get_chat_last_listing(1)
    assert listing["kind"] == "customers"
    assert len(listing["items"]) == 31


def test_resolve_library_full_listing_dispatch_none_for_unrelated_query(ns, library_fixture):
    """"เอกสารของป่าไม้เขียว" ไม่ใช่คำขอรายชื่อเอกสารทั้งหมดในคลัง -> ต้องคืน None ปล่อยให้ routing ทางอื่นตัดสินต่อ"""
    assert ns["_looks_like_full_library_listing_query"]("เอกสารของป่าไม้เขียว") is False
    assert ns["_resolve_library_full_listing_dispatch"]("เอกสารของป่าไม้เขียว", chat_id=1) is None
